"""개발 시드 — 슬라이드 36 데이터셋 (DB 검증 verify_runtime.sql 과 동일 체계).

멱등: tenant_code='nova' 가 있으면 건너뜀. 실 데이터 이행 시 이 시드는 비활성화.
"""
from __future__ import annotations

import hashlib
import json
import logging

from app.db import get_pool

logger = logging.getLogger("edim.seed")

TENANT = "nova"

KOF_ITEMS = [
    ("A", "Fan Model", ["KAD", "KFD", "KAP", "KFS", "ECE", "ECB"]),
    ("B", "Fan Size", ["13", "21", "32"]),
    ("C", "Material", ["32", "45"]),
    ("D", "Bearing Type", ["1.0505", "2.1005", "3.1010", "4.1510"]),
    ("E", "Motor", ["15", "21"]),
    ("F", "Sensor", ["1.Air Vol."]),
]

PRODUCT_CODES = [
    ("KDCR 3-13", "Fan 원심 Casing Double"),
    ("KAD 900 FW", "SF Fan"),
    ("KDP 1-21", "Fan 원심 Casing"),
    ("H 22 380", "Motor"),
    ("ECC 55", "Cooling Coil"),
    ("EFP 55", "Filter"),
    ("EHC 55", "Heating Coil"),
    ("EMX 55", "Mixing Box"),
    ("KDP 9", "Bearing Unit"),
    ("FDV-480", "Motor H22 380V"),
    ("KDC-1", "Casing 강판"),
    ("EWT-3", "Water Trap"),
]

# (mother, child, qty, sort, slot_map[(child_slot, mother_slot|None, fixed|None)])
RELATIONSHIPS = [
    ("KDCR 3-13", "KAD 900 FW", 1, 1, [("B", "B", None), ("E", "E", None)]),
    ("KAD 900 FW", "KDP 1-21", 1, 1, [("B", "B", None), ("E", "E", None)]),
    ("KAD 900 FW", "H 22 380", 1, 2, [("V", None, "380V")]),
    ("KAD 900 FW", "KDP 9", 4, 3, []),
    ("KDCR 3-13", "ECC 55", 1, 2, [("C", "C", None)]),
    ("KDCR 3-13", "EFP 55", 3, 3, []),
    ("KDCR 3-13", "EHC 55", 1, 4, [("C", "C", None)]),
    ("KDCR 3-13", "EMX 55", 1, 5, []),
]

# (code, source, supplier|None, price, from, to)
PRICES = [
    ("KDP 1-21", "QUOTE", "효성", 2_400_000, "2026-01-01", None),
    ("H 22 380", "QUOTE", "효성", 2_400_000, "2026-01-01", None),
    ("ECC 55", "QUOTE", "중원", 1_180_000, "2026-01-01", None),
    ("EFP 55", "QUOTE", "중원", 420_000, "2026-01-01", None),
    ("EHC 55", "QUOTE", "중원", 960_000, "2026-01-01", None),
    ("EMX 55", "QUOTE", "중원", 310_000, "2026-01-01", None),
    ("KDP 9", "QUOTE", "대신금속", 96_000, "2026-01-01", None),
    ("FDV-480", "QUOTE", "효성", 450_000, "2026-06-01", None),
    ("FDV-480", "QUOTE", "LG", 462_000, "2026-05-15", "2026-05-31"),
    ("KDC-1", "PURCHASE", "중원", 128_000, "2026-04-01", "2026-06-30"),
    ("KDC-1", "STOCK", None, 131_000, "2026-07-01", None),
    ("EWT-3", "APPLIED", "대신금속", 36_000, "2026-06-10", None),
]

TECH_ROWS = [
    ("710", {"pd": 11, "pt": 111, "rpm": 1000, "eff": 82, "power": 40, "sound": 57}),
    ("800", {"pd": 9, "pt": 109, "rpm": 980, "eff": 84, "power": 43, "sound": 58}),
    ("900", {"pd": 7, "pt": 107, "rpm": 910, "eff": 72, "power": 44, "sound": 59}),
    ("1000", {"pd": 6, "pt": 104, "rpm": 860, "eff": 68, "power": 47, "sound": 61}),
    ("1120", {"pd": 5, "pt": 101, "rpm": 800, "eff": 63, "power": 52, "sound": 63}),
]

SUPPLIERS = ["효성", "LG", "중원", "대신금속"]


def _ensure_dev_table(cur) -> None:
    """개발서버 운영 도구 — 운영자 요구사항 접수함 (54-테이블 설계 스키마 외, 멱등 생성)."""
    cur.execute("""CREATE TABLE IF NOT EXISTS dev_requirement (
        req_id      SERIAL PRIMARY KEY,
        tenant_id   INT NOT NULL,
        screen_id   VARCHAR(50),
        category    VARCHAR(20)  NOT NULL DEFAULT 'CHANGE',
        title       VARCHAR(200) NOT NULL,
        content     TEXT         NOT NULL DEFAULT '',
        priority    VARCHAR(10)  NOT NULL DEFAULT 'P2',
        status      VARCHAR(20)  NOT NULL DEFAULT 'OPEN',
        requester   VARCHAR(50)  NOT NULL,
        resolution  TEXT,
        created_at  TIMESTAMPTZ  NOT NULL DEFAULT now(),
        resolved_at TIMESTAMPTZ)""")
    cur.execute("""CREATE TABLE IF NOT EXISTS dev_requirement_image (
        image_id    SERIAL PRIMARY KEY,
        req_id      INT NOT NULL REFERENCES dev_requirement(req_id) ON DELETE CASCADE,
        file_name   VARCHAR(200) NOT NULL,
        file_path   VARCHAR(300) NOT NULL,
        file_size   INT NOT NULL,
        content_type VARCHAR(60) NOT NULL DEFAULT 'image/png',
        created_at  TIMESTAMPTZ NOT NULL DEFAULT now())""")


# ── seed v15 — B17: 부품 마스터(prt_part)·도면 BOM(dwg_bom)·공급자 코드(prt_supplier_code_map)·슬롯 정의 ──

# (part_no, name, spec, material_code, supplier, product_main_code, unit, weight, is_standard)
PARTS_V15 = [
    ('PRT-CAS-900', 'Casing Φ900', 'Fan 원심 Casing 용접구조', 'SPHC', None, 'KDCR 3-13', 'EA', 42.5, False),
    ('PRT-IMP-900', 'Impeller Airfoil 900', '밸런싱 등급 G6.3', 'SS400', '효성', 'KDCR 3-13', 'EA', 18.2, False),
    ('PRT-SHF-045', 'Shaft Φ45', '연마 h6', 'SS400', '중원', None, 'EA', 12.0, True),
    ('PRT-BRG-6210', 'Bearing 6210', 'Deep Groove 2RS', None, '중원', None, 'EA', 0.5, True),
]

# (part_no, qty, assembly_seq, assembly_note) — KDCR 3-13 조립순서 ◆
BOM_V15 = [
    ('PRT-BRG-6210', 2, 1, 'Bearing 압입 (양단)'),
    ('PRT-SHF-045', 1, 2, '축 조립·수평 확인'),
    ('PRT-IMP-900', 1, 5, 'Impeller 밸런싱 후 체결'),
    ('PRT-CAS-900', 1, 6, 'Casing 최종 조립'),
]

# (part_no, supplier, supplier_code, supplier_name)
SUPPLIER_CODES_V15 = [
    ('PRT-BRG-6210', '중원', 'JW-6210-2RS', 'Deep Groove Ball Bearing 6210'),
    ('PRT-IMP-900', '효성', 'HS-IMP-900A', 'Airfoil Impeller 900'),
]

# (item_slot, source item_name, is_required, sort) — 'KDP 1-21' 필수 슬롯 정의
SLOT_ITEMS_V15 = [
    ('A', 'Fan Model', True, 1),
    ('B', 'Fan Size', True, 2),
    ('C', 'Material', False, 3),
]

LABELS_V15 = {
    'menu.plm-parts': ('Part Ledger (M-4-7)', '部品台帳 (M-4-7)', '零件台账 (M-4-7)'),
    'screen.plm-parts': ('Part Ledger', '部品台帳', '零件台账'),
    'parts.title': ('Part Ledger — prt_part', '部品台帳 — prt_part', '零件台账 — prt_part'),
    'parts.registerF2': ('+ Register Part F2', '＋ 部品登録 F2', '＋ 零件登记 F2'),
    'parts.partNo': ('Part No.', '部品番号', '零件编号'),
    'parts.partName': ('Part Name', '部品名', '零件名'),
    'parts.specCol': ('Spec', '仕様', '规格'),
    'parts.materialCol': ('Material', '材質', '材质'),
    'parts.supplierCol': ('Supplier', '仕入先', '供应商'),
    'parts.weightCol': ('Weight(kg)', '重量(kg)', '重量(kg)'),
    'parts.stdChip': ('STD', 'STD', 'STD'),
    'parts.bomCol': ('BOM Refs', 'BOM 参照', 'BOM 引用'),
    'parts.supCodes': ('Supplier Code Map (ERP-018)', '仕入先コードマップ (ERP-018)', '供应商代码映射 (ERP-018)'),
    'parts.supCodeAdd': ('Add Mapping', 'マッピング追加', '添加映射'),
    'parts.selectHint': ('Select a part row to view supplier code mappings', '部品行を選択すると仕入先コードを表示', '选择零件行以查看供应商代码'),
    'parts.noSupCodes': ('No supplier code mappings', '仕入先コードマッピングなし', '无供应商代码映射'),
    'editor.bomLive': ('Sub Item DWG · Assembly Seq (dwg_bom)', 'Sub Item DWG · 組立順序 (dwg_bom)', 'Sub Item DWG · 装配顺序 (dwg_bom)'),
    'purchase.supCode': ('Supplier Code', '仕入先コード', '供应商代码'),
    'detail.slotDef': ('Required Slot Definition (product_code_item)', '必須スロット定義 (product_code_item)', '必需槽位定义 (product_code_item)'),
}


def seed_v15(cur, tid: int) -> None:
    cur.execute('SELECT 1 FROM prt_part WHERE tenant_id=%s LIMIT 1', (tid,))
    if not cur.fetchone():
        ids: dict[str, int] = {}
        for no, name, spec, mat, sup, pcode, unit, weight, std in PARTS_V15:
            mid = None
            if mat:
                cur.execute('SELECT material_id FROM mat_material WHERE tenant_id=%s AND material_code=%s',
                            (tid, mat))
                m = cur.fetchone()
                mid = m[0] if m else None
            sid = None
            if sup:
                cur.execute('SELECT company_id FROM com_company WHERE tenant_id=%s AND company_name=%s LIMIT 1',
                            (tid, sup))
                s = cur.fetchone()
                sid = s[0] if s else None
            pcid = None
            if pcode:
                cur.execute('SELECT min(product_code_id) FROM product_code WHERE tenant_id=%s AND main_code=%s',
                            (tid, pcode))
                pc = cur.fetchone()
                pcid = pc[0] if pc else None
            cur.execute(
                """INSERT INTO prt_part (tenant_id, part_no, part_name, specification, material_id,
                   supplier_id, product_code_id, unit, weight, is_standard)
                   VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s) RETURNING part_id""",
                (tid, no, name, spec, mid, sid, pcid, unit, weight, std))
            ids[no] = cur.fetchone()[0]
        cur.execute(
            "SELECT drawing_id FROM dwg_drawing WHERE tenant_id=%s AND drawing_no='KDCR 3-13'", (tid,))
        d = cur.fetchone()
        if d:
            for i, (no, qty, seq, note) in enumerate(BOM_V15, start=1):
                cur.execute(
                    """INSERT INTO dwg_bom (drawing_id, part_id, item_no, quantity, assembly_seq, assembly_note)
                       VALUES (%s,%s,%s,%s,%s,%s)""", (d[0], ids[no], i, qty, seq, note))
        for no, sup, scode, sname in SUPPLIER_CODES_V15:
            cur.execute('SELECT company_id FROM com_company WHERE tenant_id=%s AND company_name=%s LIMIT 1',
                        (tid, sup))
            s = cur.fetchone()
            if s:
                cur.execute(
                    """INSERT INTO prt_supplier_code_map (tenant_id, part_id, supplier_id,
                       supplier_code, supplier_name) VALUES (%s,%s,%s,%s,%s)""",
                    (tid, ids[no], s[0], scode, sname))
        logger.info('seed v15 — prt_part %d + bom %d + supplier_code %d',
                    len(PARTS_V15), len(BOM_V15), len(SUPPLIER_CODES_V15))
    cur.execute('SELECT 1 FROM product_code_item LIMIT 1')
    if not cur.fetchone():
        cur.execute(
            "SELECT min(product_code_id) FROM product_code WHERE tenant_id=%s AND main_code='KDP 1-21'",
            (tid,))
        pc = cur.fetchone()
        if pc and pc[0]:
            n = 0
            for slot, item_name, req, sort in SLOT_ITEMS_V15:
                cur.execute('SELECT item_id FROM code_item WHERE tenant_id=%s AND item_name=%s LIMIT 1',
                            (tid, item_name))
                ci = cur.fetchone()
                if ci:
                    cur.execute(
                        """INSERT INTO product_code_item (product_code_id, item_slot, source_item_id,
                           is_required, sort_order) VALUES (%s,%s,%s,%s,%s)""",
                        (pc[0], slot, ci[0], req, sort))
                    n += 1
            logger.info('seed v15 — product_code_item %d', n)
    for key, (en, ja, zh) in LABELS_V15.items():
        for locale, text in (('en', en), ('ja', ja), ('zh', zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v16 — B19: 창고/저장위치 계층 (erp_warehouse, ERP-020/021) ──

# (location_type, code, name, parent_code, hazard, inspection)
WAREHOUSES_V16 = [
    ('REGION', 'KR-CW', '창원 사업장', None, None, None),
    ('PLANT', 'CW-P1', '1공장 (Fan 조립)', 'KR-CW', None, None),
    ('WAREHOUSE', 'P1-WH-A', '자재창고 A', 'CW-P1', None, '6개월'),
    ('STORAGE', 'WH-A-GEN', '일반 자재 보관', 'P1-WH-A', None, '3개월'),
    ('STORAGE', 'WH-A-HAZ', '위험물 보관소', 'P1-WH-A', '액체·가스 (도료·용접가스)', '1개월'),
    ('SECTOR', 'GEN-A01', 'A-01 (표준 부품)', 'WH-A-GEN', None, None),
    ('SECTOR', 'HAZ-H01', 'H-01 (인화성 액체)', 'WH-A-HAZ', '액체', '1개월'),
]

LABELS_V16 = {
    'menu.erp-warehouse': ('Warehouse·Storage (M-8-4)', '倉庫·保管場所 (M-8-4)', '仓库·储位 (M-8-4)'),
    'screen.erp-warehouse': ('Warehouse·Storage', '倉庫·保管場所', '仓库·储位'),
    'wh.title': ('Warehouse/Storage Hierarchy — erp_warehouse', '倉庫/保管場所階層 — erp_warehouse', '仓库/储位层级 — erp_warehouse'),
    'wh.addF2': ('+ Add Location F2', '＋ 位置追加 F2', '＋ 添加位置 F2'),
    'wh.typeCol': ('Type', '種別', '类型'),
    'wh.codeCol': ('Code', 'コード', '代码'),
    'wh.nameCol': ('Name', '名称', '名称'),
    'wh.hazardCol': ('Hazard Allowed', '危険物許可', '危险品许可'),
    'wh.inspectionCol': ('Inspection Cycle', '検査周期', '检查周期'),
    'wh.hierHint': ('Hierarchy: REGION→PLANT→WAREHOUSE→STORAGE→SECTOR (order enforced)',
                    '階層: REGION→PLANT→WAREHOUSE→STORAGE→SECTOR (順序強制)',
                    '层级: REGION→PLANT→WAREHOUSE→STORAGE→SECTOR (强制顺序)'),
    'purch.poDialog': ('PO Terms (ERP-017)', 'PO 条件 (ERP-017)', 'PO 条款 (ERP-017)'),
    'purch.delivery': ('Delivery Terms', '納品条件', '交货条款'),
    'purch.transport': ('Transport', '輸送手段', '运输方式'),
    'purch.minQty': ('Min Order Qty', '最小購買数量', '最小采购量'),
    'purch.cert': ('Certificate Required', '証明書要求', '要求证书'),
}


def seed_v16(cur, tid: int) -> None:
    cur.execute('SELECT 1 FROM erp_warehouse WHERE tenant_id=%s LIMIT 1', (tid,))
    if not cur.fetchone():
        ids: dict[str, int] = {}
        for lt, code, name, parent, hazard, insp in WAREHOUSES_V16:
            cur.execute(
                """INSERT INTO erp_warehouse (tenant_id, parent_id, location_type, location_code,
                   location_name, hazard_allowed, inspection_cycle)
                   VALUES (%s,%s,%s,%s,%s,%s,%s) RETURNING warehouse_id""",
                (tid, ids.get(parent), lt, code, name, hazard, insp))
            ids[code] = cur.fetchone()[0]
        logger.info('seed v16 — erp_warehouse %d노드', len(WAREHOUSES_V16))
    for key, (en, ja, zh) in LABELS_V16.items():
        for locale, text in (('en', en), ('ja', ja), ('zh', zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v17 — B20: Macro 4-Way 필드·CODING 모드·tbx_macro_ref ──

FLOWCHART_V17 = {
    'nodes': [
        {'id': 'f1', 'name': 'Impeller 참조', 'sub': 'Table12', 'x': 20, 'y': 12, 'w': 110, 'h': 34},
        {'id': 'f2', 'name': 'Casing 폭', 'sub': 'Table12', 'x': 20, 'y': 60, 'w': 110, 'h': 34},
        {'id': 'f3', 'name': 'FES 여유', 'sub': 'Var', 'x': 20, 'y': 108, 'w': 110, 'h': 34},
        {'id': 'f4', 'name': 'Σ Shaft 길이', 'x': 150, 'y': 60, 'w': 90, 'h': 34},
    ],
    'edges': [['f1', 'f4'], ['f2', 'f4'], ['f3', 'f4']],
}

CODE_TEXT_V17 = (
    "def shaft_length(mc: float, fes: float) -> float:\n"
    "    base = table12('E', 560, 800, agg='SUM')  # Casing 구간 합\n"
    "    margin = var('FES', fes)                  # 여유 (Variant)\n"
    "    return round((base + margin) * prec(1), 1)\n")

CODING_MACRO_V17 = (
    "def weight_estimate(parts: list[dict]) -> float:\n"
    "    \"\"\"부품 중량 합산 — prt_part.weight × BOM 수량 (CODING 모드 데모).\"\"\"\n"
    "    return sum(p['weight'] * p['qty'] for p in parts if p.get('weight'))\n")


def seed_v17(cur, tid: int) -> None:
    # 4-Way 필드 보강 — Shaft 길이 계산 (미보유 시 1회)
    cur.execute(
        """SELECT macro_id, code_text FROM tbx_macro
           WHERE tenant_id=%s AND macro_name='Shaft 길이 계산'""", (tid,))
    m = cur.fetchone()
    if m and not m[1]:
        cur.execute(
            """UPDATE tbx_macro SET code_text=%s, flowchart_def=%s,
               description_text=COALESCE(description_text,
                 'Casing 구간 Table 합계에 FES 여유를 더해 Shaft 길이를 산출 — MC>500 조건 분기'),
               test_input=%s, test_result=%s WHERE macro_id=%s""",
            (CODE_TEXT_V17, json.dumps(FLOWCHART_V17),
             json.dumps({'MC': 520, 'FES': 15}), json.dumps({'value': 2685, 'ok': True}),
             m[0]))
        logger.info('seed v17 — Shaft 길이 계산 4-Way 필드 보강')
    # CODING 모드 데모 매크로
    cur.execute(
        "SELECT 1 FROM tbx_macro WHERE tenant_id=%s AND apply_type='CODING' LIMIT 1", (tid,))
    if not cur.fetchone():
        cur.execute(
            """INSERT INTO tbx_macro (tenant_id, macro_name, code_text, description_text,
               apply_type, status)
               VALUES (%s,'중량 추정 (Coding)',%s,'BOM 중량 합산 — 복잡 로직은 Coding 모드 (TBX-010)',
                       'CODING','DRAFT')""", (tid, CODING_MACRO_V17))
        logger.info('seed v17 — CODING 모드 매크로 1건')
    # tbx_macro_ref — Table 참조 매크로 (DIM D = Table12 참조)
    cur.execute('SELECT 1 FROM tbx_macro_ref LIMIT 1')
    if not cur.fetchone():
        cur.execute("SELECT table_id FROM tbl_data_table WHERE tenant_id=%s AND table_name='Table12'",
                    (tid,))
        t = cur.fetchone()
        if t:
            n = 0
            cur.execute(
                """SELECT macro_id, macro_expr FROM tbx_macro
                   WHERE tenant_id=%s AND macro_expr ILIKE '%%Table12%%'""", (tid,))
            for mid, _expr in cur.fetchall():
                cur.execute(
                    """INSERT INTO tbx_macro_ref (macro_id, ref_type, ref_target_id)
                       VALUES (%s,'TABLE',%s)""", (mid, t[0]))
                n += 1
            logger.info('seed v17 — tbx_macro_ref %d건 (Table12)', n)


def _seed_invariants(cur, tid: int) -> None:
    """버전 게이트와 무관하게 매 기동 실행되는 불변식 재확정 (자기치유).

    seed_v4 는 멱등(1회)이라, 라이브 테스트·수동 조작이 데모 계정을 강등하면
    복원되지 않는다. 여기서 레벨이 다를 때만 재승격하고, 실제 치유 시에만 감사 남긴다.
    """
    cur.execute(
        """UPDATE sys_user SET user_level='ADMIN', updated_at=now()
           WHERE tenant_id=%s AND login_id='edim' AND user_level<>'ADMIN'
           RETURNING user_id""", (tid,))
    healed = cur.fetchone()
    if healed:
        uid = healed[0]
        cur.execute(
            """INSERT INTO sys_history (tenant_id, target_table, target_id, action,
               actor_id, before_data, after_data)
               VALUES (%s,'sys_user',%s,'LEVEL_CHANGE',%s,
                       '{"level":"(demoted)","source":"seed self-heal"}'::jsonb,
                       '{"level":"ADMIN"}'::jsonb)""", (tid, uid, uid))
        logger.info("seed invariant healed — edim 레벨 → ADMIN 재확정")


def run_seed() -> None:
    pool = get_pool()
    if pool is None:
        logger.warning("seed skipped — DB unavailable")
        return
    with pool.connection() as conn, conn.cursor() as cur:
        # dev_requirement 테이블은 마이그레이션(0002)이 담당 — 시드는 데이터만 (C6)
        cur.execute("SELECT tenant_id FROM sys_tenant WHERE tenant_code=%s", (TENANT,))
        row = cur.fetchone()
        if row:
            logger.info("base seed exists (tenant_id=%s)", row[0])
            _seed_v2(cur, row[0])
            seed_v3(cur, row[0])
            seed_v4(cur, row[0])
            seed_v5(cur, row[0])
            seed_v6(cur, row[0])
            seed_v7(cur, row[0])
            seed_v8(cur, row[0])
            seed_v9(cur, row[0])
            seed_v10(cur, row[0])
            seed_v11(cur, row[0])
            seed_v12(cur, row[0])
            seed_v13(cur, row[0])
            seed_v14(cur, row[0])
            seed_v15(cur, row[0])
            seed_v16(cur, row[0])
            seed_v17(cur, row[0])
            seed_v18(cur, row[0])
            seed_v19(cur, row[0])
            seed_v20(cur, row[0])
            seed_v21(cur, row[0])
            seed_v22(cur, row[0])
            seed_v23(cur, row[0])
            seed_v24(cur, row[0])
            seed_v25(cur, row[0])
            seed_v26(cur, row[0])
            seed_v27(cur, row[0])
            seed_v28(cur, row[0])
            seed_v29(cur, row[0])
            seed_v31(cur, row[0])
            seed_v30(cur, row[0])
            seed_v32(cur, row[0])
            seed_v33(cur, row[0])
            seed_v34(cur, row[0])
            _seed_invariants(cur, row[0])
            return

        cur.execute(
            "INSERT INTO sys_tenant (tenant_code, tenant_name) VALUES (%s,%s) RETURNING tenant_id",
            (TENANT, "NOVA Solution"))
        tid = cur.fetchone()[0]

        pw = hashlib.sha256(b"edim").hexdigest()
        cur.execute(
            """INSERT INTO sys_user (tenant_id, login_id, user_name, password_hash, department, user_level)
               VALUES (%s,'edim','YS.Gang',%s,'기술연구소','SETUP')""", (tid, pw))

        suppliers: dict[str, int] = {}
        for name in SUPPLIERS:
            cur.execute(
                """INSERT INTO com_company (tenant_id, company_type, company_name)
                   VALUES (%s,'SUPPLIER',%s) RETURNING company_id""", (tid, name))
            suppliers[name] = cur.fetchone()[0]
        cur.execute(
            """INSERT INTO com_company (tenant_id, company_type, company_name)
               VALUES (%s,'CUSTOMER','Micron') RETURNING company_id""", (tid,))
        micron = cur.fetchone()[0]

        cur.execute(
            """INSERT INTO prj_project (tenant_id, project_no, project_name, customer_id, project_type)
               VALUES (%s,'PS-61313-5','Micron #7',%s,'Client') RETURNING project_id""",
            (tid, micron))
        project_id = cur.fetchone()[0]

        cur.execute(
            """INSERT INTO code_group (tenant_id, group_code, group_name, group_type,
               hierarchy_address, approval_status)
               VALUES (%s,'KOF','Specification - Fan','SPECIFICATION','/C/ENG/FAN/KOF','APPROVED')
               RETURNING group_id""", (tid,))
        gid = cur.fetchone()[0]

        for slot, name, values in KOF_ITEMS:
            cur.execute(
                """INSERT INTO code_item (tenant_id, group_id, item_slot, item_name)
                   VALUES (%s,%s,%s,%s) RETURNING item_id""", (tid, gid, slot, name))
            item_id = cur.fetchone()[0]
            for i, v in enumerate(values):
                cur.execute(
                    """INSERT INTO code_item_value (tenant_id, item_id, value_code,
                       sort_order, approval_status)
                       VALUES (%s,%s,%s,%s,'APPROVED')""", (tid, item_id, v, i))

        codes: dict[str, int] = {}
        for main, name in PRODUCT_CODES:
            cur.execute(
                """INSERT INTO product_code (tenant_id, main_code, group_id, code_name,
                   hierarchy_address, approval_status)
                   VALUES (%s,%s,%s,%s,%s,'APPROVED') RETURNING product_code_id""",
                (tid, main, gid, name, f"/C/ENG/FAN/{main.split(' ')[0]}"))
            codes[main] = cur.fetchone()[0]

        for mother, child, qty, sort, maps in RELATIONSHIPS:
            cur.execute(
                """INSERT INTO code_relationship (tenant_id, mother_code_id, child_code_id,
                   quantity, sort_order, approval_status)
                   VALUES (%s,%s,%s,%s,%s,'APPROVED') RETURNING rel_id""",
                (tid, codes[mother], codes[child], qty, sort))
            rel_id = cur.fetchone()[0]
            for child_slot, mother_slot, fixed in maps:
                cur.execute(
                    """INSERT INTO code_relationship_slot_map (rel_id, child_slot, mother_slot, fixed_value)
                       VALUES (%s,%s,%s,%s)""", (rel_id, child_slot, mother_slot, fixed))

        for code, source, supplier, price, vfrom, vto in PRICES:
            cur.execute(
                """INSERT INTO cst_price (tenant_id, product_code_id, price_source,
                   supplier_id, price, valid_from, valid_to)
                   VALUES (%s,%s,%s,%s,%s,%s,%s)""",
                (tid, codes[code], source,
                 suppliers.get(supplier) if supplier else None, price, vfrom, vto))

        cur.execute(
            """INSERT INTO tbl_data_table (tenant_id, table_name, table_type, department,
               hierarchy_address, column_def, approval_status)
               VALUES (%s,'FanTechData','TECH','Engineering','/T/ENG/TECH/FanTechData',
                       %s,'APPROVED') RETURNING table_id""",
            (tid, json.dumps({"columns": ["pd", "pt", "rpm", "eff", "power", "sound"]})))
        tech_id = cur.fetchone()[0]
        for i, (key, vals) in enumerate(TECH_ROWS):
            cur.execute(
                """INSERT INTO tbl_data_row (table_id, row_key, row_key_num, row_values, sort_order)
                   VALUES (%s,%s,%s,%s,%s)""", (tech_id, key, key, json.dumps(vals), i))

        cur.execute(
            """INSERT INTO cpq_selection (tenant_id, project_id, finished_goods_code,
               product_code_id, slot_values, status)
               VALUES (%s,%s,'KDCR 3-13-13-15',%s,%s,'DRAFT')""",
            (tid, project_id, codes["KDCR 3-13"],
             json.dumps({"B": "13", "C": "32", "E": "15"})))

        logger.info("seed complete — tenant_id=%s, %d codes", tid, len(codes))
        _seed_v2(cur, tid)
        seed_v3(cur, tid)
        seed_v4(cur, tid)
        seed_v5(cur, tid)
        seed_v6(cur, tid)
        seed_v7(cur, tid)
        seed_v8(cur, tid)
        seed_v9(cur, tid)
        seed_v10(cur, tid)
        seed_v11(cur, tid)
        seed_v12(cur, tid)
        seed_v13(cur, tid)
        seed_v14(cur, tid)
        seed_v15(cur, tid)
        seed_v16(cur, tid)
        seed_v17(cur, tid)
        seed_v18(cur, tid)
        seed_v19(cur, tid)
        seed_v20(cur, tid)
        seed_v21(cur, tid)
        seed_v22(cur, tid)
        seed_v23(cur, tid)
        seed_v24(cur, tid)
        seed_v25(cur, tid)
        seed_v26(cur, tid)
        seed_v27(cur, tid)
        seed_v28(cur, tid)
        seed_v29(cur, tid)
        seed_v31(cur, tid)
        seed_v30(cur, tid)
        seed_v32(cur, tid)
        _seed_invariants(cur, tid)


# ── seed v2 — 승인함·문서함·사용자·프로세스 이벤트·이력 (배치 A) ──

USERS_V2 = [
    ("kim01", "Kim", "기술", "GENERAL", "ACTIVE"),
    ("park.f", "Park", "재무", "ADMIN", "LOCKED"),
    ("lee.t", "Lee", "기술", "GENERAL", "ACTIVE"),
    ("jang.s", "Jang", "영업", "GENERAL", "ACTIVE"),
]

DOCS_V2 = [
    ("DF 342-234 E", "Density 계산서", "TECH_DOC", "CHECK", "KD-0.2", "Kim", None, "S-2"),
    ("DF 342-235 A", "Fan 성능 Report", "TECH_DOC", "APPROVE", "KD-1.0", "Lee", None, "S-3"),
    ("QR-61216-01", "CLT 견적서", "QUOTATION", "ACCEPTED", "1.1", "Jang", "2026-06-30", "S-3"),
    ("I011902", "승인도면 (AHU 5)", "DRAWING", "ACCEPTED", "B", "Kim", "2026-06-27", "S-1"),
    ("DF 342-236", "소음 예측 계산서", "TECH_DOC", "SET_UP", "KD-0.1", "Kim", None, "S-2"),
]

PROC_DEFS_V2 = [
    ("OR", "수주", "영업", False), ("AP", "승인 도서", "영업", True),
    ("PL", "Part List", "기술", True), ("BOM", "BOM", "기술", True),
    ("MR", "제작의뢰", "기술", False), ("PR", "발주 요청", "자재", False),
    ("IR", "기성청구", "재무", False),
]
PROC_EDGES_V2 = [("OR", "AP"), ("OR", "PL"), ("PL", "BOM"), ("BOM", "PR"), ("AP", "MR")]

# (proc, project_no, assignee, due, status)
EVENTS_V2 = [
    ("MR", "PS-61313-5", "edim", "2026-07-09", "IN_PROGRESS"),
    ("PL", "PS-61313-5", "kim01", "2026-07-05", "TODO"),
    ("BOM", "PS-598", "kim01", "2026-07-11", "TODO"),
    ("MR", "PS-612", "edim", "2026-07-07", "TODO"),
    ("IR", "PS-598", "park.f", "2026-07-04", "TODO"),
]


def _seed_v2(cur, tid: int) -> None:
    cur.execute("SELECT 1 FROM doc_control WHERE tenant_id=%s LIMIT 1", (tid,))
    if cur.fetchone():
        logger.info("seed v2 exists — skip")
        return

    for login, name, dept, level, status in USERS_V2:
        cur.execute(
            """INSERT INTO sys_user (tenant_id, login_id, user_name, password_hash,
               department, user_level, status)
               VALUES (%s,%s,%s,%s,%s,%s,%s)
               ON CONFLICT (tenant_id, login_id) DO NOTHING""",
            (tid, login, name, hashlib.sha256(b"edim").hexdigest(), dept, level, status))

    for no, name in [("PS-612", "반도체 배기 개선"), ("PS-598", "물류센터 공조")]:
        cur.execute(
            """INSERT INTO prj_project (tenant_id, project_no, project_name)
               VALUES (%s,%s,%s) ON CONFLICT (tenant_id, project_no) DO NOTHING""",
            (tid, no, name))

    for doc_no, title, dtype, status, ver, person, appdate, grade in DOCS_V2:
        cur.execute(
            """INSERT INTO doc_control (tenant_id, doc_no, title, doc_type,
               released_status, version, person, approval_date, management_grade)
               VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s)""",
            (tid, doc_no, title, dtype, status, ver, person, appdate, grade))

    defs: dict[str, int] = {}
    for code, name, dept, auto in PROC_DEFS_V2:
        cur.execute(
            """INSERT INTO erp_process_def (tenant_id, proc_code, proc_name, department, is_auto)
               VALUES (%s,%s,%s,%s,%s) RETURNING proc_def_id""",
            (tid, code, name, dept, auto))
        defs[code] = cur.fetchone()[0]
    for f, t in PROC_EDGES_V2:
        cur.execute(
            """INSERT INTO erp_process_edge (tenant_id, from_def_id, to_def_id)
               VALUES (%s,%s,%s)""", (tid, defs[f], defs[t]))

    cur.execute("SELECT login_id, user_id FROM sys_user WHERE tenant_id=%s", (tid,))
    uids = dict(cur.fetchall())
    cur.execute("SELECT project_no, project_id FROM prj_project WHERE tenant_id=%s", (tid,))
    pids = dict(cur.fetchall())

    for proc, pno, login, due, status in EVENTS_V2:
        cur.execute(
            """INSERT INTO erp_process_event (tenant_id, proc_def_id, project_id,
               status, assignee_id, due_date)
               VALUES (%s,%s,%s,%s,%s,%s)""",
            (tid, defs[proc], pids[pno], status, uids.get(login), due))

    # 승인 요청 (PENDING) — product_code / doc_control 대상
    cur.execute(
        "SELECT main_code, product_code_id FROM product_code WHERE tenant_id=%s", (tid,))
    cids = dict(cur.fetchall())
    cur.execute(
        "SELECT doc_no, doc_control_id FROM doc_control WHERE tenant_id=%s", (tid,))
    dids = dict(cur.fetchall())
    reqs = [
        ("product_code", cids["KDCR 3-13"], "UPDATE", "검토", "ysgang", "KDCR 3-13 ↔ KDI 21 관계"),
        ("product_code", cids["FDV-480"], "CREATE", "승인", "ysgang", "FDV-480 신규 등록"),
        ("doc_control", dids["DF 342-235 A"], "UPDATE", "승인", "kim01", "Fan 성능 Report v1.0"),
        ("doc_control", dids["DF 342-234 E"], "CREATE", "검토", "lee.t", "Density 계산서 v0.2"),
    ]
    requester_fallback = uids.get("edim") or next(iter(uids.values()))
    for table, target, rtype, step, login, label in reqs:
        cur.execute(
            """INSERT INTO sys_approval_request (tenant_id, target_table, target_id,
               request_type, step, requester_id, comment)
               VALUES (%s,%s,%s,%s,%s,%s,%s)""",
            (tid, table, target, rtype, step,
             uids.get(login, requester_fallback), label))

    actor = requester_fallback
    for table, target, action in [
        ("cpq_run", 1, "RUN_SUCCESS"), ("code_relationship", 1, "UPDATE"),
        ("sys_hierarchy", 1, "MOVE"),
    ]:
        cur.execute(
            """INSERT INTO sys_history (tenant_id, target_table, target_id, action,
               actor_id, after_data)
               VALUES (%s,%s,%s,%s,%s,%s)""",
            (tid, table, target, action, actor, json.dumps({"seed": True})))

    logger.info("seed v2 complete — docs/users/events/approvals")


# ── seed v3 — Table12 (Variant, W-20) ──

TABLE12_ROWS = [
    ("560", {"A": 55, "B": 45, "C": 9, "E": 656}),
    ("630", {"A": 32, "B": 45, "C": 11, "E": 656}),
    ("710", {"A": 679, "B": 760, "C": 45, "D": 700, "E": 656}),
    ("800", {"A": 721, "B": 806, "C": 45, "D": 760, "E": 702}),
]


def seed_v3(cur, tid: int) -> None:
    cur.execute(
        "SELECT 1 FROM tbl_data_table WHERE tenant_id=%s AND table_name='Table12'", (tid,))
    if cur.fetchone():
        return
    cur.execute(
        """INSERT INTO tbl_data_table (tenant_id, table_name, table_type, department,
           hierarchy_address, column_def, approval_status)
           VALUES (%s,'Table12','VARIANT','Engineering','/T/ENG/VARIANT/Table12',
                   %s,'APPROVED') RETURNING table_id""",
        (tid, json.dumps({"columns": ["A", "B", "C", "D", "E"], "keyLabel": "Size"})))
    table_id = cur.fetchone()[0]
    for i, (key, vals) in enumerate(TABLE12_ROWS):
        cur.execute(
            """INSERT INTO tbl_data_row (table_id, row_key, row_key_num, row_values, sort_order)
               VALUES (%s,%s,%s,%s,%s)""", (table_id, key, key, json.dumps(vals), i))
    logger.info("seed v3 complete — Table12 %d rows", len(TABLE12_ROWS))


# ── seed v4 — RBAC 데모(edim→ADMIN) + 알림 샘플 ──

def seed_v4(cur, tid: int) -> None:
    cur.execute("SELECT 1 FROM sys_notification WHERE tenant_id=%s LIMIT 1", (tid,))
    if cur.fetchone():
        return
    cur.execute(
        "UPDATE sys_user SET user_level='ADMIN' WHERE tenant_id=%s AND login_id='edim'", (tid,))
    cur.execute(
        "SELECT user_id FROM sys_user WHERE tenant_id=%s AND login_id='edim'", (tid,))
    uid = cur.fetchone()[0]
    for ntype, title in [
        ("APPROVAL_REQUEST", "승인 요청 — 도면 KDCR 3-13 Rev.B (Kim)"),
        ("DEADLINE_WARN", "기한 경고 — PS-612 MR 제작의뢰 초과 2일"),
    ]:
        cur.execute(
            """INSERT INTO sys_notification (tenant_id, user_id, notify_type, title, link_url)
               VALUES (%s,%s,%s,%s,'/common')""", (tid, uid, ntype, title))
    logger.info("seed v4 complete — edim=ADMIN, 알림 2건")


# ── seed v5 — 치수 정의 이행 (dwg_drawing + tbx_macro + dwg_dimension) ──

DIMS_V5 = [
    # (label, type, variant, macro_expr)
    ("A", "KEY", 670, None),
    ("B", "KEY", None, "=A+56"),
    ("C", "DETAIL", 45, None),
    ("D", "DETAIL", None, "=Table12(B,710)"),
    ("E", "DETAIL", 320, None),
    ("K", "KEY", None, "=A*1.62"),
]


def seed_v5(cur, tid: int) -> None:
    cur.execute(
        "SELECT 1 FROM dwg_drawing WHERE tenant_id=%s AND drawing_no='KDCR 3-13'", (tid,))
    if cur.fetchone():
        return
    cur.execute(
        """INSERT INTO dwg_drawing (tenant_id, drawing_no, drawing_name, drawing_type,
           dwg_kind, current_rev, status)
           VALUES (%s,'KDCR 3-13','Fan 원심 Casing 제작도','PART','MANUFACTURING','B','APPROVED')
           RETURNING drawing_id""", (tid,))
    drawing_id = cur.fetchone()[0]
    cur.execute(
        "SELECT product_code_id FROM product_code WHERE tenant_id=%s AND main_code='KDCR 3-13'",
        (tid,))
    pc = cur.fetchone()
    # Studio 데모 Macro (APPROVED — 실행 이력 포함)
    cur.execute(
        """INSERT INTO tbx_macro (tenant_id, macro_name, prompt_text, macro_expr,
           description_text, status, hierarchy_address, test_input, test_result)
           VALUES (%s,'Shaft 길이 계산',
                   'SS Fan 샤프트의 길이 계산 — Table 참조 합산',
                   'IF(MC>500, Table12(E,560:800,Cos2)+Var(FES,15), Table12(E,560:800,Cos1)+Var(FES,15))*PreC(1)',
                   'Impeller/Casing/Bearing 폭을 더해 Shaft 길이를 계산',
                   'APPROVED','/M/ENG/FAN/SHAFT',
                   %s, %s)""",
        (tid, json.dumps({"MC": 520, "FES": 15}), json.dumps({"value": 2685})))
    for label, dtype, variant, expr in DIMS_V5:
        macro_id = None
        if expr:
            cur.execute(
                """INSERT INTO tbx_macro (tenant_id, macro_name, macro_expr, status,
                   hierarchy_address)
                   VALUES (%s,%s,%s,'APPROVED',%s) RETURNING macro_id""",
                (tid, f"DIM {label} (KDCR 3-13)", expr, f"/M/ENG/FAN/DIM_{label}"))
            macro_id = cur.fetchone()[0]
        cur.execute(
            """INSERT INTO dwg_dimension (tenant_id, drawing_id, product_code_id,
               dim_label, dim_type, macro_id, variant_value)
               VALUES (%s,%s,%s,%s,%s,%s,%s)""",
            (tid, drawing_id, pc[0] if pc else None, label, dtype, macro_id, variant))
    logger.info("seed v5 complete — dwg_dimension %d건 (Macro 바인딩)", len(DIMS_V5))


# ── seed v6 — UI 번역 (sys_translation, REQ-N-015/SYS-021) ──
# key(≤40자) → (en, ja, zh) — KO 는 프론트 기본 문자열 (폴백)

UI_TRANSLATIONS: dict[str, tuple[str, str, str]] = {
    "common.query": ("Query", "照会", "查询"),
    # 9.32 — 대장 검색 UI (SearchBox·플레이스홀더·빈 결과). en/ja/zh 파리티.
    "common.search": ("Search", "検索", "搜索"),
    "common.clear": ("Clear", "クリア", "清除"),
    "grid.noSearchResults": ("No results — check your search term",
                             "検索結果なし — 検索語を確認してください", "无搜索结果 — 请检查搜索词"),
    "parts.searchPlaceholder": ("Search part no·name·spec", "部品番号・名称・仕様で検索", "按零件号·名称·规格搜索"),
    "companies.searchPlaceholder": ("Search company·notes", "取引先名・備考で検索", "按公司名·备注搜索"),
    "materials.searchPlaceholder": ("Search material code·name", "材料コード・名称で検索", "按材料代码·名称搜索"),
    "drawings.searchPlaceholder": ("Search drawing no·name", "図面番号・名称で検索", "按图号·名称搜索"),
    "products.searchPlaceholder": ("Search product code·name", "製品コード・名称で検索", "按产品代码·名称搜索"),
    "doc.searchPlaceholder": ("Search doc no·title", "文書番号・タイトルで検索", "按文档号·标题搜索"),
    "wo.searchPlaceholder": ("Search WO no·title·drawing·project", "指示番号・件名・図面・案件で検索", "按工单号·标题·图纸·项目搜索"),
    "qc.searchPlaceholder": ("Search inspection no·ref·item", "検査番号・参照・品目で検索", "按检验号·参照·物料搜索"),
    "eco.searchPlaceholder": ("Search ECO no·title·target", "ECO番号・件名・対象で検索", "按ECO号·标题·对象搜索"),
    "po.searchPlaceholder": ("Search PO no·supplier", "発注番号・仕入先で検索", "按订单号·供应商搜索"),
    "costact.searchPlaceholder": ("Search category·item·PO", "分類・品目・発注で検索", "按分类·物料·订单搜索"),
    # 9.34 — 검색화면 커버리지 확장 후 발굴한 미번역 크롬 (검사·원가·제품코드 en/ja/zh 파리티)
    "qc.colNo": ("Insp. No", "検査番号", "检验号"),
    "qc.colType": ("Type", "種別", "类型"),
    "qc.colRef": ("Target", "対象", "对象"),
    "qc.colItem": ("Item", "品目", "物料"),
    "qc.colResult": ("Result", "判定", "判定"),
    "qc.colMeasured": ("Measured", "測定", "测量"),
    "qc.colInspector": ("Inspector", "検査者", "检验员"),
    "qc.colAt": ("Inspected At", "検査日時", "检验时间"),
    "qc.typeIncoming": ("Incoming", "受入検査", "来料检验"),
    "qc.typeProcess": ("In-Process", "工程検査", "过程检验"),
    "qc.typeOutgoing": ("Outgoing", "出荷検査", "出货检验"),
    "qc.resultPass": ("Pass", "合格", "合格"),
    "qc.resultCond": ("Conditional", "条件付", "有条件"),
    "qc.resultFail": ("Fail", "不合格", "不合格"),
    "qc.phItemCode": ("Item code", "品目コード", "物料代码"),
    "qc.phItemName": ("Item name", "品名", "品名"),
    "qc.phRef": ("Ref (PO/WO)", "参照 (PO/WO)", "参照 (PO/WO)"),
    "qc.phMeasured": ("Measured value", "測定値", "测量值"),
    "qc.btnRegister": ("＋ Register Inspection", "＋ 検査登録", "＋ 登记检验"),
    "qc.btnCert": ("🖶 Certificate PDF", "🖶 成績書PDF", "🖶 检验报告PDF"),
    "qc.certTitle": ("Certificate PDF for selected (or all) inspections", "選択検査(または全体)の成績書PDF", "所选(或全部)检验的报告PDF"),
    "cost.catMaterial": ("Material", "材料費", "材料费"),
    "cost.catManufacturing": ("Manufacturing", "製造費", "制造费"),
    "cost.catDirect": ("Direct Expense", "直接経費", "直接费用"),
    "master.composeBtn": ("＋ Compose", "＋ 組合生成", "＋ 组合生成"),
    "master.origin": ("Origin", "生成", "来源"),
    "common.create": ("Create", "生成", "创建"),
    # 9.35 — 잔여 화면 미번역 크롬 (공용 메뉴·ECO대장·마일스톤·재고·재무·감사)
    "menu.erp-tenants": ("Customer Management (M-14-6C)", "顧客管理 (M-14-6C)", "客户管理 (M-14-6C)"),
    "menu.erp-heads": ("Head Management (M-14-6E)", "Head管理 (M-14-6E)", "Head管理 (M-14-6E)"),
    "common.all": ("All", "全て", "全部"),
    "common.title": ("Title", "タイトル", "标题"),
    "common.target": ("Target", "対象", "对象"),
    "common.countSuffix": ("", "件", "件"),
    "ecoL.colNo": ("ECO No", "ECO番号", "ECO编号"),
    "ecoL.colTargetType": ("Target Type", "対象種別", "对象类型"),
    "ecoL.colChangeType": ("Change Type", "変更種別", "变更类型"),
    "ecoL.colRevTransition": ("Rev Transition", "Rev遷移", "Rev迁移"),
    "ecoL.colCreatedBy": ("Created By", "登録者", "登记人"),
    "ecoL.colCreatedAt": ("Created At", "登録日時", "登记时间"),
    "ecoL.colAppliedAt": ("Applied At", "適用日時", "应用时间"),
    "ms.dDONE": ("Done", "完了", "完成"),
    "ms.dOVERDUE": ("Overdue", "遅延", "逾期"),
    "ms.dDUE_SOON": ("Due soon", "間近", "临近"),
    "ms.dPENDING": ("Pending", "待機", "待处理"),
    "common.kinds": ("", "種", "种"),
    "head.namePh": ("Task (user)", "業務 (ユーザー)", "业务 (用户)"),
    "audit.targetPh": ("e.g. cst_price", "例: cst_price", "例: cst_price"),
    "common.new": ("New", "新規", "新建"),
    "common.delete": ("Delete", "削除", "删除"),
    "common.save": ("Save", "保存", "保存"),
    "common.apply": ("Apply", "適用", "应用"),
    "common.approve": ("Approve", "承認", "批准"),
    "common.reject": ("Reject", "差戻", "驳回"),
    "common.upload": ("Upload", "アップロード", "上传"),
    "common.download": ("Download", "ダウンロード", "下载"),
    "common.total": ("Total", "合計", "合计"),
    "shell.file": ("File", "ファイル", "文件"),
    "shell.edit": ("Edit", "編集", "编辑"),
    "shell.view": ("View", "照会", "查看"),
    "shell.tools": ("Tools", "ツール", "工具"),
    "shell.window": ("Window", "ウィンドウ", "窗口"),
    "shell.help": ("Help", "ヘルプ", "帮助"),
    "shell.common": ("Common", "共通", "公共"),
    "shell.todo": ("To-Do", "To-Do", "待办"),
    "shell.alerts": ("Notifications", "通知", "通知"),
    "shell.searchPh": ("Search screen·code·drawing (⌘K)", "画面・コード・図面検索 (⌘K)", "搜索画面·代码·图纸 (⌘K)"),
    "shell.openHint": ("Open a screen from the left menu", "左メニューから画面を開いてください", "请从左侧菜单打开画面"),
    "shell.openHint2": ("Double-click = new tab (MDI)", "ダブルクリック＝新タブ (MDI)", "双击＝新标签页 (MDI)"),
    "shell.pending": ("Pending approvals", "承認待ち", "待审批"),
    "login.title": ("Sign in", "ログイン", "登录"),
    "login.userId": ("User ID", "社員番号", "工号"),
    "login.password": ("Password", "パスワード", "密码"),
    "login.tenant": ("Tenant", "テナント", "租户"),
    "login.submit": ("Sign in (Enter)", "ログイン (Enter)", "登录 (Enter)"),
    "login.checking": ("Checking…", "確認中…", "验证中…"),
    "cpq.airflow": ("Airflow", "風量", "风量"),
    "cpq.pressure": ("Static P.", "静圧", "静压"),
    "cpq.applyF8": ("Apply F8", "適用 F8", "应用 F8"),
    "cpq.finishedCode": ("Finished Goods Code", "完成品コード", "成品代码"),
    "cpq.slotSpec": ("Selected Spec (Slot)", "選択仕様 (Slot)", "选择规格 (Slot)"),
    "cpq.bomTitle": ("BOM · Live Pricing", "BOM・リアルタイム価格", "BOM·实时价格"),
    "cpq.quotePreview": ("Quote Preview", "見積プレビュー", "报价预览"),
    "cpq.specExcel": ("Spec Excel ⬆", "仕様 Excel ⬆", "规格 Excel ⬆"),
    "cpq.qty": ("Qty", "数量", "数量"),
    "cpq.name": ("Name", "品名", "品名"),
    "cpq.amount": ("Amount(K)", "金額(千)", "金额(千)"),
}


def seed_v6(cur, tid: int) -> None:
    cur.execute(
        "SELECT 1 FROM sys_translation WHERE tenant_id=%s AND entity_type='UI' LIMIT 1", (tid,))
    if cur.fetchone():
        return
    n = 0
    for key, (en, ja, zh) in UI_TRANSLATIONS.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                   VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))
            n += 1
    logger.info("seed v6 complete — UI 번역 %d행 (en/ja/zh)", n)


# ── seed v7 — UI 번역 확장: 화면 제목·메뉴 트리·공통 버튼·CAD (REQ-N-015) ──
# key 체계: screen.<screenId> = MDI 탭 제목 / menu.<nodeId> = 좌측 트리 라벨 / common·shell·cad

UI_TRANSLATIONS_V7: dict[str, tuple[str, str, str]] = {
    # 화면 제목 (한글 포함 화면만 — 영문 제목은 폴백 사용)
    "screen.cpq-selection": ("Product Selection — AHU 5", "製品選定 — AHU 5", "产品选定 — AHU 5"),
    "screen.cpq-techdata": ("Technical Data", "技術データ", "技术数据"),
    "screen.code-subcode": ("Sub Code Registration", "Sub Code 登録", "Sub Code 注册"),
    "screen.code-datatable": ("Data Table", "データ Table", "数据 Table"),
    "screen.cpq-docmgmt": ("Document Box", "文書ボックス", "文档箱"),
    "screen.plm-duct": ("Building Duct", "建築設備 Duct", "建筑设备 Duct"),
    "screen.erp-access": ("Users & Permissions", "ユーザー・権限", "用户·权限"),
    "screen.com-approval": ("Approval Inbox", "承認箱", "审批箱"),
    "screen.com-tasks": ("Dept. Task Box", "部署業務箱", "部门任务箱"),
    "screen.com-mobile": ("Mobile Preview", "Mobile プレビュー", "Mobile 预览"),
    "screen.erp-project": ("Project Registration", "Project 登録", "Project 注册"),
    "screen.erp-price": ("Price Management", "単価管理", "单价管理"),
    "screen.erp-purchase": ("Purchase & PO", "購買・発注", "采购·下单"),
    # 좌측 메뉴 트리 (리프 = 화면코드 유지)
    "menu.cpq-selection": ("Product Selection (C-1)", "製品選定 (C-1)", "产品选定 (C-1)"),
    "menu.cpq-techdata": ("Technical Data (C-2)", "技術データ (C-2)", "技术数据 (C-2)"),
    "menu.cpq-docmgmt": ("Document Box (M-5-4)", "文書ボックス (M-5-4)", "文档箱 (M-5-4)"),
    "menu.plm-material": ("Material — planned", "Material — 予定", "Material — 待定"),
    "menu.plm-quality": ("Quality — planned", "Quality — 予定", "Quality — 待定"),
    "menu.plm-arr": ("Arrangement Set-Up — planned", "Arrangement Set-Up — 予定", "Arrangement Set-Up — 待定"),
    "menu.plm-duct": ("Building Duct (M-4-3)", "建築設備 Duct (M-4-3)", "建筑设备 Duct (M-4-3)"),
    "menu.code-subcode": ("Sub Code Registration (S-1-1)", "Sub Code 登録 (S-1-1)", "Sub Code 注册 (S-1-1)"),
    "menu.code-raw": ("Raw Material·GPI — planned", "Raw Material·GPI — 予定", "Raw Material·GPI — 待定"),
    "menu.code-datatable": ("Data Table Mgmt (M-3-7)", "データ Table 管理 (M-3-7)", "数据 Table 管理 (M-3-7)"),
    "menu.code-variant": ("Variant·Constant — planned", "Variant·Constant — 予定", "Variant·Constant — 待定"),
    "menu.erp-purchase": ("PR·PO (M-8-2)", "発注 PR·PO (M-8-2)", "采购 PR·PO (M-8-2)"),
    "menu.erp-price": ("Price Management (M-12-5)", "単価管理 (M-12-5)", "单价管理 (M-12-5)"),
    "menu.erp-access": ("Users & Permissions (M-14-6)", "ユーザー・権限 (M-14-6)", "用户·权限 (M-14-6)"),
    "menu.tbx-templet": ("Templet Mgmt — planned", "Templet 管理 — 予定", "Templet 管理 — 待定"),
    "menu.com-approval": ("Approval Inbox (M-15-2)", "承認箱 (M-15-2)", "审批箱 (M-15-2)"),
    "menu.com-tasks": ("Dept. Task Box (M-15-3)", "部署業務箱 (M-15-3)", "部门任务箱 (M-15-3)"),
    "menu.com-folder": ("Project Folder·History (M-15-8/9)", "Project Folder·履歴 (M-15-8/9)", "Project Folder·历史 (M-15-8/9)"),
    "menu.com-mobile": ("Mobile App Preview (M-16)", "Mobile App プレビュー (M-16)", "Mobile App 预览 (M-16)"),
    "menu.com-search": ("Global Search — planned", "統合検索 — 予定", "全局搜索 — 待定"),
    "menu.moduleCommon": ("Common", "共通", "公共"),
    # 신규 화면 메뉴 라벨 (한글 잔존 해소 — check_i18n_en)
    "menu.cpq-xreview": ("X-code Review (C-1X)", "X-codeレビュー (C-1X)", "X-code 审核 (C-1X)"),
    "menu.plm-eco": ("Design Change ECO/ECN (D-5)", "設計変更 ECO/ECN (D-5)", "设计变更 ECO/ECN (D-5)"),
    "menu.plm-eco-ledger": ("Change History Ledger (D-5L)", "変更履歴台帳 (D-5L)", "变更历史台账 (D-5L)"),
    "menu.plm-bom-compare": ("BOM Compare (G-3B)", "BOM比較 (G-3B)", "BOM 比较 (G-3B)"),
    "menu.code-master": ("Product Code Master (M-3-8)", "製品コードマスタ (M-3-8)", "产品编码主数据 (M-3-8)"),
    "menu.i18n-data": ("Data Translation (M-13-2)", "データ翻訳 (M-13-2)", "数据翻译 (M-13-2)"),
    "menu.erp-sales-order": ("Sales Order (D-1)", "受注管理 (D-1)", "订单管理 (D-1)"),
    "menu.erp-milestone": ("Schedule / Milestone (D-7)", "日程・マイルストーン (D-7)", "进度・里程碑 (D-7)"),
    "menu.erp-calendar": ("Workday / Holiday Calendar (M-8-6)", "稼働日・休日カレンダー (M-8-6)", "工作日・假日日历 (M-8-6)"),
    "menu.erp-finance": ("Multi-currency / Tax Master (M-13-1)", "多通貨・税マスタ (M-13-1)", "多币种・税务主数据 (M-13-1)"),
    "menu.erp-po": ("PO Lifecycle (G-3)", "発注ライフサイクル (G-3)", "采购生命周期 (G-3)"),
    "menu.erp-inventory": ("Inventory (D-2)", "在庫管理 (D-2)", "库存管理 (D-2)"),
    "menu.erp-work-order": ("Work Order (D-3)", "作業指示 (D-3)", "工单 (D-3)"),
    "menu.erp-quality": ("Inspection / Quality (D-4)", "検査・品質 (D-4)", "检验・质量 (D-4)"),
    "menu.erp-cost-actual": ("Cost Actual / Variance (D-6)", "原価実績・差異 (D-6)", "成本实绩・差异 (D-6)"),
    "menu.erp-anomaly": ("Anomaly Event (M-14-4A)", "異常イベント (M-14-4A)", "异常事件 (M-14-4A)"),
    "menu.erp-audit": ("Audit Log (M-14-6A)", "監査照会 (M-14-6A)", "审计查询 (M-14-6A)"),
    "menu.tbx-runs": ("Run History / Cleanup (E-3)", "Run履歴・整理 (E-3)", "Run 历史・清理 (E-3)"),
    # 공통 버튼 확장
    "common.preview": ("Preview", "プレビュー", "预览"),
    "common.close": ("Close", "閉じる", "关闭"),
    "common.edit": ("Edit", "編集", "编辑"),
    "common.print": ("Print", "印刷", "打印"),
    "common.tempSave": ("Temp Save F12", "一時保存 F12", "暂存 F12"),
    "common.requestApproval": ("Request Approval", "承認依頼", "请求审批"),
    "common.runningTest": ("Running Test", "Running Test", "Running Test"),
    "common.backToList": ("Back to List", "一覧へ戻る", "返回列表"),
    "shell.todoApproval": ("Approvals", "承認確認", "审批确认"),
    "shell.todoPl": ("PL delayed", "PL 遅延", "PL 延迟"),
    # CAD
    "cad.hint": ("Wheel zoom · drag pan · dblclick fit · click = props",
                 "ホイール拡大・ドラッグ移動・Wクリック合わせ・クリック＝属性",
                 "滚轮缩放·拖动平移·双击适配·单击=属性"),
    "cad.measureHint": ("Click two points = distance · endpoint/center snap",
                        "2点クリック＝距離計測・端点/中心スナップ",
                        "点击两点=测距·端点/中心捕捉"),
    "cad.measure": ("Measure", "計測", "测量"),
    "cad.layer": ("Layer", "レイヤー", "图层"),
    "cad.fitTitle": ("Fit (double-click)", "フィット (Wクリック)", "适配 (双击)"),
}


def seed_v7(cur, tid: int) -> None:
    # 키 단위 멱등 — v7 사전에 키가 추가되면 누락분만 삽입
    cur.execute(
        """SELECT DISTINCT field FROM sys_translation
           WHERE tenant_id=%s AND entity_type='UI' AND locale='en'""", (tid,))
    have = {r[0] for r in cur.fetchall()}
    n = 0
    for key, (en, ja, zh) in UI_TRANSLATIONS_V7.items():
        if key in have:
            continue
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                   VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))
            n += 1
    if n:
        logger.info("seed v7 — UI 번역 확장 %d행 삽입 (화면·메뉴·CAD)", n)


# ── seed v8 — 도면 대장 (B7): Rev 이력·Supersedure 실데이터 + Run DXF 도면 연결 ──

def seed_v8(cur, tid: int) -> None:
    cur.execute(
        "SELECT drawing_id FROM dwg_drawing WHERE tenant_id=%s AND drawing_no='KDCR 3-13'",
        (tid,))
    row = cur.fetchone()
    if not row:
        return   # v5 도면이 아직 없음 (다음 기동 시 처리)
    did = row[0]
    # Run 산출 DXF 는 KDCR 3-13 제작도(build_part_dxf) — 미연결분 도면에 연결 (매 기동 멱등)
    cur.execute(
        """UPDATE dwg_file SET drawing_id=%s
           WHERE tenant_id=%s AND drawing_id IS NULL AND folder='DWG' AND file_type='DXF'""",
        (did, tid))
    cur.execute("SELECT 1 FROM dwg_revision WHERE drawing_id=%s", (did,))
    if cur.fetchone():
        return
    for rev, rdate, reason, by in [
        ("A", "2026-06-12", "최초 발행", "YS.Gang"),
        ("B", "2026-06-28", "흡입콘 치수 보정 (E 310→320)", "Kim"),
    ]:
        cur.execute(
            """INSERT INTO dwg_revision (drawing_id, rev_no, rev_date, rev_reason, revised_by)
               VALUES (%s,%s,%s,%s,%s)""", (did, rev, rdate, reason, by))
    # 구형 도면 + 대체 이력 (Supersedure 화면 실데이터)
    cur.execute(
        """INSERT INTO dwg_drawing (tenant_id, drawing_no, drawing_name, drawing_type,
           dwg_kind, current_rev, status)
           VALUES (%s,'KDCR 3-12','Fan 원심 Casing 제작도 (구형)','PART','MANUFACTURING',
                   'C','RELEASED')
           RETURNING drawing_id""", (tid,))
    old_id = cur.fetchone()[0]
    cur.execute(
        """INSERT INTO dwg_revision (drawing_id, rev_no, rev_date, rev_reason, revised_by)
           VALUES (%s,'C','2026-05-30','최종 개정 (단종)','YS.Gang')""", (old_id,))
    cur.execute(
        """INSERT INTO dwg_supersedure (tenant_id, old_drawing_id, new_drawing_id,
           reason, superseded_date)
           VALUES (%s,%s,%s,'설계 개정 — 신형 Casing 계열(3-13)로 대체','2026-06-12')""",
        (tid, old_id, did))
    logger.info("seed v8 complete — dwg_revision 3건 · dwg_supersedure 1건 · Run DXF 연결")


# ── seed v9 — UI 번역 전면 확장 (B9): 24화면 크롬 라벨 (REQ-N-015) ──
# 생성 유틸: tools/gen_i18n_bundles.py 가 이 사전들에서 프론트 OFFLINE_BUNDLES 를 생성

UI_TRANSLATIONS_V9: dict[str, tuple[str, str, str]] = {
    "access.accountActions": ("Account Actions — {n}", "アカウント操作 — {n}", "账户操作 — {n}"),
    "access.action": ("Action", "操作", "操作"),
    "access.addUser": ("＋ Add User", "＋ ユーザー登録", "＋ 添加用户"),
    "access.at": ("Timestamp", "日時", "时间"),
    "access.auditLog": ("Recent Audit Log", "最近の監査ログ", "最近审计日志"),
    "access.changeLevel": ("Change Level (Audited)", "レベル変更 (監査)", "变更级别 (审计)"),
    "access.deactivate": ("Deactivate", "無効化", "停用"),
    "access.inviteMail": ("Invite Mail", "招待メール", "邀请邮件"),
    "access.level": ("Level", "レベル", "级别"),
    "access.name": ("Name", "氏名", "姓名"),
    "access.principle1": ("4-tier levels PLATFORM/ADMIN/SETUP/GENERAL",
        "レベル4段階 PLATFORM/ADMIN/SETUP/GENERAL",
        "级别分4级 PLATFORM/ADMIN/SETUP/GENERAL"),
    "access.principle2": ("Unauthorized menus are hidden — front-end hiding is not security (re-checked in service)",
        "権限のないメニューは非表示 — フロント非表示はセキュリティではない(サービス再検査)",
        "无权限菜单不显示 — 前端隐藏不等于安全(服务端复检)"),
    "access.principle3": ("Tenant management (ADM-001) is Platform-only",
        "テナント管理(ADM-001)は Platform 専用",
        "租户管理(ADM-001)仅限 Platform"),
    "access.principles": ("Principles", "原則", "原则"),
    "access.resource": ("Resource", "リソース", "资源"),
    "access.role": ("Role", "ロール", "角色"),
    "access.roleEdit": ("Edit Role — {n} (default: permission approval spec 98-menu matrix)",
        "ロール編集 — {n} (既定: 権限承認定義書 98メニューマトリクス)",
        "角色编辑 — {n} (默认: 权限审批定义书 98菜单矩阵)"),
    "access.search": ("Search", "検索", "搜索"),
    "access.searchPh": ("login·name", "login・氏名", "login·姓名"),
    "access.unlock": ("Unlock", "ロック解除", "解锁"),
    "access.userCount": ("Users — {n}", "ユーザー — {n}名", "用户 — {n}人"),
    "appr.after": ("After", "変更後", "变更后"),
    "appr.assetType": ("Asset Type", "資産種別", "资产类型"),
    "appr.before": ("Before", "変更前", "变更前"),
    "appr.commentPh": ("Comment (required on reject)", "コメント (差戻時必須)", "备注 (驳回时必填)"),
    "appr.date": ("Date", "日付", "日期"),
    "appr.delegatedReqs": ("Delegated Requests", "委任された依頼", "受托请求"),
    "appr.detail": ("Detail", "詳細", "详情"),
    "appr.drawing": ("Drawing", "図面", "图纸"),
    "appr.history": ("History", "履歴", "历史"),
    "appr.inbox": ("Approval Inbox", "承認箱", "审批箱"),
    "appr.input": ("Input", "入力", "输入"),
    "appr.mobileSync": ("Mobile Sync", "モバイル同期", "移动同步"),
    "appr.mobileSyncDesc": ("Same data·rules as mobile approval (M-16) (APP-002)",
        "モバイル承認(M-16)と同一データ・ルール (APP-002)",
        "与移动审批(M-16)相同数据·规则 (APP-002)"),
    "appr.myReqs": ("My Requests", "自分の依頼", "我的请求"),
    "appr.noCircular": ("No Circular References", "循環参照なし", "无循环引用"),
    "appr.refIntegrity": ("Referential Integrity", "参照整合性", "引用完整性"),
    "appr.reqDate": ("Request Date", "依頼日", "请求日期"),
    "appr.reqKind": ("Request Kind", "依頼区分", "请求类别"),
    "appr.requester": ("Requester", "依頼者", "请求人"),
    "appr.result": ("Result", "結果", "结果"),
    "appr.revCompare": ("Rev Compare: Rev.A → Rev.B (dimension B re-review applied)",
        "Rev 比較: Rev.A → Rev.B (寸法B再検討反映)",
        "Rev 比较: Rev.A → Rev.B (尺寸B复审反映)"),
    "appr.rule1": ("On approve, transitions approval_status=APPROVED",
        "承認時 approval_status=APPROVED へ遷移",
        "审批通过时迁移为 approval_status=APPROVED"),
    "appr.rule2": ("System DB-affecting work requires Platform approval",
        "System DB 影響作業は Platform 承認必須",
        "影响 System DB 的操作需 Platform 审批"),
    "appr.rule3": ("AI outputs approvable only after passing verification (Test)",
        "AI 生成物は検証(Test)通過後のみ承認可能",
        "AI 生成物仅在通过验证(Test)后可审批"),
    "appr.rules": ("Approval Rules", "承認ルール", "审批规则"),
    "appr.selectReq": ("Select a request", "依頼を選択してください", "请选择请求"),
    "appr.slotReview": ("Review Slot definition·value list changes", "Slot 定義・値リスト変更レビュー", "Slot 定义·值列表变更审查"),
    "appr.stage": ("Stage", "段階", "阶段"),
    "appr.status": ("Status", "状態", "状态"),
    "appr.tableFormDoc": ("Table · Form · Document", "Table · Form · 文書", "Table · Form · 文档"),
    "appr.target": ("Target", "対象", "对象"),
    "appr.toProcess": ("Requests to Process", "処理する依頼", "待处理请求"),
    "appr.toProcessN": ("Requests to Process — {n}", "処理する依頼 — {n}件", "待处理请求 — {n}件"),
    "appr.type": ("Type", "種別", "类型"),
    "bell.empty": ("No notifications", "通知なし", "无通知"),
    "bell.readAll": ("Mark All Read", "すべて既読", "全部已读"),
    "bell.title": ("Notifications", "通知", "通知"),
    "bell.unreadPoll": ("Unread {n} · 60s polling", "未読 {n} · 60s ポーリング", "未读 {n} · 60s 轮询"),
    "codrel.childGroup": ("[ Child Group ] — double-click = Code detail",
        "[ Child Group ] — ダブルクリック＝コード詳細",
        "[ Child Group ] — 双击=代码详情"),
    "codrel.codeSourceHint": ("Code Source: shows each Slot's item definitions · 3D ☑ 2D ☐",
        "Code Source: 各 Slot の項目定義を表示 · 3D ☑ 2D ☐",
        "Code Source: 显示各 Slot 的项目定义 · 3D ☑ 2D ☐"),
    "codrel.runHint": ("Select Mother Slot combo (B·C·E) and Run — verifies full expansion of matching Children",
        "Mother Slot 組合せ(B・C・E)を選択して Run — 条件一致 Child 全量展開を検証",
        "选择 Mother Slot 组合(B·C·E)并 Run — 验证条件匹配 Child 全量展开"),
    "codrel.slotMap": ("Slot Mapping", "Slot マッピング", "Slot 映射"),
    "codrel.slotMapHint": ("Slot mapping (slot_map) propagates Mother values into Child code digits (CODE-008)",
        "Slot マッピング(slot_map)で Mother 値が Child コード桁へ伝播 (CODE-008)",
        "通过 Slot 映射(slot_map)将 Mother 值传播到 Child 代码位 (CODE-008)"),
    "codrel.testNeeded": ("Test Required", "Test 必要", "需 Test"),
    "codrel.testPassed": ("Running Test Passed", "Running Test 合格", "Running Test 通过"),
    "cvs.cmd": ("Command:", "コマンド:", "命令:"),
    "dash.alertsTitle": ("Anomaly Alerts (Time·Cash) — double-click = event detail",
        "異常警告 (時間・資金) — ダブルクリック＝イベント詳細",
        "异常警报 (时间·资金) — 双击=事件详情"),
    "dash.blkCost": ("Cost 18.8", "原価 18.8", "成本 18.8"),
    "dash.blkIn": ("Inflow +9.1", "入金 +9.1", "收款 +9.1"),
    "dash.blkOrder": ("Orders 23.0", "受注 23.0", "接单 23.0"),
    "dash.blkOut": ("Outflow -4.2", "出金 -4.2", "付款 -4.2"),
    "dash.cashTitle": ("Cash Flow (Purchase Out · Sales In)", "資金フロー (仕入出金 · 売上入金)", "资金流 (采购付款 · 销售收款)"),
    "dash.content": ("Content", "内容", "内容"),
    "dash.dept": ("Dept.", "部署", "部门"),
    "dash.deptEvents": ("Events by Dept.", "部署別 Event 状況", "各部门 Event 状况"),
    "dash.doneWeek": ("Done (wk)", "完了(週)", "完成(周)"),
    "dash.flowAp": ("AP Approval Docs", "AP 承認図書", "AP 审批文档"),
    "dash.flowApp": ("APP Customer Approval", "APP 顧客承認", "APP 客户审批"),
    "dash.flowDf": ("DF Delivery", "DF 納品", "DF 交付"),
    "dash.flowFf": ("FF Completion", "FF 完成", "FF 完工"),
    "dash.flowHint": ("Currently at OR order stage — automatic (☑) processes transition via EDIM Run outputs",
        "現在 OR 受注段階 — 自動(☑)プロセスは EDIM Run 成果物で状態遷移",
        "当前处于 OR 接单阶段 — 自动(☑)流程由 EDIM Run 产出物驱动状态迁移"),
    "dash.flowIr": ("IR Progress Billing", "IR 出来高請求", "IR 进度请款"),
    "dash.flowMi": ("MI Receiving", "MI 入庫", "MI 入库"),
    "dash.flowMp": ("MP Production Plan", "MP 生産計画", "MP 生产计划"),
    "dash.flowMr": ("MR Mfg. Request", "MR 製作依頼", "MR 制作委托"),
    "dash.flowOr": ("OR Order", "OR 受注", "OR 接单"),
    "dash.flowPcr": ("PCR Quote Review", "PCR 見積検討", "PCR 报价审查"),
    "dash.flowPo": ("PO Ordering", "PO 発注", "PO 下单"),
    "dash.flowPr": ("PR Order Request", "PR 発注依頼", "PR 下单申请"),
    "dash.flowPs": ("PS Registration", "PS 登録", "PS 注册"),
    "dash.flowQcr": ("QCR Quote", "QCR 見積", "QCR 报价"),
    "dash.flowWr": ("WR Work Order", "WR 作業指示", "WR 工单"),
    "dash.kind": ("Kind", "区分", "类别"),
    "dash.kindMoney": ("Cash", "資金", "资金"),
    "dash.kindTime": ("Time", "時間", "时间"),
    "dash.kpiAlerts": ("Anomaly Alerts (Time·Cash)", "異常警告 (時間・資金)", "异常警报 (时间·资金)"),
    "dash.kpiApprovalWait": ("Pending Approval", "承認待ち", "待审批"),
    "dash.kpiMonthOrder": ("Orders This Month", "今月の受注", "本月接单"),
    "dash.kpiRunning": ("Active Project", "進行中 Project", "进行中 Project"),
    "dash.monAug": ("Aug", "08月", "08月"),
    "dash.monSep": ("Sep", "09月", "09月"),
    "dash.processStatus": ("Project Up/Downstream Process Status — {n}",
        "Project 前後工程状況 — {n}",
        "Project 前后工序状况 — {n}"),
    "dash.profitTitle": ("Profit/Loss by Project (Order·Cost·Margin)",
        "Project 別損益 (受注額・原価・マージン)",
        "各 Project 损益 (接单额·成本·毛利)"),
    "dash.sourceHint": ("Source: erp_process_event state machine (ERP-014) · roles ADMIN·Mgmt",
        "集計元: erp_process_event 状態機械 (ERP-014) · 権限 ADMIN・経営",
        "汇总来源: erp_process_event 状态机 (ERP-014) · 权限 ADMIN·经营"),
    "docmgmt.addDoc": ("＋ Add Document", "＋ 文書登録", "＋ 添加文档"),
    "docmgmt.approveWaiting": ("Approve Pending", "Approve 待ち", "Approve 待定"),
    "docmgmt.cancel": ("Cancel", "キャンセル", "取消"),
    "docmgmt.category": ("Category", "分類", "分类"),
    "docmgmt.docLedger": ("Document Ledger — {n} (double-click = doc detail)",
        "文書台帳 — {n}件 (ダブルクリック＝文書詳細)",
        "文档台账 — {n}件 (双击=文档详情)"),
    "docmgmt.docType": ("Document Type", "文書種別", "文档类型"),
    "docmgmt.gradeHint": ("S-1/S-2 force CONFIDENTIAL watermark · masked when under-privileged (DOC-002)",
        "S-1/S-2 は CONFIDENTIAL 透かし強制 · 権限不足時マスキング (DOC-002)",
        "S-1/S-2 强制 CONFIDENTIAL 水印 · 权限不足时脱敏 (DOC-002)"),
    "docmgmt.previewGrade": ("Preview — Grade Control", "プレビュー — Grade 統制", "预览 — Grade 管控"),
    "docmgmt.printHint": ("Watermark·print controls applied (real render)",
        "透かし・出力統制適用 (実レンダー)",
        "水印·打印管控生效 (实渲染)"),
    "docmgmt.regTitle": ("Add Document — doc_control", "文書登録 — doc_control", "文档登记 — doc_control"),
    "docmgmt.registerF12": ("Register F12", "登録 F12", "登记 F12"),
    "docmgmt.rendering": ("Rendering… (backend required)", "レンダリング中… (バックエンド必要)", "渲染中… (需要后端)"),
    "docmgmt.search": ("Search", "検索", "搜索"),
    "docmgmt.selDocProgress": ("Selected Document — {n} · Progress",
        "選択文書 — {n} · Progress",
        "所选文档 — {n} · Progress"),
    "docmgmt.selectDoc": ("Select a document", "文書を選択", "选择文档"),
    "docmgmt.statusFilter": ("Status Filter", "状態フィルター", "状态筛选"),
    "docmgmt.titleReq": ("Title *", "タイトル *", "标题 *"),
    "docmgmt.type": ("Type", "種別", "类型"),
    "doctpl.calcRefTable": ("Calc Reference Table", "計算参照 Table", "计算引用 Table"),
    "doctpl.chartTable": ("[Chart] · [Table]", "[チャート] · [Table]", "[图表] · [Table]"),
    "doctpl.correction": ("Correction", "補正", "校正"),
    "doctpl.densityCalcDoc": ("Density Calc Sheet", "密度計算書", "密度计算书"),
    "doctpl.fanPerfReport": ("Fan Performance Report", "Fan 性能 Report", "Fan 性能 Report"),
    "doctpl.formDef": ("Form definition: Document Set-up S-3-3 (CPQ-012)",
        "Form 定義: Document Set-up S-3-3 (CPQ-012)",
        "Form 定义: Document Set-up S-3-3 (CPQ-012)"),
    "doctpl.inputOutput": ("Input Data → Output Data (Macro calc)",
        "Input Data → Output Data (Macro 計算)",
        "Input Data → Output Data (Macro 计算)"),
    "doctpl.macroCalcF9": ("▶ Macro Calc F9", "▶ Macro 計算 F9", "▶ Macro 计算 F9"),
    "doctpl.printPreview": ("Print Preview", "Print プレビュー", "Print 预览"),
    "doctpl.psychroChart": ("Psychrometric Chart — calc points plotted (graph Templet)",
        "湿り空気線図 — 計算点表示 (グラフ Templet)",
        "焓湿图 — 显示计算点 (图形 Templet)"),
    "doctpl.spreadsheet": ("Spreadsheet (A~L — intermediate calc)",
        "スプレッドシート (A~L — 中間計算)",
        "电子表格 (A~L — 中间计算)"),
    "dtable.addRowF2": ("＋ Add Row F2", "＋ 行追加 F2", "＋ 添加行 F2"),
    "dtable.addressTitle": ("Table Address — Hierarchy = DB address",
        "Table Address — Hierarchy = DB アドレス",
        "Table Address — Hierarchy = DB 地址"),
    "dtable.apprStatus": ("Approval Status", "承認状態", "审批状态"),
    "dtable.gridTitle": ("Table12 — double-click cell = edit (Excel syntax)",
        "Table12 — セルダブルクリック ＝ 編集 (Excel 文法)",
        "Table12 — 双击单元格 = 编辑 (Excel 语法)"),
    "dtable.importRule1": ("Only structured-format Excel (row-1 header = column names) allowed",
        "定型様式(1行目ヘッダー = 列名)の Excel のみ許可",
        "仅允许规范格式(首行表头 = 列名)的 Excel"),
    "dtable.importRule2": ("Duplicate Key rows rejected — download error report",
        "Key 重複行は拒否 — エラーレポートをダウンロード",
        "Key 重复行拒绝 — 下载错误报告"),
    "dtable.importRule3": ("Text in numeric column ignores that cell (warning)",
        "数値列にテキスト混入時は該当セル無視(警告)",
        "数值列混入文本时忽略该单元格(警告)"),
    "dtable.importRules": ("Import Rules", "Import ルール", "Import 规则"),
    "dtable.keyHint": ("Numeric Key auto-parses row_key_num — range query (10:25) sort guaranteed",
        "Key が数値なら row_key_num 自動パース — 範囲照会(10:25)のソート保証",
        "Key 为数字时自动解析 row_key_num — 范围查询(10:25)保证排序"),
    "dtable.macroRefHint": ("Macro references via this path:", "Macro がこのパスで参照:", "Macro 通过此路径引用:"),
    "dtable.reapprovalHint": ("Re-approval required on save (Grade policy)",
        "変更保存時に再承認対象 (Grade ポリシー)",
        "变更保存时需重新审批 (Grade 政策)"),
    "dtable.refMacroTitle": ("Referencing Macro — 4 impacted", "参照 Macro — 影響度 4件", "引用 Macro — 影响 4件"),
    "dtable.screen": ("Screen", "画面", "画面"),
    "dtable.tableDef": ("Table Definition", "Table 定義", "Table 定义"),
    "dtable.unsaved": ("{n} changes unsaved", "変更 {n}件 未保存", "变更 {n}件 未保存"),
    "dtable.usage": ("Usage Formula", "使用式", "使用公式"),
    "duct.accessTurning": ("Access Door·Turning", "点検口・Turning", "检修口·Turning"),
    "duct.ach": ("Air Changes", "換気回数", "换气次数"),
    "duct.addDiffuser": ("Add Diffuser", "Diffuser 追加", "添加 Diffuser"),
    "duct.aiReadHint": ("AI reads architectural drawing: ramps·beams·fire zones·rooms → marks no-install areas (DUCT-001)",
        "建築図 AI 判読: ランプ・梁・防火区画・室区分 → 設置不可エリア表示 (DUCT-001)",
        "建筑图 AI 判读: 坡道·梁·消防分区·房间区分 → 标记不可安装区域 (DUCT-001)"),
    "duct.airCond": ("Air Conditions", "空気条件", "空气条件"),
    "duct.airflowBasis": ("Airflow Basis", "風量基準", "风量基准"),
    "duct.autoPlace": ("▶ Auto Place (shortest path·fluid flow)",
        "▶ 自動配置 (最短経路・流体フロー)",
        "▶ 自动布置 (最短路径·流体流动)"),
    "duct.basic": ("Basic", "基本", "基本"),
    "duct.calcHint": ("Calculates pressure loss·Leak·load after placement",
        "配置後に圧力損失・Leak・荷重を計算",
        "布置后计算压力损失·Leak·载荷"),
    "duct.ceiling": ("Ceiling Ht/Beam/Tile", "階高/梁/テックス", "层高/梁/吊顶板"),
    "duct.designCond": ("Design Conditions", "設計条件", "设计条件"),
    "duct.diffuserPlaced": ("{n} Diffusers auto-placed (qty adjustable)",
        "Diffuser {n}個を自動配置 (数量調整可)",
        "自动布置 Diffuser {n}个 (可调数量)"),
    "duct.drawingCall": ("Load Drawing (AI Read)", "図面呼出 (AI 判読)", "调用图纸 (AI 判读)"),
    "duct.edimHint1": ("Eng: BOM/Code · Mfg: min scrap", "技術: BOM/Code・製造: 最小スクラップ", "技术: BOM/Code·制造: 最少废料"),
    "duct.edimHint2": ("Material: purchasing · Sales: quote", "資材: 購買・営業: 見積", "物料: 采购·销售: 报价"),
    "duct.edimLink": ("EDIM Link", "EDIM 連携", "EDIM 关联"),
    "duct.equipment": ("Equipment", "機器", "设备"),
    "duct.exhaust": ("Exhaust", "排気", "排风"),
    "duct.floor": ("Floor", "階", "楼层"),
    "duct.heatLoad": ("Heat Load", "発熱", "发热"),
    "duct.highVel": ("High Vel.", "高速", "高速"),
    "duct.itemCol": ("Item", "項目", "项目"),
    "duct.lowVel": ("Low Vel.", "低速", "低速"),
    "duct.manualAdjust": ("Manual Adjust (Drag·Click)", "手動調整 (Drag・Click)", "手动调整 (Drag·Click)"),
    "duct.materialDb": ("Duct Material DB", "Duct 資材 DB", "Duct 材料 DB"),
    "duct.materialHint": ("Type·loss·weight·max length per Size",
        "種類・損失・重量・Size別最大長",
        "种类·损失·重量·按 Size 最大长度"),
    "duct.multiFloor": ("Multi Floor (Point XYZ)", "複数階 (Point XYZ)", "多楼层 (Point XYZ)"),
    "duct.occupancy": ("Occupancy", "人員", "人员"),
    "duct.room": ("Room", "室", "房间"),
    "duct.runAutoHint": ("▶ Run auto place", "▶ 自動配置を実行してください", "▶ 请执行自动布置"),
    "duct.scopeChip": ("Scope confirmation target (Suppl. Note §3.3)",
        "事業範囲確定対象 (補完ノート §3.3)",
        "业务范围待确定 (补充说明 §3.3)"),
    "duct.startEnd": ("Start–End", "始点–終点", "起点–终点"),
    "duct.stdSelect": ("Std Select", "Std 選択", "Std 选择"),
    "duct.supply": ("Supply Air", "給気", "送风"),
    "duct.techCalc": ("Technical Data Calc", "技術データ計算", "技术资料计算"),
    "duct.usage": ("Usage", "用途", "用途"),
    "duct.valueCol": ("Value", "値", "值"),
    "duct.velocityBasis": ("Velocity Basis", "流速基準", "流速基准"),
    "duct.vent": ("Ventilation", "換気", "通风"),
    "dwg.byCol": ("Handler", "処理者", "处理人"),
    "dwg.cancel": ("Cancel", "キャンセル", "取消"),
    "dwg.countChip": ("dwg_drawing {n} rows", "dwg_drawing {n}件", "dwg_drawing {n}件"),
    "dwg.dateCol": ("Date", "日付", "日期"),
    "dwg.dblOpenHint": ("Double-click = linked DXF CAD viewer",
        "ダブルクリック＝連結 DXF CAD ビューア",
        "双击＝关联 DXF CAD 查看器"),
    "dwg.drawingName": ("Drawing Name", "図面名", "图纸名称"),
    "dwg.drawingNo": ("Drawing No.", "図面番号", "图纸编号"),
    "dwg.fileCol": ("File (DXF)", "ファイル (DXF)", "文件 (DXF)"),
    "dwg.ledger": ("Drawing Ledger", "図面台帳", "图纸台账"),
    "dwg.listTitle": ("Drawing List — {n} (click = Rev history)",
        "図面一覧 — {n}件 (クリック＝Rev 履歴)",
        "图纸列表 — {n}件 (单击＝Rev 历史)"),
    "dwg.newNo": ("New Drawing", "新図面", "新图纸"),
    "dwg.nextRev": ("Next", "次", "下一"),
    "dwg.noPlaceholder": ("e.g. KDCR 3-15", "例: KDCR 3-15", "例: KDCR 3-15"),
    "dwg.noSup": ("No supersedure history", "代替履歴なし", "无替代历史"),
    "dwg.oldNo": ("Old Drawing", "旧図面", "旧图纸"),
    "dwg.queryF8": ("Query F8", "照会 F8", "查询 F8"),
    "dwg.reasonCol": ("Reason", "事由", "事由"),
    "dwg.regTitle": ("Drawing Registration — dwg_drawing", "図面登録 — dwg_drawing", "图纸注册 — dwg_drawing"),
    "dwg.registerF12": ("Register F12", "登録 F12", "注册 F12"),
    "dwg.registerF2": ("＋ Register Drawing F2", "＋ 図面登録 F2", "＋ 图纸注册 F2"),
    "dwg.revCount": ("Rev Count", "Rev 数", "Rev 数"),
    "dwg.revHistory": ("Rev History", "Rev 履歴", "Rev 历史"),
    "dwg.revReasonPh": ("Revision reason", "改訂事由", "修订事由"),
    "dwg.revUp": ("Rev Up", "Rev 上げ", "Rev 升版"),
    "dwg.selectHint": ("Select a drawing row on the left to view Rev history",
        "左の図面行を選択すると Rev 履歴が表示されます",
        "选择左侧图纸行即可显示 Rev 历史"),
    "dwg.selectOpt": ("(select)", "(選択)", "(选择)"),
    "dwg.statusCol": ("Status", "状態", "状态"),
    "dwg.supHint": ("Old drawing → new drawing supersede", "旧図面 → 新図面 代替", "旧图纸 → 新图纸 替代"),
    "dwg.supRegister": ("Register Supersedure", "代替登録", "替代注册"),
    "dwg.supTitle": ("Supersedure — {n} (dwg_supersedure)",
        "Supersedure — {n}件 (dwg_supersedure)",
        "Supersedure — {n}件 (dwg_supersedure)"),
    "dwg.superseded": ("Superseded", "代替済", "已替代"),
    "dwg.typeCol": ("Type", "種別", "类型"),
    "editor.asmSeq": ("Assembly Sequence", "組立順序", "装配顺序"),
    "editor.blockDblHint": ("Block double-click = part info detail",
        "Block ダブルクリック＝部品情報詳細",
        "Block 双击＝零件信息详情"),
    "editor.cadNeedsBackend": ("CAD server unavailable — please refresh (temporary connection error)",
        "CAD サーバー接続不可 — 更新してください (一時的な接続エラー)",
        "CAD 服务器不可用 — 请刷新 (临时连接错误)"),
    "editor.dblEdit": ("Double-click = edit", "ダブルクリック＝編集", "双击＝编辑"),
    "editor.designRule": ("Design Rule — Dimension Set-up",
        "Design Rule — 寸法 Set-up",
        "Design Rule — 尺寸 Set-up"),
    "editor.drawing": ("Drafting…", "作図中…", "绘图中…"),
    "editor.evaluated": ("Evaluated ✓", "評価 ✓", "评估 ✓"),
    "editor.exportDxf": ("Export DXF", "DXF エクスポート", "导出 DXF"),
    "editor.kindCol": ("Kind", "区分", "区分"),
    "editor.macroHint": ("EDIM Macro call → shows formula·direct input allowed",
        "EDIM Macro 呼出 → 計算式表示・直接入力可",
        "调用 EDIM Macro → 显示计算式·可直接输入"),
    "editor.openDxf": ("Open DXF", "DXF を開く", "打开 DXF"),
    "editor.relCond": ("Cond1: vertical·horizontal·center·middle / Cond2: contact(face·line·point)·coords·angle",
        "条件1: 垂直・水平・中心・中央 / 条件2: 接触(面・線・点)・座標・角度",
        "条件1: 垂直·水平·中心·中央 / 条件2: 接触(面·线·点)·坐标·角度"),
    "editor.relPriority": ("Priority: A/B ① → B/C ② (auto cycle check)",
        "優先順位: A/B ① → B/C ② (循環自動チェック)",
        "优先级: A/B ① → B/C ② (循环自动检查)"),
    "editor.relValue": ("Relation Value", "関係値", "关系值"),
    "editor.snapOn": ("Snap: endpoint·middle·center", "Snap: 端点・中央・中心", "Snap: 端点·中点·中心"),
    "editor.subItemDwg": ("Sub Item DWG · Assembly Sequence", "Sub Item DWG・組立順序", "Sub Item DWG·装配顺序"),
    "editor.toolCopy": ("Copy CO", "コピー CO", "复制 CO"),
    "editor.toolDim": ("Dimension DI", "寸法 DI", "尺寸 DI"),
    "editor.toolErase": ("Erase E", "削除 E", "删除 E"),
    "editor.toolExtend": ("Extend", "延長", "延伸"),
    "editor.toolMirror": ("Mirror", "反転", "镜像"),
    "editor.toolMove": ("Move", "移動", "移动"),
    "editor.toolProps": ("Properties CH", "属性 CH", "属性 CH"),
    "editor.toolRotate": ("Rotate RO", "回転 RO", "旋转 RO"),
    "editor.toolTrim": ("Trim TR", "トリム TR", "修剪 TR"),
    "editor.verifyMacro": ("Verify Macro", "検証 Macro", "验证 Macro"),
    "enum.active": ("Active", "適用中", "应用中"),
    "enum.all": ("All", "全体", "全部"),
    "enum.delayed": ("Delayed", "遅延", "延迟"),
    "enum.done": ("Done", "完了", "完成"),
    "enum.expired": ("Expired", "満了", "过期"),
    "enum.progress": ("In Progress", "進行", "进行"),
    "enum.purchase": ("Purchase", "購買", "采购"),
    "enum.quote": ("Quote", "見積", "报价"),
    "enum.quoteApplied": ("Quote Applied", "見積適用", "报价应用"),
    "enum.stock": ("Stock", "在庫", "库存"),
    "enum.waiting": ("Waiting", "待機", "待定"),
    "folder.action": ("Action", "作業", "操作"),
    "folder.actor": ("Actor", "作業者", "操作人"),
    "folder.at": ("Timestamp", "日時", "时间"),
    "folder.autoSaveHint": ("EDIM Run output auto-save convention (Overview §6) — dwg_file.folder CHECK 5 types",
        "EDIM Run 成果物自動保存規約 (概要 §6) — dwg_file.folder CHECK 5種",
        "EDIM Run 产出物自动保存规约 (概要 §6) — dwg_file.folder CHECK 5种"),
    "folder.batchOps": ("Batch Ops", "一括作業", "批量操作"),
    "folder.countDxfHint": ("{n} items (DXF double-click = CAD viewer)",
        "{n}件 (DXF ダブルクリック＝CAD ビューア)",
        "{n}件 (DXF 双击＝CAD 查看器)"),
    "folder.customerExport": ("Export for Customer", "顧客提出用エクスポート", "客户交付导出"),
    "folder.fileName": ("File Name", "ファイル名", "文件名"),
    "folder.historyTitle": ("History Query (sys_history) — diff = before/after JSON compare",
        "履歴照会 (sys_history) — diff = before/after JSON 比較",
        "历史查询 (sys_history) — diff = before/after JSON 比较"),
    "folder.linkedDrawing": ("Linked Drawing", "連結図面", "关联图纸"),
    "folder.receivedDesc": ("Received materials (S-3-5) + data migration originals preserved (Migration Plan §2-4)",
        "受付資料(S-3-5) + データ移行原本保存 (移行計画書 §2-4)",
        "接收资料(S-3-5) + 数据迁移原件保存 (迁移计划书 §2-4)"),
    "folder.receivedTitle": ("RECEIVED Folder", "RECEIVED フォルダ", "RECEIVED 文件夹"),
    "folder.selectFile": ("Select a file", "ファイルを選択してください", "请选择文件"),
    "folder.selectedFile": ("Selected File", "選択ファイル", "所选文件"),
    "folder.targetAll": ("Target: All", "対象: 全体", "对象: 全部"),
    "folder.zipDownload": ("ZIP Download", "ZIP ダウンロード", "ZIP 下载"),
    "mobile.approvalCaption": ("Approval Detail — QR entry", "承認詳細 — QR 進入", "审批详情 — QR 进入"),
    "mobile.arNote": ("AR view deferred until Digital Twin integration (APP-008)",
        "AR ビューは Digital Twin 連携後の対応 (APP-008)",
        "AR 视图待 Digital Twin 集成后处理 (APP-008)"),
    "mobile.arView": ("AR View", "AR ビュー", "AR 视图"),
    "mobile.attachPhoto": ("Attach Photo", "写真添付", "附加照片"),
    "mobile.designNote": ("Design Notes", "設計ノート", "设计笔记"),
    "mobile.designNoteBody": ("① QR entry (APP-001) ② Approvals share web inbox data (APP-002) ③ Project-centric chat History (APP-003) ④ Material·inspection with photo attach + offline cache (APP-004/005) — 32px touch-target exception noted",
        "① QR 進入 (APP-001) ② 承認はウェブ承認箱と同一データ (APP-002) ③ Project 中心の会話 History (APP-003) ④ 資材・検収は写真添付 + オフラインキャッシュ (APP-004/005) — タッチターゲット 32px 例外明記",
        "① QR 进入 (APP-001) ② 审批与网页审批箱同一数据 (APP-002) ③ 以 Project 为中心的对话 History (APP-003) ④ 物料·验收支持照片附件 + 离线缓存 (APP-004/005) — 明确 32px 触控目标例外"),
    "mobile.homeCaption": ("Home — Menu Grid", "ホーム — メニュー Grid", "首页 — 菜单 Grid"),
    "mobile.inspection": ("Inspection", "検収", "验收"),
    "mobile.maintenance": ("Maintenance", "保守", "维护"),
    "mobile.materialInOut": ("Material In/Out", "資材入出庫", "物料出入库"),
    "mobile.materialQrScan": ("Material QR Scan", "資材 QR スキャン", "物料 QR 扫描"),
    "mobile.notice": ("Notices", "お知らせ", "公告"),
    "mobile.projectChat": ("Project Chat", "Project コミュニケーション", "Project 沟通"),
    "mobile.projectChatHist": ("Project Chat (History)", "Project 会話 (History)", "Project 对话 (History)"),
    "mobile.projectInfo": ("Project Info", "Project 情報", "Project 信息"),
    "mobile.qrHint": ("QR is the single entry point — straight to drawings·documents·Project·tasks (for areas without PC access)",
        "QR がすべての入口 — 図面・書類・Project・業務へ直行 (PC 利用不可エリア対応)",
        "QR 是所有入口 — 直达图纸·文件·Project·业务 (应对无法使用电脑的区域)"),
    "mobile.qrScan": ("QR Scan", "QR スキャン", "QR 扫描"),
    "mobile.qrScanDesc": ("Access drawings·documents·Project·tasks",
        "図面・書類・Project・業務へアクセス",
        "访问图纸·文件·Project·业务"),
    "mobile.qty": ("Qty", "数量", "数量"),
    "mobile.receive": ("Receiving", "入庫処理", "入库处理"),
    "mobile.receiveCaption": ("Material Receiving — On-site", "資材入庫 — 現場処理", "物料入库 — 现场处理"),
    "mobile.request": ("Request", "依頼", "请求"),
    "mobile.taskApproval": ("Task Approval", "業務承認", "业务审批"),
    "mobile.warehouse": ("Warehouse", "倉庫", "仓库"),
    "price.addPrice": ("＋ Add Unit Price", "＋ 単価登録", "＋ 单价登记"),
    "price.applyDate": ("Apply Date", "適用日", "适用日"),
    "price.avg": ("Avg", "平均", "平均"),
    "price.baseDate": ("Base Date", "基準日", "基准日"),
    "price.cancel": ("Cancel", "キャンセル", "取消"),
    "price.currency": ("Currency", "通貨", "货币"),
    "price.last": ("Latest", "直近", "最近"),
    "price.ledger": ("Price Ledger — {n} rows (dblclick = code detail)",
        "単価台帳 — {n}件 (Wクリック＝コード詳細)",
        "单价台账 — {n}条 (双击=代码详情)"),
    "price.max": ("Max", "最高", "最高"),
    "price.min": ("Min", "最低", "最低"),
    "price.noExpiry": ("(no expiry)", "(無期限)", "(无期限)"),
    "price.noPrice": ("No unit price", "単価なし", "无单价"),
    "price.priceLbl": ("Unit Price", "単価", "单价"),
    "price.priceTable": ("Price Table", "単価 Table", "单价 Table"),
    "price.prio1": ("① Quote Applied", "① 見積適用", "① 报价适用"),
    "price.prio2": ("② Purchase History", "② 購買履歴", "② 采购记录"),
    "price.prio3": ("③ Stock Price", "③ 在庫単価", "③ 库存单价"),
    "price.prio4": ("④ Quote", "④ 見積", "④ 报价"),
    "price.priority": ("Apply Priority (Pricing Run resolve)",
        "適用優先順位 (Pricing Run resolve)",
        "适用优先级 (Pricing Run resolve)"),
    "price.priorityHint": ("Code·period match first — EXCLUDE constraint blocks period overlap (DB v0.5)",
        "Code・期間一致を優先 — EXCLUDE 制約で期間重複を遮断 (DB v0.5)",
        "Code·期间匹配优先 — EXCLUDE 约束阻止期间重叠 (DB v0.5)"),
    "price.queryF8": ("Query F8", "照会 F8", "查询 F8"),
    "price.regTitle": ("Add Unit Price — cst_price", "単価登録 — cst_price", "单价登记 — cst_price"),
    "price.registerF12": ("Register F12", "登録 F12", "登记 F12"),
    "price.resolveSim": ("Resolve Simulation", "Resolve シミュレーション", "Resolve 模拟"),
    "price.stockCalc": ("Stock Price Calc — {n} (auto from in/out, ERP-021)",
        "在庫単価算出 — {n} (入出庫ベース自動, ERP-021)",
        "库存单价计算 — {n} (基于出入库自动, ERP-021)"),
    "price.supplier": ("Supplier", "仕入先", "供应商"),
    "price.tableAll": ("All (4 types)", "全体 (4種)", "全部 (4种)"),
    "price.tableApplied": ("4. Quote Applied", "4. 見積適用", "4. 报价适用"),
    "price.tablePurchase": ("2. Purchase History", "2. 購買履歴", "2. 采购记录"),
    "price.tableQuote": ("1. Quote", "1. 見積", "1. 报价"),
    "price.tableStock": ("3. Stock Price", "3. 在庫単価", "3. 库存单价"),
    "price.validFrom": ("Valid From", "適用開始", "适用开始"),
    "price.validTo": ("Valid To", "適用終了", "适用结束"),
    "printsetup.a4Landscape": ("A4 Landscape", "A4 横", "A4 横向"),
    "printsetup.a4Portrait": ("A4 Portrait", "A4 縦", "A4 纵向"),
    "printsetup.approvalPending": ("Pending Approval", "承認待ち", "待审批"),
    "printsetup.color": ("Color", "色", "颜色"),
    "printsetup.colorOpt": ("Color", "カラー", "彩色"),
    "printsetup.dataBind": ("Data Placement", "Data 位置指定", "Data 位置指定"),
    "printsetup.dataCall": ("Data Call", "Data 呼出", "Data 调用"),
    "printsetup.dataCustomer": ("[Data] Customer·Date·Owner",
        "[Data] Customer·Date·担当",
        "[Data] Customer·Date·负责人"),
    "printsetup.defaultLayout": ("Default Form Layout", "基本様式配置", "默认样式布局"),
    "printsetup.export": ("Export", "エクスポート", "导出"),
    "printsetup.flowHint": ("Place placeholders → bind data paths → published Form rendered by SVC-11 (shared by quote·PCR·work order)",
        "プレースホルダー配置 → データパスバインド → 公開 Form は SVC-11 がレンダリング (見積書・PCR・作業指示書共通)",
        "放置占位符 → 绑定数据路径 → 已发布 Form 由 SVC-11 渲染 (报价单·PCR·工单通用)"),
    "printsetup.fontSize": ("Font / Size", "Font / サイズ", "Font / 大小"),
    "printsetup.footerPh": ("Footer — page · company info", "フッター — ページ・会社情報", "页脚 — 页码·公司信息"),
    "printsetup.gradeHint": ("Grade-controlled documents enforce watermark·print limits (DOC-002)",
        "Grade 統制文書は透かし・出力制限を強制 (DOC-002)",
        "Grade 管控文档强制水印·打印限制 (DOC-002)"),
    "printsetup.graphPerf": ("[Graph] Performance Curve", "[グラフ] 性能曲線", "[图表] 性能曲线"),
    "printsetup.headerFooter": ("Header·Footer", "ヘッダー・フッター", "页眉·页脚"),
    "printsetup.headerPh": ("Header — logo · Title · DOC No.",
        "ヘッダー — ロゴ・Title・DOC No.",
        "页眉 — 标志·Title·DOC No."),
    "printsetup.inspectReport": ("Inspection Report", "検査成績書", "检验报告"),
    "printsetup.loadGraph": ("Load Graph", "グラフ読込", "加载图表"),
    "printsetup.margin": ("Margin", "余白", "页边距"),
    "printsetup.monoOpt": ("B&W", "モノクロ", "黑白"),
    "printsetup.none": ("None", "なし", "无"),
    "printsetup.outputSettings": ("Output Settings", "出力設定", "输出设置"),
    "printsetup.paper": ("Paper", "用紙", "纸张"),
    "printsetup.quoteForm": ("Quote (CLT)", "見積書 (CLT)", "报价单 (CLT)"),
    "printsetup.requestPublish": ("Request Approval → Publish", "承認依頼 → 公開", "请求审批 → 发布"),
    "printsetup.stdTemplet": ("Standard Templet", "標準 Templet", "标准 Templet"),
    "printsetup.tableTech": ("[Table] Technical Data", "[Table] 技術 Data", "[Table] 技术 Data"),
    "printsetup.watermark": ("Watermark", "透かし", "水印"),
    "printsetup.workOrder": ("Work Order", "作業指示書", "工单"),
    "prj.addUpload": ("＋ Upload", "＋ アップロード", "＋ 上传"),
    "prj.approvalWaiting": ("Approval ☑ Pending", "Approval ☑ 待ち", "Approval ☑ 等待"),
    "prj.autoNumber": ("Auto Numbering (PS-)", "自動採番 (PS-)", "自动编号 (PS-)"),
    "prj.clientContact": ("Client Contact", "Client 担当者", "Client 负责人"),
    "prj.date": ("Date", "日付", "日期"),
    "prj.dupCheck": ("Duplicate Check", "重複確認", "重复检查"),
    "prj.fileName": ("File Name", "ファイル名", "文件名"),
    "prj.fileType": ("Type", "種類", "类型"),
    "prj.owner": ("Owner", "担当者", "负责人"),
    "prj.printSetup": ("Print Settings", "Print 設定", "Print 设置"),
    "prj.printSetupCall": ("🖨 Open Print Set-up", "🖨 Print Set-up 呼出", "🖨 调用 Print Set-up"),
    "prj.receivedFiles": ("Register Received Files (File)", "受付資料登録 (File)", "接收资料登记 (File)"),
    "prj.registeredAt": ("Registered", "登録日", "登记日"),
    "prj.registrant": ("Registered By", "登録者", "登记人"),
    "prj.requirements": ("Requirements·Pain Points", "要求事項・Pain Point", "需求·Pain Point"),
    "prj.salesStage": ("Sales Stage", "営業段階", "销售阶段"),
    "prj.saveF12": ("Save F12", "保存 F12", "保存 F12"),
    "prj.stageClosed": ("Closed", "終了", "结束"),
    "prj.stageContract": ("Contract", "契約", "合同"),
    "prj.stageContractChange": ("Contract Change", "契約変更", "合同变更"),
    "prj.stageHint": ("PS starts → then links to PCR→QCR→OR process state machine (ERP-001/002)",
        "PS 開始 → 以降 PCR→QCR→OR プロセス状態機械に連結 (ERP-001/002)",
        "PS 开始 → 之后连接 PCR→QCR→OR 流程状态机 (ERP-001/002)"),
    "prj.stageHistory": ("Sales Stage — change history logged", "営業段階 — 変更履歴記録", "销售阶段 — 记录变更历史"),
    "prj.stageNegotiation": ("Negotiation", "協議", "协商"),
    "prj.stageQuote": ("Quote", "見積", "报价"),
    "prj.stageTechProposal": ("Tech Proposal", "技術提案", "技术提案"),
    "prj.status": ("Status", "状態", "状态"),
    "prj.todoTitle": ("To-do (my processes)", "To-do (自分担当プロセス)", "To-do (本人负责流程)"),
    "prj.unsaved": ("Unsaved", "未保存", "未保存"),
    "procset.auto": ("Auto", "自動", "自动"),
    "procset.autoTransition": ("☑ Run·Event Transition", "☑ Run・イベント遷移", "☑ Run·事件转移"),
    "procset.caution": ("Caution", "注意", "注意"),
    "procset.cautionText": ("※ Changes affecting the System DB require Platform approval",
        "※ System DB に影響する変更は Platform 承認が必要",
        "※ 影响 System DB 的变更需 Platform 审批"),
    "procset.dbSetupHint": ("Define data fields per process", "プロセス別データ項目定義", "按流程定义数据项"),
    "procset.deadlineRule": ("Deadline Rule", "期限ルール", "期限规则"),
    "procset.defTitle": ("Process Definition (erp_process_def / erp_process_edge)",
        "プロセス定義 (erp_process_def / erp_process_edge)",
        "流程定义 (erp_process_def / erp_process_edge)"),
    "procset.editDef": ("Edit Definition — {n}", "定義編集 — {n}", "编辑定义 — {n}"),
    "procset.form": ("Processing Form", "処理 Form", "处理 Form"),
    "procset.formApproval": ("Approval Docs Form", "承認図書 Form", "审批文件 Form"),
    "procset.formOrder": ("Order Form v2", "受注 Form v2", "接单 Form v2"),
    "procset.formPr": ("PO Request Form", "発注依頼 Form", "下单申请 Form"),
    "procset.formSetupHint": ("Processing screens link to EDIM Toolbox UI Form",
        "処理画面は EDIM Toolbox UI Form に連結",
        "处理画面连接 EDIM Toolbox UI Form"),
    "procset.jsonbSchema": ("JSONB Schema", "JSONB スキーマ", "JSONB 模式"),
    "procset.mapHint": ("Select node → edit definition below · 40 preloaded (Slide 10) customized per tenant",
        "ノード選択 → 下段で定義編集・初期搭載 40種 (スライド10) テナント別カスタマイズ",
        "选择节点 → 下方编辑定义 · 初始加载 40种 (幻灯片10) 按租户自定义"),
    "procset.next": ("Successor", "後続", "后续"),
    "procset.notifyHint1": ("Set notification targets on state transition", "状態遷移時の通知対象設定", "设置状态转移时的通知对象"),
    "procset.notifyHint2": ("Overdue → anomaly alert (Dashboard rollup)",
        "期限超過 → 異常警告 (Dashboard 集計)",
        "超过期限 → 异常警告 (Dashboard 汇总)"),
    "procset.notifyRule": ("Notification Rule", "通知ルール", "通知规则"),
    "procset.ownerDeptHead": ("Dept. Head", "部門長", "部门主管"),
    "procset.ownerProject": ("Project Owner", "Project 担当者", "Project 负责人"),
    "procset.ownerPurchase": ("Purchase Owner", "購買担当", "采购负责人"),
    "procset.ownerRule": ("Owner Rule", "担当ルール", "负责规则"),
    "procset.ownerTech": ("Tech Owner", "技術担当", "技术负责人"),
    "procset.precondition": ("Precondition", "先行条件", "前置条件"),
    "procset.prev": ("Predecessor", "先行", "前置"),
    "procset.processMap": ("Process Map — {n}", "プロセスマップ — {n}", "流程图 — {n}"),
    "procset.processName": ("Process Name", "プロセス名", "流程名称"),
    "procset.saveApprovalF12": ("Save → Request Approval F12", "保存 → 承認依頼 F12", "保存 → 请求审批 F12"),
    "purch.createPoF12": ("Create PO → PO-61313-2 F12", "発注作成 → PO-61313-2 F12", "创建 PO → PO-61313-2 F12"),
    "purch.filterGeneral": ("General Orders", "一般発注", "普通订单"),
    "purch.filterIncoming": ("Incoming (5)", "入庫予定 (5)", "待入库 (5)"),
    "purch.filterPoDone": ("Ordered (PO 12)", "発注完了 (PO 12)", "已下单 (PO 12)"),
    "purch.filterPrWait": ("Pending Order (PR 3)", "発注待ち (PR 3)", "待下单 (PR 3)"),
    "purch.flowOrder": ("Order", "発注", "下单"),
    "purch.flowQuoteReq": ("Quote Request", "見積依頼", "报价请求"),
    "purch.poCreated": ("PO-61313-2 created — awaiting RA approval",
        "PO-61313-2 作成 — RA 承認待ち",
        "PO-61313-2 已创建 — 等待 RA 审批"),
    "purch.prTitle": ("Purchase Request", "発注依頼", "采购申请"),
    "purch.priceResolve": ("Price Resolve (CST-001)", "単価 Resolve (CST-001)", "单价 Resolve (CST-001)"),
    "purch.priorityHint": ("Priority: quote-applied → purchase → stock → quote (Price Mgmt M-12-5)",
        "優先順位: 見積適用→購買→在庫→見積 (単価管理 M-12-5)",
        "优先级: 报价应用→采购→库存→报价 (单价管理 M-12-5)"),
    "purch.processTitle": ("Process (per W-14 definition)", "プロセス (W-14 定義準拠)", "流程 (遵循 W-14 定义)"),
    "purch.purchaseTerms": ("Purchase Terms (ERP-017)", "購買条件 (ERP-017)", "采购条件 (ERP-017)"),
    "purch.qcrBtn": ("Quote Request (QCR)", "見積依頼 (QCR)", "报价请求 (QCR)"),
    "purch.selectedCount": ("{n} selected", "選択 {n}件", "已选 {n}项"),
    "purch.stockOk": ("Stock OK", "在庫充足", "库存充足"),
    "purch.supplierCode": ("Supplier Code", "仕入先コード", "供应商代码"),
    "purch.supplierMap": ("Supplier Code Mapping (ERP-018)", "仕入先コードマッピング (ERP-018)", "供应商代码映射 (ERP-018)"),
    "purch.supplierMapHint": ("Show supplier-side code on the PO", "発注書に仕入先側コードを表記", "在采购单上标注供应商侧代码"),
    "purch.termsDelivery": ("Delivery: EXW / FOB / CIP · named place",
        "納品: EXW / FOB / CIP・指定場所",
        "交货: EXW / FOB / CIP · 指定地点"),
    "purch.termsPayment": ("Payment · currency · MOQ · certificates", "支払・通貨・最小数量・証明書要求", "付款·货币·最小数量·证书要求"),
    "run.apRequest": ("AP Request", "AP 依頼", "AP 请求"),
    "run.custApprovalWait": ("Awaiting Customer Approval", "顧客承認待ち", "等待客户批准"),
    "run.elapsed": ("Elapsed", "所要", "耗时"),
    "run.erpSend": ("ERP Send", "ERP 送信", "ERP 传送"),
    "run.execLog": ("Run Log", "実行ログ", "执行日志"),
    "run.file": ("File", "ファイル", "文件"),
    "run.folder": ("Folder", "フォルダ", "文件夹"),
    "run.generated": ("Generated", "生成", "生成"),
    "run.measured": ("Measured", "実測", "实测"),
    "run.message": ("Message", "メッセージ", "消息"),
    "run.nextAction": ("Next Action", "次のアクション", "下一步操作"),
    "run.open": ("Open", "開く", "打开"),
    "run.openFolder": ("Open Folder", "フォルダを開く", "打开文件夹"),
    "run.outputsTitle": ("Outputs — PS-61313-5 (double-click = doc detail)",
        "成果物 — PS-61313-5 (Wクリック＝文書詳細)",
        "产出物 — PS-61313-5 (双击=文档详情)"),
    "run.qcrIssue": ("Issue QCR", "QCR 発行", "签发 QCR"),
    "run.rerunF5": ("Rerun F5", "再実行 F5", "重新执行 F5"),
    "run.running": ("Running", "実行中", "执行中"),
    "run.status": ("Status", "状態", "状态"),
    "run.step": ("Step", "段階", "阶段"),
    "run.successChip": ("SUCCESS 8m 32s · -86% vs 1h target",
        "SUCCESS 8m 32s・目標 1h 比 -86%",
        "SUCCESS 8m 32s · 较目标 1h -86%"),
    "run.task": ("Task", "作業", "作业"),
    "run.time": ("Time", "時刻", "时刻"),
    "run.type": ("Type", "種別", "类型"),
    "run.waitPipeline": ("Shown after the pipeline completes…", "パイプライン完了後に表示されます…", "流水线完成后显示…"),
    "studio.findFeature": ("[ Find Feature ]", "[ 機能検索 ]", "[ 查找功能 ]"),
    "studio.fnWizard": ("[ Function Wizard ]", "[ 関数ウィザード ]", "[ 函数向导 ]"),
    "studio.generateAi": ("▶ Generate (AI)", "▶ 生成 (AI)", "▶ 生成 (AI)"),
    "studio.generating": ("Generating…", "生成中…", "生成中…"),
    "studio.macroTitle": ("Macro — Excel-compatible syntax (editable)",
        "Macro — Excel 互換構文 (編集可)",
        "Macro — Excel 兼容语法 (可编辑)"),
    "studio.notTested": ("Not Tested", "未検証", "未验证"),
    "studio.saveVer": ("Save (v0.3)", "保存 (v0.3)", "保存 (v0.3)"),
    "studio.searchNl": ("Search features in natural language…", "自然言語で機能検索…", "用自然语言搜索功能…"),
    "studio.syncHint": ("4-Way Sync — editing one side suggests syncing the rest · simple calc = Macro / complex = Coding(AI) (TBX-008/010)",
        "4-Way Sync — 一方修正時に残りの同期を提案・簡易計算＝Macro / 複雑＝Coding(AI) (TBX-008/010)",
        "4-Way Sync — 修改一侧时建议同步其余 · 简单计算=Macro / 复杂=Coding(AI) (TBX-008/010)"),
    "studio.syncSuggest": ("Mismatch with Prompt — sync suggested",
        "Prompt と不一致 — 同期を提案",
        "与 Prompt 不一致 — 建议同步"),
    "studio.verifyApprove": ("Verify · Request Approval", "検証・承認依頼", "验证·请求审批"),
    "subcode.addValue": ("＋ Add Value", "＋ 値追加", "＋ 添加值"),
    "subcode.apprStatus": ("Approval Status", "承認状態", "审批状态"),
    "subcode.approved": ("Approved", "承認", "已批准"),
    "subcode.codeAsset": ("Code Assets — KDCR 3-13", "コード資産 — KDCR 3-13", "代码资产 — KDCR 3-13"),
    "subcode.count": ("Count", "件数", "件数"),
    "subcode.desc": ("Description", "説明", "说明"),
    "subcode.dupCheck": ("Duplicate Check", "重複検討", "重复检查"),
    "subcode.item": ("Item", "項目", "项目"),
    "subcode.newF2": ("New F2", "新規 F2", "新建 F2"),
    "subcode.newItemTitle": ("New Item ({n}) — required = yellow cells",
        "新規項目 ({n}) — 必須は黄色セル",
        "新建项目 ({n}) — 必填为黄色单元格"),
    "subcode.none": ("— None", "— なし", "— 无"),
    "subcode.open": ("Open", "開く", "打开"),
    "subcode.queryF8": ("Query F8", "照会 F8", "查询 F8"),
    "subcode.refTable": ("Ref Table", "参照 Table", "引用 Table"),
    "subcode.saveF12": ("Save F12", "保存 F12", "保存 F12"),
    "subcode.status": ("Status", "状態", "状态"),
    "subcode.valueList": ("Value List", "値一覧", "值列表"),
    "taskbox.alerts": ("Anomaly Alerts", "異常警告", "异常警告"),
    "taskbox.apprBox": ("Approval Request Box", "承認依頼箱", "审批请求箱"),
    "taskbox.code": ("Code", "コード", "代码"),
    "taskbox.completeDone": ("Mark Complete (DONE)", "完了処理 (DONE)", "完成处理 (DONE)"),
    "taskbox.content": ("Content", "内容", "内容"),
    "taskbox.deadline": ("Due", "期限", "期限"),
    "taskbox.deadlineRule": ("Due rule: OR+7 days (per W-14)",
        "期限ルール: OR+7日 (W-14 定義)",
        "期限规则: OR+7天 (W-14 定义)"),
    "taskbox.deptAll": ("All Departments", "部署全体", "全部门"),
    "taskbox.filter": ("Filter", "フィルタ", "筛选"),
    "taskbox.fri": ("Fri", "金", "周五"),
    "taskbox.gotoInbox": ("→ Go to Approval Inbox", "→ 承認箱へ移動", "→ 转到审批箱"),
    "taskbox.kind": ("Kind", "区分", "类别"),
    "taskbox.mon": ("Mon", "月", "周一"),
    "taskbox.mrProc": ("MR Production Request", "MR 製作依頼", "MR 制作委托"),
    "taskbox.mySchedule": ("My Schedule (Weekly)", "マイスケジュール (週間)", "我的日程 (周)"),
    "taskbox.myTasks": ("My Tasks", "マイ業務", "我的任务"),
    "taskbox.noAlerts": ("No alerts", "警告なし", "无警告"),
    "taskbox.openForm": ("Open Handling Form (Toolbox)", "処理 Form を開く (Toolbox)", "打开处理 Form (Toolbox)"),
    "taskbox.orDone": ("OR Order DONE", "OR 受注 DONE", "OR 接单 DONE"),
    "taskbox.overdueWarn": ("Over 2 days — Dashboard anomaly alert raised",
        "2日超過、Dashboard 異常警告発生",
        "超过2天，触发 Dashboard 异常警告"),
    "taskbox.owner": ("Owner", "担当", "负责人"),
    "taskbox.plOverdue": ("PL overdue 2 days", "PL 期限超過 2日", "PL 逾期 2天"),
    "taskbox.process": ("Process", "プロセス", "流程"),
    "taskbox.tasksN": ("Tasks — {n}", "業務 — {n}件", "任务 — {n}项"),
    "taskbox.thu": ("Thu", "木", "周四"),
    "taskbox.timeKind": ("Time", "時間", "时间"),
    "taskbox.title": ("Subject", "件名", "标题"),
    "taskbox.tue": ("Tue", "火", "周二"),
    "taskbox.wed": ("Wed", "水", "周三"),
    "techdata.densityCalc": ("Density Correction Sheet", "密度補正計算書", "密度修正计算书"),
    "techdata.designOptions": ("Design Options — approved Sub Code values only (CODE-003)",
        "設計オプション — 承認済 Sub Code 値のみ (CODE-003)",
        "设计选项 — 仅限已批准 Sub Code 值 (CODE-003)"),
    "techdata.eff": ("Efficiency", "効率", "效率"),
    "techdata.fanPerfTable": ("Fan Performance Table (PDF)", "Fan 性能表 (PDF)", "Fan 性能表 (PDF)"),
    "techdata.generated": ("Generated", "生成", "生成"),
    "techdata.graphWizard": ("Graph Wizard Templet (TBX-011)",
        "グラフウィザード Templet (TBX-011)",
        "图表向导 Templet (TBX-011)"),
    "techdata.notSelected": ("Not Selected", "未選定", "未选定"),
    "techdata.perfCurve": ("Performance Curve — selection point highlighted",
        "性能曲線 — 選定点ハイライト",
        "性能曲线 — 高亮选定点"),
    "techdata.queryF8": ("Query F8", "照会 F8", "查询 F8"),
    "techdata.selFanDrawing": ("Selected Fan Drawing — {n}", "選定 Fan 図面 — {n}", "选定 Fan 图纸 — {n}"),
    "techdata.selPoint": ("Selection Point", "選定点", "选定点"),
    "techdata.soundForecast": ("Sound Forecast (Sound {n} dB)", "騒音予測 (Sound {n} dB)", "噪音预测 (Sound {n} dB)"),
    "techdata.techDataTitle": ("Technical Data — Table range query + Macro (TBL-004)",
        "Technical Data — Table 範囲照会 + Macro (TBL-004)",
        "Technical Data — Table 范围查询 + Macro (TBL-004)"),
    "techdata.valid": ("Valid", "有効", "有效"),
    "uidsn.aiDraft": ("UI Draft Proposal", "UI 草案提案", "UI 草案建议"),
    "uidsn.aiHint": ("→ Organize purpose / items / required DB Tables, then propose Templet · invoke for Customizing",
        "→ 用途 / 項目 / 必要 DB Table 整理後 Templet 提案・呼び出して Customizing",
        "→ 整理用途 / 项目 / 所需 DB Table 后建议 Templet · 调用后 Customizing"),
    "uidsn.aiPlaceholder": ("Describe the Application to develop…",
        "開発する Application の説明を入力…",
        "输入要开发的 Application 说明…"),
    "uidsn.aiTitle": ("[ UI Dev AI ]", "[ UI 開発 AI ]", "[ UI 开发 AI ]"),
    "uidsn.bindHint": ("Button → Command Templet (copy · target · Data) · Combo → Data Set-up (Table binding)",
        "Button → Command Templet (コピー・対象・Data)・Combo → Data Set-up (Table バインド)",
        "Button → Command Templet (复制·目标·Data) · Combo → Data Set-up (Table 绑定)"),
    "uidsn.flowHint": ("Click palette to place → edit props → bind action Templet → save·version → approve·publish",
        "パレットクリック配置 → 属性編集 → 動作 Templet バインド → 保存・バージョン → 承認・公開",
        "点击面板放置 → 编辑属性 → 绑定动作 Templet → 保存·版本 → 审批·发布"),
    "uidsn.publish": ("Publish (after approval)", "公開 (承認後)", "发布 (审批后)"),
    "uidsn.saveF12": ("Save F12", "保存 F12", "保存 F12"),
    "uidsn.unsaved": ("Unsaved Changes", "未保存の変更", "未保存更改"),
    "uidsn.widgetBox": ("Widget Box — click = place", "Widget Box — クリック＝配置", "Widget Box — 单击=放置"),
    "wp.buyPath": ("BUY {n} → purchase price resolve", "BUY {n}件 → 購買単価 resolve", "BUY {n}项 → 采购单价 resolve"),
    "wp.costPath": ("Cost Path (Pricing Run)", "原価経路 (Pricing Run)", "成本路径 (Pricing Run)"),
    "wp.makeBuy": ("Make/Buy (double-click to toggle)", "Make/Buy (Wクリック切替)", "Make/Buy (双击切换)"),
    "wp.makePath": ("MAKE {n} → mfg cost (time × labor rate)",
        "MAKE {n}件 → 製造費 (時間×賃率)",
        "MAKE {n}项 → 制造费 (时间×工时费率)"),
    "wp.processDef": ("Process Definition — KDCR 3-13 · {p} @ {w}",
        "工程定義 — KDCR 3-13・{p} @ {w}",
        "工序定义 — KDCR 3-13 · {p} @ {w}"),
    "wp.saveF12": ("Save F12", "保存 F12", "保存 F12"),
    "wp.supplier": ("Supplier", "仕入先", "供应商"),
    "wp.timeMin": ("Time(min)", "Time(分)", "Time(分)"),
    "wp.toggleHint": ("Toggle MAKE/BUY by double-clicking Material Data",
        "MAKE/BUY 切替は Material Data をWクリック",
        "双击 Material Data 切换 MAKE/BUY"),
    "wp.unsaved": ("Unsaved Changes", "未保存の変更", "未保存更改"),
    "menu.plm-drawings": ("Drawing Ledger (M-4-1)", "図面台帳 (M-4-1)", "图纸台账 (M-4-1)"),
    "menu.erp-project": ("Project Registration (S-3-5)", "Project 登録 (S-3-5)", "Project 注册 (S-3-5)"),
    "screen.plm-drawings": ("Drawing Ledger", "図面台帳", "图纸台账"),
}


def seed_v9(cur, tid: int) -> None:
    # 키 단위 멱등 — v7 과 동일 방식 (누락 키만 삽입)
    cur.execute(
        """SELECT DISTINCT field FROM sys_translation
           WHERE tenant_id=%s AND entity_type='UI' AND locale='en'""", (tid,))
    have = {r[0] for r in cur.fetchall()}
    n = 0
    for key, (en, ja, zh) in UI_TRANSLATIONS_V9.items():
        if key in have:
            continue
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                   VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))
            n += 1
    if n:
        logger.info("seed v9 — UI 번역 전면 확장 %d행 삽입 (24화면 크롬, B9)", n)


# ── seed v10 — B13: Arrangement·Templet 샘플 + 메뉴 라벨 갱신 ──

ARRANGEMENTS_V10 = [
    ('ARR-DD2', 'Double Deck 2 표준 구성', 'AHU', '좌우 대칭·상부 배기', '실내 설치·방진 마운트',
     [('KAD 900 FW', 'CENTER', 1), ('KDP 1-21', 'LEFT', 1), ('H 22 380', 'RIGHT', 1)]),
    ('ARR-SD1', 'Single Deck 표준 구성', 'AHU', '단층·전면 흡입', '옥상 설치·방수 커버',
     [('KAD 900 FW', 'CENTER', 1), ('KDP 9', 'REAR', 2)]),
]

TEMPLETS_V10 = [
    ('CMD-COPY', 'COMMAND', {'action': 'copy', 'target': 'selection', 'data': 'clipboard'}),
    ('DATA-TBL12', 'DATA', {'table': 'Table12', 'binding': 'combo', 'keyColumn': 'Key'}),
    ('FORM-QUOTE', 'FORM', {'form': 'CLT', 'placeholders': ['project.no', 'customer', 'bom']}),
]

MENU_UPDATES_V10 = {
    'menu.plm-arr': ('Arrangement Set-Up (M-4-2)', 'Arrangement Set-Up (M-4-2)', 'Arrangement Set-Up (M-4-2)'),
    'menu.tbx-templet': ('Templet Mgmt (S-2-3)', 'Templet 管理 (S-2-3)', 'Templet 管理 (S-2-3)'),
    'screen.plm-arr': ('Arrangement Set-Up', 'Arrangement Set-Up', 'Arrangement Set-Up'),
    'screen.tbx-templet': ('Templet Mgmt', 'Templet 管理', 'Templet 管理'),
}


def seed_v10(cur, tid: int) -> None:
    cur.execute('SELECT 1 FROM arrangement_code WHERE tenant_id=%s LIMIT 1', (tid,))
    if not cur.fetchone():
        for code, name, family, direction, install, comps in ARRANGEMENTS_V10:
            cur.execute(
                """INSERT INTO arrangement_code (tenant_id, arrangement_code, arrangement_name,
                   product_family, direction_option, install_option, approval_status)
                   VALUES (%s,%s,%s,%s,%s,%s,'APPROVED') RETURNING arrangement_id""",
                (tid, code, name, family, direction, install))
            arr_id = cur.fetchone()[0]
            for i, (pc_code, pos, qty) in enumerate(comps):
                cur.execute(
                    'SELECT product_code_id FROM product_code WHERE tenant_id=%s AND main_code=%s',
                    (tid, pc_code))
                pc = cur.fetchone()
                if pc:
                    cur.execute(
                        """INSERT INTO arrangement_component (arrangement_id, product_code_id,
                           position_key, quantity, sort_order) VALUES (%s,%s,%s,%s,%s)""",
                        (arr_id, pc[0], pos, qty, i))
        logger.info('seed v10 — arrangement %d건', len(ARRANGEMENTS_V10))
    cur.execute('SELECT 1 FROM tbx_templet WHERE tenant_id=%s LIMIT 1', (tid,))
    if not cur.fetchone():
        for name, ttype, definition in TEMPLETS_V10:
            cur.execute(
                """INSERT INTO tbx_templet (tenant_id, templet_name, templet_type, definition,
                   approval_status) VALUES (%s,%s,%s,%s,'APPROVED')""",
                (tid, name, ttype, json.dumps(definition)))
        logger.info('seed v10 — templet %d건', len(TEMPLETS_V10))
    # 메뉴 라벨 갱신 (— 예정 해제) + 신규 화면 키 — 키 단위 멱등
    for key, (en, ja, zh) in MENU_UPDATES_V10.items():
        for locale, text in (('en', en), ('ja', ja), ('zh', zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v11 — B13-2: 재질·검증 규칙 샘플 + 메뉴 라벨 갱신 ──

MATERIALS_V11 = [
    ('SS400', '일반 구조용 압연강', 'STEEL', 7.85, 'KS D 3503', None),
    ('STS304', '스테인리스 강판', 'STEEL', 7.93, 'KS D 3698', None),
    ('AL6061', '알루미늄 합금', 'ALUMINUM', 2.70, 'KS D 6701', None),
    ('SPHC', '열간압연 강판', 'STEEL', 7.85, 'KS D 3501', None),
]

VERIFICATIONS_V11 = [
    ('B 치수 상한 검증', 'DIM B (KDCR 3-13)', 'B(H) 가 900 을 초과하면 Casing 표준 강판 폭 초과 — 설계 재검토'),
    ('K 전장 간섭 검증', 'DIM K (KDCR 3-13)', 'K 전장이 1,500 초과 시 표준 베이스 프레임 간섭 — Arrangement 확인'),
]

MENU_UPDATES_V11 = {
    'menu.code-variant': ('Variant·Constant (S-1-2)', 'Variant·Constant (S-1-2)', 'Variant·Constant (S-1-2)'),
    'menu.code-raw': ('Raw Material·GPI (M-3-2)', 'Raw Material·GPI (M-3-2)', 'Raw Material·GPI (M-3-2)'),
    'menu.plm-material': ('Material (M-4-4)', 'Material (M-4-4)', 'Material (M-4-4)'),
    'menu.plm-quality': ('Quality (M-4-5)', 'Quality (M-4-5)', 'Quality (M-4-5)'),
    'screen.code-variant': ('Variant·Constant', 'Variant·Constant', 'Variant·Constant'),
    'screen.code-raw': ('Raw Material·GPI', 'Raw Material·GPI', 'Raw Material·GPI'),
    'screen.plm-material': ('Material', 'Material', 'Material'),
    'screen.plm-quality': ('Quality', 'Quality', 'Quality'),
    'menu.com-search': ('Global Search (⌘K)', '統合検索 (⌘K)', '全局搜索 (⌘K)'),
}


def seed_v11(cur, tid: int) -> None:
    cur.execute('SELECT 1 FROM mat_material WHERE tenant_id=%s LIMIT 1', (tid,))
    if not cur.fetchone():
        for code, name, mtype, density, std, hazard in MATERIALS_V11:
            cur.execute(
                """INSERT INTO mat_material (tenant_id, material_code, material_name,
                   material_type, density, standard, hazard_class)
                   VALUES (%s,%s,%s,%s,%s,%s,%s)""",
                (tid, code, name, mtype, density, std, hazard))
        logger.info('seed v11 — mat_material %d건', len(MATERIALS_V11))
    cur.execute('SELECT 1 FROM dwg_verification WHERE tenant_id=%s LIMIT 1', (tid,))
    if not cur.fetchone():
        cur.execute(
            """SELECT drawing_id FROM dwg_drawing
               WHERE tenant_id=%s AND drawing_no='KDCR 3-13' LIMIT 1""", (tid,))
        d = cur.fetchone()
        if d:
            n = 0
            for rule, macro_name, warning in VERIFICATIONS_V11:
                cur.execute(
                    'SELECT macro_id FROM tbx_macro WHERE tenant_id=%s AND macro_name=%s',
                    (tid, macro_name))
                m = cur.fetchone()
                if m:
                    cur.execute(
                        """INSERT INTO dwg_verification (tenant_id, drawing_id, rule_name,
                           macro_id, warning_message) VALUES (%s,%s,%s,%s,%s)""",
                        (tid, d[0], rule, m[0], warning))
                    n += 1
            logger.info('seed v11 — dwg_verification %d건', n)
    for key, (en, ja, zh) in MENU_UPDATES_V11.items():
        for locale, text in (('en', en), ('ja', ja), ('zh', zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v12 — B14: 역할·권한 매트릭스 + Hierarchy 주소 체계 ──

ROLES_V12 = [
    ('PLATFORM', '플랫폼 관리 — 시스템 정의 변경'),
    ('ADMIN', '관리자 — 사용자·권한·잠금해제'),
    ('SETUP', '설계 Set-up — 코드·도면·Macro 편집'),
    ('GENERAL', '일반 — 조회·업무 처리'),
]

# (role, screenId, action) — WRITE ⊃ READ
PERMS_V12 = [
    ('PLATFORM', '*', 'WRITE'),
    ('ADMIN', 'erp-access', 'WRITE'), ('ADMIN', 'com-approval', 'WRITE'),
    ('ADMIN', 'erp-dashboard', 'READ'), ('ADMIN', 'cpq-selection', 'READ'),
    ('SETUP', 'cpq-selection', 'WRITE'), ('SETUP', 'plm-design', 'WRITE'),
    ('SETUP', 'code-subcode', 'WRITE'), ('SETUP', 'code-datatable', 'WRITE'),
    ('SETUP', 'tbx-macro', 'WRITE'), ('SETUP', 'erp-price', 'WRITE'),
    ('SETUP', 'com-approval', 'WRITE'), ('SETUP', 'erp-access', 'READ'),
    ('GENERAL', 'cpq-selection', 'READ'), ('GENERAL', 'erp-dashboard', 'READ'),
    ('GENERAL', 'com-tasks', 'WRITE'), ('GENERAL', 'com-folder', 'READ'),
]

# (tree_type, name, symbol, address, parent_address)
HIERARCHY_V12 = [
    ('PRODUCT', 'Code', 'C', '/C', None),
    ('PRODUCT', 'Engineering', 'ENG', '/C/ENG', '/C'),
    ('PRODUCT', 'Fan', 'FAN', '/C/ENG/FAN', '/C/ENG'),
    ('PRODUCT', 'Macro', 'M', '/M', None),
    ('PRODUCT', 'ENG Macro', 'ENG', '/M/ENG', '/M'),
    ('PRODUCT', 'Fan Macro', 'FAN', '/M/ENG/FAN', '/M/ENG'),
    ('GENERAL_DB', 'Table', 'T', '/T', None),
    ('GENERAL_DB', 'ENG Table', 'ENG', '/T/ENG', '/T'),
    ('GENERAL_DB', 'Variant', 'VAR', '/T/ENG/VARIANT', '/T/ENG'),
]


def seed_v12(cur, tid: int) -> None:
    cur.execute('SELECT 1 FROM sys_role WHERE tenant_id=%s LIMIT 1', (tid,))
    if not cur.fetchone():
        role_ids = {}
        for name, desc in ROLES_V12:
            cur.execute(
                """INSERT INTO sys_role (tenant_id, role_name, description)
                   VALUES (%s,%s,%s) RETURNING role_id""", (tid, name, desc))
            role_ids[name] = cur.fetchone()[0]
        for role, key, action in PERMS_V12:
            cur.execute(
                """INSERT INTO sys_role_permission (role_id, resource_type, resource_key, action)
                   VALUES (%s,'MENU',%s,%s)""", (role_ids[role], key, action))
        logger.info('seed v12 — sys_role %d + permission %d', len(ROLES_V12), len(PERMS_V12))
    cur.execute('SELECT 1 FROM sys_hierarchy WHERE tenant_id=%s LIMIT 1', (tid,))
    if not cur.fetchone():
        ids = {}
        for ttype, name, symbol, address, parent in HIERARCHY_V12:
            cur.execute(
                """INSERT INTO sys_hierarchy (tenant_id, parent_id, tree_type, node_name,
                   symbol, address, approval_status)
                   VALUES (%s,%s,%s,%s,%s,%s,'APPROVED') RETURNING hierarchy_id""",
                (tid, ids.get(parent), ttype, name, symbol, address))
            ids[address] = cur.fetchone()[0]
        logger.info('seed v12 — sys_hierarchy %d', len(HIERARCHY_V12))


# ── seed v13 — B13/B14 신규 화면 메뉴·화면 번역 키 (sys_translation) ──

MENU_UPDATES_V13 = {
    'menu.code-hierarchy': ('Hierarchy Address (M-3-1)', 'Hierarchy アドレス (M-3-1)', 'Hierarchy 地址 (M-3-1)'),
    'menu.erp-company-master': ('Suppliers & Customers (M-14-2)', '仕入先·取引先 (M-14-2)', '供应商·客户 (M-14-2)'),
    'screen.code-hierarchy': ('Hierarchy Address', 'Hierarchy アドレス', 'Hierarchy 地址'),
    'screen.erp-company-master': ('Suppliers & Customers', '仕入先·取引先', '供应商·客户'),
}


def seed_v13(cur, tid: int) -> None:
    # 키 단위 멱등 — UPDATE 후 미존재 시 INSERT (v10/v11 과 동일 패턴)
    for key, (en, ja, zh) in MENU_UPDATES_V13.items():
        for locale, text in (('en', en), ('ja', ja), ('zh', zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v14 — B16: 도면 블록(dwg_document)·단계별 승인(dwg_approval)·부품 관계(dwg_part_relation) ──

# Design Editor Block 패널 원천 (KDCR 3-13) — content = 캔버스 좌표/치수
BLOCKS_V14 = [
    ('Casing', {'x': 190, 'y': 60, 'w': 220, 'h': 180}, 190, 60),
    ('Impeller', {'x': 230, 'y': 80, 'w': 140, 'h': 100, 'sub': 'Airfoil 900'}, 230, 80),
    ('Shaft', {'x': 80, 'y': 150, 'w': 440, 'h': 12}, 80, 150),
    ('Brg-L', {'x': 96, 'y': 138, 'w': 26, 'h': 34}, 96, 138),
    ('Brg-R', {'x': 478, 'y': 138, 'w': 26, 'h': 34}, 478, 138),
    ('Inlet Cone-L', {'x': 150, 'y': 96, 'w': 40, 'h': 110, 'dashed': True}, 150, 96),
    ('Inlet Cone-R', {'x': 410, 'y': 96, 'w': 40, 'h': 110, 'dashed': True}, 410, 96),
]

# (block_a, block_b, align, contact, macro_name, priority)
RELATIONS_V14 = [
    ('Impeller', 'Shaft', 'center', 'face', 'Shaft 길이 계산', 1),
    ('Casing', 'Inlet Cone-L', 'vertical', 'line', None, 2),
    ('Casing', 'Inlet Cone-R', 'vertical', 'line', None, 2),
]

TAB_LABELS_V14 = {
    'dwg.detailTitle': ('Drawing Detail', '図面詳細', '图纸详情'),
    'dwg.tabRev': ('Rev History', 'Rev 履歴', 'Rev 历史'),
    'dwg.tabApproval': ('Approval Steps', '承認ステップ', '审批步骤'),
    'dwg.tabVariants': ('Variants', 'Variants', 'Variants'),
    'dwg.tabReferencers': ('Referencers', 'Referencers', 'Referencers'),
    'dwg.tabAttachment': ('Attachment', '添付', '附件'),
    'dwg.needBackend': ('Backend connection required', 'バックエンド接続が必要', '需要后端连接'),
    'dwg.noVariants': ('No drawings in the same family', '同一ファミリの図面なし', '无同族图纸'),
    'dwg.noRefs': ('No parent references this drawing (code)', 'この図面(コード)を参照する上位なし', '无上级引用此图纸(代码)'),
    'dwg.noFiles': ('No linked files (auto-linked when Run outputs are generated)',
                    'リンクされたファイルなし (Run 生成時に自動リンク)', '无关联文件 (Run 生成时自动关联)'),
    'dwg.motherCol': ('Parent Code', '上位コード', '上级代码'),
    'dwg.descCol': ('Description', '説明', '说明'),
    'dwg.stepCol': ('Step', 'ステップ', '步骤'),
    'dwg.resultCol': ('Result', '結果', '结果'),
    'dwg.chainDone': ('Approval chain complete', '承認チェーン完了', '审批链完成'),
    'dwg.noSteps': ('No approval history — start from WRITE step', '承認履歴なし — 作成から進行', '无审批记录 — 从编写开始'),
    'dwg.stepCommentPh': ('Comment (required on reject)', 'コメント (却下時必須)', '备注 (驳回时必填)'),
    'dwg.approveBtn': ('Approve', '承認', '批准'),
    'dwg.rejectBtn': ('Reject', '却下', '驳回'),
    'editor.simPanel': ('Simulation', 'シミュレーション', '仿真'),
    'editor.simHint': ('Change a VARIANT dim — MACRO dims re-evaluate instantly (no save)',
                       'VARIANT 寸法を変更 — MACRO 寸法が即時再計算 (保存なし)',
                       '修改 VARIANT 尺寸 — MACRO 尺寸即时重算 (不保存)'),
    'editor.simApply': ('Apply (commit dims)', '適用 (寸法反映)', '应用 (提交尺寸)'),
    'editor.relAlign': ('Align', '整列', '对齐'),
    'editor.relContact': ('Contact', '接触', '接触'),
}


def seed_v14(cur, tid: int) -> None:
    cur.execute(
        """SELECT drawing_id FROM dwg_drawing WHERE tenant_id=%s AND drawing_no='KDCR 3-13'""",
        (tid,))
    d = cur.fetchone()
    if d:
        did = d[0]
        cur.execute('SELECT 1 FROM dwg_document WHERE tenant_id=%s LIMIT 1', (tid,))
        if not cur.fetchone():
            for name, content, ox, oy in BLOCKS_V14:
                cur.execute(
                    """INSERT INTO dwg_document (tenant_id, drawing_id, block_name, content,
                       origin_x, origin_y) VALUES (%s,%s,%s,%s,%s,%s)""",
                    (tid, did, name, json.dumps(content), ox, oy))
            logger.info('seed v14 — dwg_document %d블록', len(BLOCKS_V14))
        cur.execute('SELECT 1 FROM dwg_approval WHERE drawing_id=%s LIMIT 1', (did,))
        if not cur.fetchone():
            cur.execute("SELECT user_id FROM sys_user WHERE tenant_id=%s AND login_id='edim'", (tid,))
            uid = cur.fetchone()[0]
            for step in ('WRITE', 'REVIEW', 'APPROVE'):
                cur.execute(
                    """INSERT INTO dwg_approval (drawing_id, step, approver_id, approval_date,
                       result, comment) VALUES (%s,%s,%s,CURRENT_DATE,'APPROVED','최초 발행 승인 체인')""",
                    (did, step, uid))
            logger.info('seed v14 — dwg_approval 3단계 (KDCR 3-13)')
        cur.execute('SELECT 1 FROM dwg_part_relation WHERE tenant_id=%s LIMIT 1', (tid,))
        if not cur.fetchone():
            for a, b, align, contact, macro_name, prio in RELATIONS_V14:
                mid = None
                if macro_name:
                    cur.execute('SELECT macro_id FROM tbx_macro WHERE tenant_id=%s AND macro_name=%s',
                                (tid, macro_name))
                    m = cur.fetchone()
                    mid = m[0] if m else None
                cur.execute(
                    """INSERT INTO dwg_part_relation (tenant_id, drawing_id, block_a, block_b,
                       condition_align, condition_contact, value_macro_id, priority, approval_status)
                       VALUES (%s,%s,%s,%s,%s,%s,%s,%s,'APPROVED')""",
                    (tid, did, a, b, align, contact, mid, prio))
            logger.info('seed v14 — dwg_part_relation %d건', len(RELATIONS_V14))
    for key, (en, ja, zh) in TAB_LABELS_V14.items():
        for locale, text in (('en', en), ('ja', ja), ('zh', zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v18 — F1 프로젝트 도메인: 접수 자료 실데이터화 + 신규 UI 키 ──

UI_TRANSLATIONS_V18: dict[str, tuple[str, str, str]] = {
    "prj.name": ("Project Name", "プロジェクト名", "项目名称"),
    "prj.newF2": ("＋ New F2", "＋ 新規 F2", "＋ 新建 F2"),
    "prj.ledger": ("Project Ledger — {n}", "プロジェクト台帳 — {n}件", "项目台账 — {n}件"),
    "prj.activeCtx": ("Selection = titlebar context", "選択 = タイトルバーのコンテキスト", "选择 = 标题栏上下文"),
    "prj.noFiles": ("No received files — register via + Upload (MinIO RECEIVED + dwg_file)",
                    "受領資料なし — ＋アップロードで登録 (MinIO RECEIVED + dwg_file)",
                    "暂无接收资料 — 通过＋上传登记 (MinIO RECEIVED + dwg_file)"),
    "prj.regTitle": ("Register Project — prj_project", "プロジェクト登録 — prj_project", "项目登记 — prj_project"),
    "prj.regNeedName": ("Enter the project name", "プロジェクト名を入力してください", "请输入项目名称"),
    "prj.regClientPh": ("Client — auto-created if absent (com_company)",
                        "顧客 — 未登録なら自動作成 (com_company)",
                        "客户 — 不存在时自动创建 (com_company)"),
    "prj.regSubmit": ("Register F12", "登録 F12", "登记 F12"),
    "common.needBackend": ("Backend connection required (mock mode)",
                           "バックエンド接続が必要 (mock モード)", "需要后端连接 (mock 模式)"),
    "shell.noProject": ("No project selected", "プロジェクト未選択", "未选择项目"),
}


def seed_v18(cur, tid: int) -> None:
    # 신규 UI 키 업서트 (키 단위 멱등)
    for key, (en, ja, zh) in UI_TRANSLATIONS_V18.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))
    # 접수 자료 실데이터화 — mock RECEIVED_FILES → dwg_file + MinIO 실객체
    cur.execute(
        """SELECT 1 FROM dwg_file WHERE tenant_id=%s AND folder='RECEIVED'
           AND file_name='Micron7_사양서_v2.xlsx'""", (tid,))
    if cur.fetchone():
        return
    cur.execute(
        "SELECT project_id FROM prj_project WHERE tenant_id=%s AND project_no='PS-61313-5'",
        (tid,))
    prj = cur.fetchone()
    if not prj:
        return
    try:
        import io

        from openpyxl import Workbook
        from reportlab.pdfgen import canvas

        from app.services import storage

        wb = Workbook()
        ws = wb.active
        ws.title = "사양"
        for row in (["항목", "값"], ["풍량", "690 m3/min"], ["정압", "85 mmAq"],
                    ["전원", "440V 60Hz"], ["비고", "Micron #7 Pre-Sales 접수 사양"]):
            ws.append(row)
        xbuf = io.BytesIO()
        wb.save(xbuf)

        pbuf = io.BytesIO()
        c = canvas.Canvas(pbuf, pagesize=(595, 842))
        c.setFont("Helvetica-Bold", 14)
        c.drawString(80, 780, "Site Layout - Micron #7 (received document)")
        c.setFont("Helvetica", 10)
        c.rect(80, 400, 430, 320)
        c.rect(100, 560, 140, 110)
        c.drawString(110, 610, "AHU #1")
        c.rect(280, 560, 140, 110)
        c.drawString(290, 610, "Fan Room")
        c.drawString(90, 380, "Received 2026-07-07 / Pre-Sales")
        c.showPage()
        c.save()

        objects = (
            ("Micron7_사양서_v2.xlsx", "XLSX", xbuf.getvalue(),
             "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
            ("현장 배치도.pdf", "PDF", pbuf.getvalue(), "application/pdf"),
        )
        # 스토리지 업로드 먼저 전부 성공해야 INSERT (부분 등재 방지)
        for fname, _t, data, ctype in objects:
            storage.put_object(f"PS-61313-5/RECEIVED/{fname}", data, ctype)
        for fname, ftype, data, _c in objects:
            cur.execute(
                """INSERT INTO dwg_file (tenant_id, project_id, folder, file_name, file_type,
                   file_path, file_size, uploaded_date)
                   VALUES (%s,%s,'RECEIVED',%s,%s,%s,%s,'2026-07-07')""",
                (tid, prj[0], fname, ftype, f"PS-61313-5/RECEIVED/{fname}", len(data)))
        logger.info("seed v18 — 접수 자료 2건 실데이터화 (RECEIVED, MinIO)")
    except Exception as e:  # noqa: BLE001 — storage 불가 시 다음 기동에서 재시도
        logger.warning("seed v18 skip: %s", e)


# ── seed v19 — F2 사용자 등록·수정 UI 키 ──

UI_TRANSLATIONS_V19: dict[str, tuple[str, str, str]] = {
    "access.editUser": ("Edit Info", "情報修正", "编辑信息"),
    "access.regTitle": ("Register User — sys_user", "ユーザー登録 — sys_user", "用户登记 — sys_user"),
    "access.regLoginPh": ("lowercase·digits·._- min 3 chars", "小文字・数字・._- 3文字以上", "小写·数字·._- 至少3位"),
    "access.email": ("Email", "メール", "邮箱"),
    "access.initPw": ("Initial Password", "初期パスワード", "初始密码"),
    "access.initPwPh": ("min 4 chars — change after handover", "4文字以上 — 引き渡し後に変更推奨", "至少4位 — 交付后建议修改"),
    "access.regNeedFields": ("login, name and initial password are required",
                             "login・氏名・初期パスワードは必須です", "login·姓名·初始密码为必填"),
    "access.regNeedName": ("Name cannot be empty", "氏名は空にできません", "姓名不能为空"),
    "access.editTitle": ("Edit Info — {n}", "情報修正 — {n}", "编辑信息 — {n}"),
    "access.editSubmit": ("Save F12", "保存 F12", "保存 F12"),
}


def seed_v19(cur, tid: int) -> None:
    for key, (en, ja, zh) in UI_TRANSLATIONS_V19.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v20 — F3 권한 게이팅 UI 키 ──

UI_TRANSLATIONS_V20: dict[str, tuple[str, str, str]] = {
    "perm.needSetup": ("Insufficient permission — SETUP or higher required (SYS-005)",
                       "権限不足 — SETUP 以上が必要 (SYS-005)", "权限不足 — 需要 SETUP 及以上 (SYS-005)"),
    "perm.needAdmin": ("Insufficient permission — ADMIN only (SYS-005)",
                       "権限不足 — ADMIN 専用 (SYS-005)", "权限不足 — 仅限 ADMIN (SYS-005)"),
    "perm.deniedTitle": ("Access denied (403)", "アクセス権限なし (403)", "无访问权限 (403)"),
    "perm.deniedBody": ("{s} is accessible to {n} or higher — request permission from an administrator (SYS-005)",
                        "{s} は {n} 以上のみアクセス可能 — 管理者に権限を申請してください (SYS-005)",
                        "{s} 仅限 {n} 及以上访问 — 请向管理员申请权限 (SYS-005)"),
    "appr.myReqsN": ("My Requests — {n}", "自分の申請 — {n}件", "我的申请 — {n}件"),
    "appr.noMine": ("No pending requests of mine (requester = me)",
                    "自分が申請した保留中の件はありません (requester = 本人)",
                    "没有我提交的待处理申请 (requester = 本人)"),
    "appr.noPending": ("No pending requests", "保留中の申請はありません", "没有待处理的申请"),
    "appr.readOnly": ("Read-only — no decision permission", "読み取り専用 — 決裁権限なし", "只读 — 无审批权限"),
}


def seed_v20(cur, tid: int) -> None:
    for key, (en, ja, zh) in UI_TRANSLATIONS_V20.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v21 — F4 무반응 일소 UI 키 ──

UI_TRANSLATIONS_V21: dict[str, tuple[str, str, str]] = {
    "subcode.addValueHint": ('Add value — type the new value at the end of Sub Item (separated by ·)', '値追加 — Sub Item 末尾に新しい値を入力 (·区切り)', '添加值 — 在 Sub Item 末尾输入新值 (·分隔)'),
    "dtable.editNeedRow": ('Edit — select a row first', '編集 — 行を先に選択してください', '编辑 — 请先选择行'),
    "printsetup.bindNeedSel": ('Bind data — select a widget on the canvas first', 'Data 位置指定 — キャンバスでウィジェットを先に選択', '数据绑定 — 请先在画布中选择控件'),
    "printsetup.defaultDone": ('Default layout ✓ — 6 widgets (header·Data×2·graph·table·footer)', '基本様式配置 ✓ — 6ウィジェット', '默认版式 ✓ — 6个控件'),
    "printsetup.dataAdded": ("Data widget added ✓ — bind a path via [Bind Data]", 'Data ウィジェット追加 ✓ — Data 位置指定でパスをバインド', '已添加数据控件 ✓ — 通过数据绑定指定路径'),
    "printsetup.graphAdded": ('Graph widget added ✓ — [graph:performance]', 'グラフウィジェット追加 ✓ — [グラフ:performance]', '已添加图表控件 ✓ — [图表:performance]'),
    "printsetup.printerHint": ('Real render in new window — print via browser (Ctrl+P)', '実レンダー新ウィンドウ — ブラウザ印刷 (Ctrl+P)', '新窗口实渲染 — 浏览器打印 (Ctrl+P)'),
    "printsetup.officeHint": ('After customer templates are finalized (P4-1 — DOCX/XLSX forms)', '顧客様式確定後 (P4-1 — DOCX/XLSX)', '客户模板确认后 (P4-1 — DOCX/XLSX)'),
    "printsetup.bindTitle": ('Bind Data — placeholder binding', 'Data 位置指定 — バインディング', '数据绑定 — 占位符绑定'),
    "printsetup.bindTarget": ('Target widget', '対象ウィジェット', '目标控件'),
    "printsetup.bindPath": ('Data path', 'データパス', '数据路径'),
    "printsetup.bindSubmit": ('Bind F12', 'バインド F12', '绑定 F12'),
}


def seed_v21(cur, tid: int) -> None:
    for key, (en, ja, zh) in UI_TRANSLATIONS_V21.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v22 — F5 마스터 수정 UI 키 ──

UI_TRANSLATIONS_V22: dict[str, tuple[str, str, str]] = {
    "qedit.required": ('Required: {n}', '必須入力: {n}', '必填: {n}'),
    "qedit.saveF12": ('Save F12', '保存 F12', '保存 F12'),
    "parts.editBtn": ('Edit', '修正', '编辑'),
    "wh.editBtn": ('Edit', '修正', '编辑'),
    "price.closeBtn": ('Close Validity', '適用終了', '结束适用'),
    "price.closeSubmit": ('Close F12', '締切 F12', '结束 F12'),
    "docmgmt.editMeta": ('Edit Meta', 'メタ修正', '编辑元数据'),
    "docmgmt.title": ('Title', 'タイトル', '标题'),
}


def seed_v22(cur, tid: int) -> None:
    for key, (en, ja, zh) in UI_TRANSLATIONS_V22.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v23 — F6 통합 검색 확장 UI 키 ──

UI_TRANSLATIONS_V23: dict[str, tuple[str, str, str]] = {
    "search.parts": ('Parts', '部品', '零件'),
    "search.companies": ('Companies', '取引先', '供应商·客户'),
    "search.warehouses": ('Warehouse·Location', '倉庫・位置', '仓库·位置'),
    "search.projects": ('Projects', 'プロジェクト', '项目'),
    "search.users": ('Users', 'ユーザー', '用户'),
    "shell.searchPh": ('Search screens·codes·parts·companies·docs (⌘K)', '画面・コード・部品・取引先・文書検索 (⌘K)', '搜索画面·代码·零件·供应商·文档 (⌘K)'),
}


def seed_v23(cur, tid: int) -> None:
    for key, (en, ja, zh) in UI_TRANSLATIONS_V23.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v24 — F7 이력 diff 뷰어 UI 키 ──

UI_TRANSLATIONS_V24: dict[str, tuple[str, str, str]] = {
    "folder.diffNoPayload": ('diff — this action has no change payload (read-only event)', 'diff — この操作には変更ペイロードがありません (照会イベント)', 'diff — 此操作没有变更载荷 (查询类事件)'),
    "folder.diffTitle": ('diff — {t} · {a}', 'diff — {t} · {a}', 'diff — {t} · {a}'),
    "folder.diffField": ('Field', 'フィールド', '字段'),
    "folder.diffEmpty": ('No recorded fields', '記録されたフィールドなし', '无记录字段'),
}


def seed_v24(cur, tid: int) -> None:
    for key, (en, ja, zh) in UI_TRANSLATIONS_V24.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# ── seed v25 — F10 조회 UX 소품 UI 키 ──

UI_TRANSLATIONS_V25: dict[str, tuple[str, str, str]] = {
    "dash.kpiDrillHint": ('Click = open the owning screen', 'クリック = 担当画面を開く', '点击 = 打开对应画面'),
    "appr.etcTypes": ('Relation·Doc·Others', '関係・文書・その他', '关系·文档·其他'),
    "appr.searchPh": ('Search target·requester', '対象・申請者検索', '搜索对象·申请人'),
}


def seed_v25(cur, tid: int) -> None:
    for key, (en, ja, zh) in UI_TRANSLATIONS_V25.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


UI_TRANSLATIONS_V27: dict[str, tuple[str, str, str]] = {
    "menu.common": ('Common', '共通', '公共'),
    "menu.cpq-run": ('Run Pipeline (C-1)', 'Run パイプライン (C-1)', 'Run 流水线 (C-1)'),
    "menu.tbx-assistant": ('Internal Q&A (AI-08)', '社内 Q&A (AI-08)', '内部 Q&A (AI-08)'),
    "menu.com-detail": ('Detail Views (Drill-down)', '詳細照会 (ドリルダウン)', '详情查询 (下钻)'),
    "menu.detail-code": ('Code Detail', 'コード詳細', '代码详情'),
    "menu.detail-part": ('Part Detail (G3-b)', '部品詳細 (G3-b)', '部件详情 (G3-b)'),
    "menu.detail-event": ('Event Detail (E-4)', 'イベント詳細 (E-4)', '事件详情 (E-4)'),
    "menu.detail-output": ('Output Doc Detail (G3-a)', '成果物文書詳細 (G3-a)', '产出文档详情 (G3-a)'),
    "shell.logout": ('Logout', 'ログアウト', '登出'),
    "login.title": ('EDIM Login', 'EDIM ログイン', 'EDIM 登录'),
    "login.userId": ('Login ID', '社員番号', '工号'),
    "login.password": ('Password', 'パスワード', '密码'),
    "login.submit": ('Login (Enter)', 'ログイン (Enter)', '登录 (Enter)'),
    "login.checking": ('Checking…', '確認中…', '验证中…'),
    "editor.blockNone": ('(no registered blocks)', '(登録 Block なし)', '(无已注册 Block)'),
    "editor.blockTitle": ('Block — save unit · call from parent', 'Block — 単位保存・上位呼び出し', 'Block — 单元保存·上级调用'),
    "editor.blockNamePh": ('Block name (alphanumeric)', 'Block 名 (英数字)', 'Block 名称 (字母/数字)'),
    "macro.namePh": ('Macro name (TBX-…)', 'Macro 名 (TBX-…)', 'Macro 名称 (TBX-…)'),
    "macro.promptPh": ('Prompt (natural language)', 'Prompt (自然言語説明)', 'Prompt (自然语言说明)'),
    "macro.varsPh": ('Test vars: A=560, B=2', 'Test 変数: A=560, B=2', '测试变量: A=560, B=2'),
    "macro.descCol": ('Description', '説明', '说明'),
    "macro.exprCol": ('Expression', '数式', '公式'),
    "macro.saveF12": ('Save (F12)', '保存 (F12)', '保存 (F12)'),
    "templet.namePh": ('Templet name', 'Templet 名', 'Templet 名称'),
    "templet.sysCol": ('System', 'システム', '系统'),
    "templet.custom": ('Custom', 'カスタム', '自定义'),
    "templet.saveDraft": ('Save (DRAFT)', '保存 (DRAFT)', '保存 (DRAFT)'),
    "docmgmt.keepGrade": ('Keep grade', 'Grade 維持', '保持 Grade'),
    "docmgmt.docNo": ('Doc No', '文書番号', '文档编号'),
    "docmgmt.appDate": ('Approved on', '承認日', '批准日'),
    "docmgmt.approver": ('Approver', '承認者', '批准人'),
    "docmgmt.newTitlePh": ('New title (meta edit)', '新タイトル (メタ修正)', '新标题 (元数据修改)'),
    "printsetup.publish": ('Publish (after approval)', '公開 (承認後)', '发布 (审批后)'),
    "printsetup.form": ('Form', '様式', '表单'),
    "printsetup.canvasTitle": ('Form canvas — [Data]/[Chart]/[Table] placeholders', '様式キャンバス — [Data]/[グラフ]/[Table] プレースホルダ', '表单画布 — [Data]/[图表]/[Table] 占位符'),
    "printsetup.placeholderList": ('Placeholder list', 'プレースホルダ一覧', '占位符列表'),
    "uidsn.publish": ('Publish (after approval)', '公開 (承認後)', '发布 (审批后)'),
    "cpq.quoteLoad": ('Load quotation…', '見積案読込…', '加载报价方案…'),
    "cpq.quoteSave": ('Save F12', '保存 F12', '保存 F12'),
    "cpq.unitPriceK": ('Unit (K)', '単価(K)', '单价(K)'),
    "cpq.liveDb": ('Live DB', '実 DB', '实库'),
    "cpq.block": ('Block', 'ブロック', '块'),
    "techdata.totalPt": ('Total Pt', '全圧 Pt', '全压 Pt'),
    "techdata.staticPd": ('Static Pd', '静圧 Pd', '静压 Pd'),
    "approval.commentPh": ('Comment (required to reject)', '決裁意見 (差戻し時必須)', '审批意见 (驳回时必填)'),
    "approval.testFail": ('Fail', '未通過', '未通过'),
    "appr.drawing": ('Drawing', '図面', '图纸'),
    "appr.filterOff": ('Hide type', '種別非表示', '隐藏类型'),
    "appr.filterOn": ('Show type', '種別表示', '显示类型'),
    "task.commentPh": ('Comment (optional)', '処理意見 (任意)', '处理意见 (可选)'),
    "folder.uploadBtn": ('⬆ Upload', '⬆ アップロード', '⬆ 上传'),
    "folder.zipAll": ('⬇ ZIP (all)', '⬇ ZIP (全体)', '⬇ ZIP (全部)'),
    "folder.registrant": ('Registrant', '登録者', '登记人'),
    "folder.dateCol": ('Date', '日付', '日期'),
    "folder.kindCol": ('Kind', '種別', '种类'),
    "mobile.qty": ('Qty', '数量', '数量'),
    "mobile.complete": ('Mark complete', '完了処理', '完成处理'),
    "mobile.delayed": ('Delayed', '遅延', '延迟'),
    "enum.delayed": ('Delayed', '遅延', '延迟'),
    "panel.upload": ('Upload', 'アップロード', '上传'),
    "detail.supplier": ('Supplier', '仕入先', '供应商'),
    "detail.procCol": ('Process', '工程', '工序'),
    "detail.process": ('Process', '工程', '工序'),
    "detail.unit": ('Unit', '単位', '单位'),
    "detail.qty": ('Qty', '数量', '数量'),
    "detail.spec": ('Spec', '規格', '规格'),
    "detail.kindCol": ('Kind', '種別', '种类'),
    "detail.dateCol": ('Date', '日付', '日期'),
    "detail.completeEvent": ('Mark complete', '完了処理', '完成处理'),
    "hier.address": ('Hierarchy address', '階層アドレス', '层级地址'),
    "hier.groupCode": ('Group code', 'グループコード', '组代码'),
    "hier.groupName": ('Group name', 'グループ名', '组名'),
    "hier.groupType": ('Type', '種類', '类型'),
    "hier.status": ('Status', '状態', '状态'),
    "hier.groupUnit": ('groups', 'グループ', '组'),
    "subcode.groupNamePh": ('Group name', 'グループ名', '组名'),
    "subcode.newGroupCodePh": ('New group code', '新規グループコード', '新组代码'),
    "subcode.valuesPh": ('Values (· or , separated)', '値リスト (・/ , 区切り)', '值列表 (· 或 , 分隔)'),
    "subcode.saveF12": ('Save (F12)', '保存 (F12)', '保存 (F12)'),
    "subcode.item": ('Item', '項目名', '项目名'),
    "raw.addBtn": ('＋ Add material', '＋ 材質登録', '＋ 新增材质'),
    "raw.codeCol": ('Material code', '材質コード', '材质代码'),
    "raw.name": ('Material', '材質名', '材质名称'),
    "raw.type": ('Type', '種類', '类型'),
    "raw.density": ('Density', '密度', '密度'),
    "raw.standard": ('Spec', '規格', '规格'),
    "raw.hazardCol": ('Hazard', '危険', '危险'),
    "raw.editNamePh": ('New material name (edit)', '新しい材質名 (修正)', '新材质名 (修改)'),
    "raw.editDensityPh": ('New density', '新しい密度', '新密度'),
    "variant.addBtn": ('＋ Add value', '＋ 値登録', '＋ 新增值'),
    "variant.itemName": ('Item', '項目名', '项目名'),
    "variant.valueCode": ('Value code', '値コード', '值代码'),
    "variant.valueName": ('Value name', '値名', '值名称'),
    "variant.editNamePh": ('New value name (edit)', '新しい値名 (修正)', '新值名称 (修改)'),
    "dtable.addRow": ('＋ Add row', '＋ 行追加', '＋ 添加行'),
    "dtable.newKeyPh": ('New key', '新規 Key', '新 Key'),
    "di18n.statusCol": ('Status', '状態', '状态'),
    "master.group": ('Group', 'グループ', '组'),
    "master.approve": ('Approve', '承認', '批准'),
    "master.inactive": ('Deactivate', '非活性', '停用'),
    "quality.ruleCol": ('Rule', 'ルール', '规则'),
    "quality.macroCol": ('Macro expr', 'マクロ式', '宏表达式'),
    "quality.warningCol": ('Warning message', '警告メッセージ', '警告消息'),
    "quality.ruleNamePh": ('Rule name', 'ルール名', '规则名'),
    "quality.macroPh": ('Macro expr (=A>=…)', 'Macro 式 (=A>=…)', 'Macro 表达式 (=A>=…)'),
    "wp.minStockCol": ('Safety stock', '安全在庫', '安全库存'),
    "dwg.registerBtn": ('＋ Add drawing', '＋ 図面登録', '＋ 新增图纸'),
    "dwg.kindCol": ('Kind', '種別', '种类'),
    "dwg.dateCol": ('Date', '日付', '日期'),
    "parts.registerBtn": ('＋ Add part', '＋ 部品登録', '＋ 新增部件'),
    "grid.rowSelect": ('Select', '選択', '选择'),
    "purch.poConfirm": ('Confirm PO', 'PO 発注確定', 'PO 下单确认'),
    "purch.qcrNotePh": ('QCR note', 'QCR 備考', 'QCR 备注'),
    "purch.available": ('Avail', '利用可', '可用'),
    "purch.onHand": ('On hand', '保有', '在库'),
    "purch.reqDate": ('Req. date', '所要日', '需求日'),
    "purch.stockJudge": ('Stock check', '在庫判定', '库存判定'),
    "inv.onHand": ('On hand', '保有', '在库'),
    "wh.addLoc": ('＋ Add location', '＋ 位置登録', '＋ 新增位置'),
    "wh.locCodeCol": ('Loc code', '位置コード', '位置代码'),
    "wh.name": ('Location', '位置名', '位置名称'),
    "wh.hazardChip": ('Hazard', '危険', '危险'),
    "wh.hazardShort": ('Hazmat', '危険物', '危险品'),
    "company.registerBtn": ('＋ Add company', '＋ 取引先登録', '＋ 新增往来单位'),
    "company.importBtn": ('⬆ Bulk import', '⬆ 一括登録', '⬆ 批量导入'),
    "company.name": ('Company', '会社名', '公司名'),
    "company.nation": ('Country', '国', '国家'),
    "supp.delivery": ('Delivery', '納期', '交期'),
    "supp.note": ('Note', '備考', '备注'),
    "access.levelChange": ('Change level', 'レベル変更', '级别变更'),
    "access.inactivate": ('Deactivate', '非活性', '停用'),
    "access.addRole": ('＋ Role', '＋ ロール', '＋ 角色'),
    "access.newRolePh": ('New role name', '新しいロール名', '新角色名'),
    "price.closeToPh": ('End date YYYY-MM-DD', '終了日 YYYY-MM-DD', '结束日期 YYYY-MM-DD'),
    "procset.manual": ('Manual', '手動', '手动'),
    "audit.login": ('Login ID', '社員番号', '工号'),
    "ms.note": ('Note', '備考', '备注'),
    "ms.complete": ('Mark complete', '完了処理', '完成处理'),
    "act.qty": ('Qty', '数量', '数量'),
    "prj.addBtn": ('＋ Add project', '＋ プロジェクト登録', '＋ 新增项目'),
    "po.supplier": ('Supplier', '仕入先', '供应商'),
    "run.elapsed": ('Elapsed', '所要', '耗时'),
    "rel.addCodePh": ('e.g. KDI 21', '例: KDI 21', '例: KDI 21'),
    "rpt.catalog": ('Report catalog', 'レポートカタログ', '报表目录'),
}

def seed_v27(cur, tid: int) -> None:
    """N7 — Next 크롬 EN/JA/ZH 커버리지 1차 (라이브 스캔 잔존분)."""
    for key, (en, ja, zh) in UI_TRANSLATIONS_V27.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


UI_TRANSLATIONS_V28: dict[str, tuple[str, str, str]] = {
    # Next 포팅기(P1~P3)에 bundles.ts 에 직접 추가됐던 키 회수분 — 단일 진실(시드) 복원
    "arr.installNone": ('— none —', '— 未指定 —', '— 未指定 —'),
    "audit.action": ('Action', '操作', '操作'),
    "audit.after": ('After', '変更後', '变更后'),
    "audit.at": ('Timestamp', '日時', '时间'),
    "audit.before": ('Before', '変更前', '变更前'),
    "audit.by": ('Actor', '実行者', '执行者'),
    "audit.detail": ('Change detail — before / after', '変更詳細 — before / after', '变更详情 — before / after'),
    "audit.empty": ('No results — adjust filters and press F8', '結果なし — フィルタ調整後 F8', '无结果 — 调整筛选后按 F8'),
    "audit.period": ('Period', '期間', '期间'),
    "audit.reset": ('Reset', 'リセット', '重置'),
    "audit.selectedCsv": ('Selected CSV', '選択 CSV', '所选 CSV'),
    "audit.target": ('Target', '対象', '对象'),
    "audit.title": ('Audit log — sys_history (all-domain change history)', '監査ログ — sys_history (全ドメイン変更履歴)', '审计日志 — sys_history (全域变更历史)'),
    "audit.top500": ('top 500', '上位 500', '前 500'),
    "audit.user": ('User', 'ユーザー', '用户'),
    "codrel.ebomHint": ('EBOM Run — expands BOM-related items only (slide 70)', 'EBOM Run — BOM関連のみ展開（スライド70）', 'EBOM Run — 仅展开BOM相关项（幻灯片70）'),
    "codrel.openDetail": ('Code detail (Tech / Variant / Drawing)', 'コード詳細（Tech・Variant・図面）', '代码详情（Tech·Variant·图纸）'),
    "common.xlsxExport": ('Export full ledger to XLSX', '台帳全体をXLSXでエクスポート', '将台账导出为XLSX'),
    "cpq.detailSelect": ('Module detail selection', 'モジュール詳細選定', '模块详细选定'),
    "cpq.mirrorHint": ('Mirror selected block (MI)', '選択ブロック左右反転 (MI)', '镜像所选块 (MI)'),
    "cpq.rotateHint": ('Rotate selected block 90\\u00b0 (RO)', '選択ブロック90\\u00b0回転 (RO)', '旋转所选块90\\u00b0 (RO)'),
    "cpq.snapOff": ('Snap OFF', 'スナップ OFF', '捕捉 OFF'),
    "dash.deptDrillHint": ('Click = department task box (M-15-3)', 'クリック = 部門業務箱 (M-15-3)', '点击 = 部门任务箱 (M-15-3)'),
    "docmgmt.alloc": ('Auto number', '自動採番', '自动编号'),
    "docmgmt.allocHint": ('Rule-based auto numbering — {DEPT}·{TYPE}·{YYYY}·{SEQ} (slide 53)', 'ルールベース自動採番 — {DEPT}・{TYPE}・{YYYY}・{SEQ}（スライド53）', '基于规则的自动编号 — {DEPT}·{TYPE}·{YYYY}·{SEQ}（幻灯片53）'),
    "docmgmt.numberingRule": ('Numbering rule', '採番ルール', '编号规则'),
    "duct.airflow": ('Unit airflow', '装置風量', '设备风量'),
    "duct.bomItem": ('Item', '品目', '品目'),
    "duct.calcPdf": ('Calc sheet PDF', '計算書 PDF', '计算书 PDF'),
    "duct.calcTitle": ('Engineering calc sheet', '技術計算表', '技术计算表'),
    "duct.condensation": ('Condensation', '結露', '结露'),
    "duct.dewPoint": ('dew pt', '露点', '露点'),
    "duct.flowCompare": ('Flow compare', '風量比較', '风量对比'),
    "duct.hanger": ('hanger', 'ハンガー', '吊架'),
    "duct.load": ('Load', '荷重', '荷载'),
    "duct.perDiff": ('Per diffuser', 'ディフューザ当り', '每散流器'),
    "duct.pressureLoss": ('Pressure loss', '圧力損失', '压力损失'),
    "duct.purchaseLink": ('to purchasing', '購買連携', '采购联动'),
    "duct.risk": ('RISK', '危険', '风险'),
    "duct.safe": ('Safe', '安全', '安全'),
    "duct.size": ('Duct W\\u00d7H', 'ダクト W\\u00d7H', '风管 W\\u00d7H'),
    "duct.supplyT": ('Supply temp', '給気温度', '送风温度'),
    "duct.toPr": ('Material PR screen', '資材発注要請画面', '物料请购画面'),
    "duct.totalLen": ('Total duct length', 'ダクト総延長', '风管总长'),
    "duct.velocity": ('Velocity', '風速', '风速'),
    "editor.cycleFound": ('Cycle {n}', '循環 {n}', '循环 {n}'),
    "editor.cycleWarn": ('Circular reference \\u2014 wrong data extraction risk (slide 67)', '循環参照 \\u2014 誤データ抽出リスク', '循环引用 \\u2014 错误数据提取风险'),
    "editor.evalOrder": ('Eval order', '評価順序', '求值顺序'),
    "editor.noCycle": ('No cycle', '循環なし', '无循环'),
    "editor.priorityCheck": ('Priority / cycle check (U2)', '優先順位・循環点検 (U2)', '优先级·循环检查 (U2)'),
    "editor.priorityNote": ('Refs are token heuristics \\u2014 may include same-name Table column args', '参照はトークン推定 \\u2014 同名Table列引数を含む場合あり', '引用为令牌推定 \\u2014 可能包含同名Table列参数'),
    "inv.balance": ('Balance', '残量', '余量'),
    "inv.expiry": ('Expiry', '有効期限', '有效期'),
    "inv.lotPanel": ('Lot expiry', 'ロット有効期限', '批次有效期'),
    "inv.noLots": ('No lot stock', 'ロット在庫なし', '无批次库存'),
    "menu.common-direct": ('Common', '共通', '通用'),
    "menu.erp-mrp": ('MRP Requirements (M-8-5)', 'MRP 所要計画 (M-8-5)', 'MRP 需求计划 (M-8-5)'),
    "menu.erp-tenant-menu": ('Tenant Menu Admin (M-14-6B)', 'テナントメニュー管理 (M-14-6B)', '租户菜单管理 (M-14-6B)'),
    "menu.plm-direct": ('PLM Set-up', 'PLM Set-up', 'PLM Set-up'),
    "menu.toolbox-direct": ('EDIM Toolbox', 'EDIM Toolbox', 'EDIM Toolbox'),
    "mrp.code": ('Item code', '資材コード', '物料代码'),
    "mrp.dueDate": ('Earliest due', '最短納期', '最早交期'),
    "mrp.empty": ('No ORDERED material lines — convert a quotation in Sales Orders', 'ORDERED 資材ラインなし', '无 ORDERED 物料行'),
    "mrp.leadDays": ('Lead (days)', 'リード(日)', '提前期(天)'),
    "mrp.name": ('Name', '品名', '品名'),
    "mrp.onHand": ('On hand', '在庫', '库存'),
    "mrp.orderBy": ('Order by', '発注推奨日', '建议下单日'),
    "mrp.orderCount": ('Orders', '受注', '订单'),
    "mrp.orders": ('Orders', '関連受注', '相关订单'),
    "mrp.recalc": ('Recalculate', '再計算', '重新计算'),
    "mrp.required": ('Required', '所要量', '需求量'),
    "mrp.shortCount": ('Short items', '不足品目', '短缺品目'),
    "mrp.shortage": ('Short', '不足', '短缺'),
    "mrp.title": ('MRP Material Requirements', 'MRP 資材所要計画', 'MRP 物料需求计划'),
    "mrp.toPurchase": ('To purchase', '発注画面', '采购画面'),
    "panel.dept": ('Dept.', '部門', '部门'),
    "panel.file": ('File', 'ファイル', '文件'),
    "panel.name": ('Name', '名前', '名称'),
    "panel.preview": ('Preview (top 4)', 'プレビュー (上位4行)', '预览 (前4行)'),
    "panel.tools": ('Tools', 'ツール', '工具'),
    "panel.type": ('Type', '種類', '类型'),
    "parts.noSubst": ('No substitutes', '代替関係なし', '无替代关系'),
    "parts.note": ('Note', '備考', '备注'),
    "parts.substLink": ('Link', '連結', '关联'),
    "parts.substNo": ('Substitute', '代替品番', '替代件号'),
    "parts.substPh": ('Substitute part no', '代替部品番号', '替代件号'),
    "parts.substTitle": ('Substitute parts', '代替部品', '替代物料'),
    "printsetup.applyHint": ('Applied to Print Test·PDF·Office instantly', '設定は出力に即時反映', '设置即时应用于输出'),
    "printsetup.colorful": ('Color', 'カラー', '彩色'),
    "printsetup.footer": ('Footer', 'フッター', '页脚'),
    "printsetup.grayscale": ('Grayscale', 'モノクロ', '黑白'),
    "printsetup.landscape": ('Landscape', '横', '横向'),
    "printsetup.orientation": ('Orientation', '方向', '方向'),
    "printsetup.portrait": ('Portrait', '縦', '纵向'),
    "printsetup.printDialog": ('Print', '印刷', '打印'),
    "prj.commentPh": ('Write a comment (Enter)', 'コメント入力 (Enter)', '输入评论 (Enter)'),
    "prj.commentTitle": ('Project chat', 'プロジェクト対話', '项目沟通'),
    "prj.noComments": ('No comments yet', 'コメントなし', '暂无评论'),
    "prj.send": ('Post', '登録', '发布'),
    "qr.hint": ('QR for on-site scan — open this screen on mobile', '現場スキャン用QR — モバイルでこの画面を開く', '现场扫码QR — 在手机打开此画面'),
    "qr.scanHint": ('Scan with a mobile camera to open this screen (login required)', 'モバイルカメラでスキャンするとこの画面へ (ログイン必要)', '用手机相机扫码即可打开此画面 (需登录)'),
    "shell.addItem": ('Add items', '項目追加', '添加项目'),
    "shell.collapse": ('Collapse', '折りたたむ', '折叠'),
    "shell.doneItems": ('Done items', 'Done items', 'Done items'),
    "shell.expand": ('Expand', '展開', '展开'),
    "shell.favAdd": ('Add current screen to favorites', '現在の画面をお気に入りに追加', '收藏当前画面'),
    "shell.favNoScreen": ('Favorites — open a screen first', 'お気に入り — 先に画面を開いてください', '收藏 — 请先打开画面'),
    "shell.favRemove": ('Remove from favorites', 'お気に入り解除', '取消收藏'),
    "shell.favorite": ('Favorite', 'お気に入り', '收藏'),
    "shell.headEdit": ('Edit header menu — add / remove / reorder Head items', 'ヘッダーメニュー編集 — Head Item の追加・削除・並べ替え', '编辑标题菜单 — 添加·删除·重新排序 Head Item'),
    "shell.headEditTitle": ('Edit Header Menu', 'ヘッダーメニュー編集', '标题菜单编辑'),
    "shell.leftNavEmpty": ('No menus to show — ✎ Edit menu', '表示するメニューがありません — ✎ メニュー編集', '没有可显示的菜单 — ✎ 菜单编辑'),
    "shell.logoHint": ('Logo image for the title bar (PNG/SVG, under 48KB, height 18px scale)', 'タイトルバーのロゴ画像 (PNG/SVG, 48KB以下)', '标题栏Logo图片 (PNG/SVG, 48KB以下)'),
    "shell.logoRemove": ('Remove logo (default E mark)', 'ロゴ削除 (既定 E マーク)', '移除Logo (默认 E 标记)'),
    "shell.logoSetting": ('Company logo', '会社ロゴ設定', '公司Logo设置'),
    "shell.menuEdit": ('Edit menu', 'メニュー編集', '菜单编辑'),
    "shell.menuEditTitle": ('Edit left menu', '左メニュー編集', '左侧菜单编辑'),
    "shell.restoreDefault": ('Restore default', '既定に戻す', '恢复默认'),
    "shell.saveTenant": ('Save as tenant default', 'テナント既定として保存', '保存为租户默认'),
    "shell.saveTenantHint": ('Save this list as the tenant default — applied to all users without a personal setting (admin)', 'このリストをテナント既定として保存 — 個人設定のない全ユーザーに適用（管理者）', '将此列表保存为租户默认 — 应用于所有无个人设置的用户（管理员）'),
    "shell.scEsc": ('Close dialog', 'ダイアログを閉じる', '关闭对话框'),
    "shell.scFkeyCrud": ('New \\u00b7 Delete \\u00b7 Query (active screen)', '新規\\u00b7削除\\u00b7照会', '新建\\u00b7删除\\u00b7查询'),
    "shell.scFkeyRun": ('Run \\u00b7 Save (active screen)', 'Run\\u00b7保存', 'Run\\u00b7保存'),
    "shell.scGlobalSearch": ('Focus global search', '統合検索フォーカス', '聚焦全局搜索'),
    "shell.scGridFind": ('Find in grid', 'グリッド内検索', '网格内查找'),
    "shell.scTabClose": ('Close tab', 'タブを閉じる', '关闭标签'),
    "shell.scTabJump": ('Jump to tab N', 'タブN選択', '跳转标签N'),
    "shell.scTabMove": ('Move MDI tab', 'MDIタブ移動', '切换MDI标签'),
    "shell.schedule": ('Schedule', 'Schedule', 'Schedule'),
    "shell.scheduleHint": ('Open milestones', 'マイルストーンを開く', '打开里程碑'),
    "shell.shortcuts": ('Keyboard shortcuts', 'ショートカット', '快捷键'),
    "shell.shownItems": ('Shown items', '表示項目', '显示项目'),
    "shell.theme": ('Color theme', 'カラーテーマ', '颜色主题'),
    "shell.todoApprovalHint": ('Open approval inbox', '承認箱を開く', '打开审批箱'),
    "shell.todoPlHint": ('Open dashboard', 'ダッシュボードを開く', '打开仪表板'),
    "techdata.curveHint": ('Click a point to select the model (selected in red)', '点クリック=モデル選定 (選定点は赤)', '点击圆点=选定型号 (选定点红色)'),
    "techdata.curveTitle": ('Performance curve', '性能曲線', '性能曲线'),
    "uidsn.bindEmpty": ('No values', '値なし', '无值'),
    "uidsn.bindTest": ('Query', '照会', '查询'),
    "uidsn.columnPh": ('(column)', '（列）', '（列）'),
    "uidsn.dataPh": ('Data (field / parameter)', 'Data（フィールド・パラメータ）', 'Data（字段·参数）'),
    "uidsn.dblSetup": ('Double-click = Set-up (action / binding)', 'ダブルクリック＝Set-up（動作・バインディング）', '双击＝Set-up（动作·绑定）'),
    "uidsn.opNone": ('(none)', '（なし）', '（无）'),
    "uidsn.setupAction": ('Action (Commend set-up macro)', '動作（Commend set-up macro）', '动作（Commend set-up macro）'),
    "uidsn.setupBind": ('Combo Data set-up — table column binding', 'Combo Data set-up — テーブル列バインディング', 'Combo Data set-up — 表列绑定'),
    "uidsn.setupBtn": ('Set-up (action / binding)', 'Set-up（動作・バインディング）', 'Set-up（动作·绑定）'),
    "uidsn.setupTitle": ('Widget Set-up', 'Widget Set-up', 'Widget Set-up'),
    "uidsn.tablePh": ('(table)', '（テーブル）', '（表）'),
    "uidsn.targetPh": ('Target (table / screen / URL)', '対象（テーブル・画面・URL）', '目标（表·画面·URL）'),
    "wh.inspAdd": ('Record', '記録', '记录'),
    "wh.inspAt": ('At', '日時', '时间'),
    "wh.inspBy": ('By', '点検者', '巡检人'),
    "wh.inspNotePh": ('Inspection note', '点検備考', '巡检备注'),
    "wh.inspResult": ('Result', '判定', '判定'),
    "wh.inspTitle": ('Inspection records', '点検実績', '巡检记录'),
    "wh.noInsp": ('No inspections', '点検記録なし', '无巡检记录'),
    "wo.age": ('Age', '経過', '经过'),
    "wo.capBasis": ('persons\\u00d7480min/day approx.', '人員\\u00d7480分/日 近似', '人数\\u00d7480分/天 近似'),
    "wo.capTitle": ('Workshop capacity', '作業場 Capacity', '车间产能'),
    "wo.days": ('d', '日', '天'),
    "wo.daysNeeded": ('Needs', '所要', '需'),
    "wo.drawing": ('Drawing', '図面', '图纸'),
    "wo.min": ('min', '分', '分'),
    "wo.noCap": ('No load \\u2014 aggregated when Work Process times (U3) are set', '負荷なし \\u2014 工程工数(U3) 入力時に集計', '无负荷 \\u2014 录入工艺工时(U3)后汇总'),
    "wo.noOpen": ('No open work orders', '未完了作業指示なし', '无未完工单'),
    "wo.persons": ('p', '名', '人'),
    "wo.schedTitle": ('Open WO schedule', '未完了WOスケジュール', '未完工单排程'),
    "wo.workMin": ('Work', '工数', '工时'),
    "wp.designPriorityTitle": ('Design priority (slide 44)', '設計優先順位', '设计优先级'),
    "wp.dpBase": ('Base point', '設計基準点', '设计基准点'),
    "wp.dpError": ('Error check', '設計エラーチェック', '设计错误检查'),
    "wp.dpKind": ('Kind', '区分', '区分'),
    "wp.dpPriority": ('Design priority', '設計優先順位', '设计优先级'),
    "wp.dpUpper": ('Upper-design data', '上位設計優先資料', '上级设计优先资料'),
    "wp.paramHint": ('Inline-edit process params (workshop·persons·skill·W.Time·warehouse·stock) — F12 to save', '工程パラメータをインライン編集 — F12 保存', '内联编辑工艺参数 — F12 保存'),
    "wp.personCol": ('Persons', '人員', '人数'),
    "wp.workshopCol": ('Workshop', '作業場', '车间'),
}

def seed_v28(cur, tid: int) -> None:
    """N7 — Next 포팅기 직접 추가 번역 회수분 (bundles.ts → 시드 단일화)."""
    for key, (en, ja, zh) in UI_TRANSLATIONS_V28.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


UI_TRANSLATIONS_V29: dict[str, tuple[str, str, str]] = {
    "hier.title": ('Hierarchy Address (M-3-1)', '階層アドレス (M-3-1)', '层级地址 (M-3-1)'),
    "hier.nodeUnit": (' nodes', 'ノード', '节点'),
    "hier.parentPh": ('Parent address (empty=root)', '上位アドレス (なし=ルート)', '上级地址 (空=根)'),
    "hier.addressPh": ('Address (1.2.3)', 'アドレス (1.2.3)', '地址 (1.2.3)'),
    "hier.namePh": ('Node name', 'ノード名', '节点名称'),
    "hier.addNode": ('＋ Node', '＋ ノード', '＋ 节点'),
    "hier.searchPh": ('Search (name·address)', '検索 (名前・アドレス)', '搜索 (名称·地址)'),
    "hier.validateHint": ('Pre-save consistency check — dup address · orphan · parent mismatch (57-⑧)', '保存前整合チェック — アドレス重複・孤児・親不一致 (57-⑧)', '保存前一致性检查 — 地址重复·孤儿·父级不一致 (57-⑧)'),
    "hier.validate": ('Check', 'チェック', '检查'),
    "hier.renamePh": ('New name (rename)', '新しい名前 (改名)', '新名称 (改名)'),
    "hier.rename": ('Rename', '改名', '改名'),
    "po.collect": ('Collect', '着払い', '到付'),
    "po.land": ('Land', '陸上', '陆运'),
    "po.sea": ('Sea', '海上', '海运'),
    "po.air": ('Air', '航空', '空运'),
    "purch.code": ('Code', 'コード', '代码'),
    "purch.reqQty": ('Req', '所要', '需求'),
}

def seed_v29(cur, tid: int) -> None:
    """N7 2차 — 하드코딩 래핑분 (Hierarchy·발주 조건 옵션)."""
    for key, (en, ja, zh) in UI_TRANSLATIONS_V29.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


UI_TRANSLATIONS_V31: dict[str, tuple[str, str, str]] = {
    "kind.dwgApproval": ('Approval dwg', '承認図', '批准图'),
    "kind.quoteCost": ('Quote/Cost', '見積/原価', '报价/成本'),
    "kind.techData": ('Tech data', '技術資料', '技术资料'),
    "kind.received": ('Received', '受領資料', '接收资料'),
    "kind.upload": ('Upload', 'アップロード', '上传'),
    "kind.quoteApplied": ('Quote-applied', '見積適用', '报价适用'),
    "kind.quote": ('Quote', '見積', '报价'),
    "kind.stock": ('Stock', '在庫', '库存'),
    "kind.purchase": ('Purchase', '購買', '采购'),
    "stage.proposal": ('Proposal', '技術提案', '技术提案'),
    "stage.negotiation": ('Negotiation', '協議', '协商'),
    "stage.contract": ('Contract', '契約', '合同'),
    "stage.contractChange": ('Contract change', '契約変更', '合同变更'),
    "stage.execution": ('Execution', '遂行', '执行'),
    "stage.closed": ('Closed', '終了', '结束'),
    "stage.review": ('Review', '検討', '评审'),
    "enum.inProgress": ('In progress', '進行', '进行中'),
    "docstat.waitApprove": ('Awaiting approve', 'Approve 待ち', '待批准'),
    "docstat.drafting": ('Drafting', '作成中', '编写中'),
    "access.personUnit": (' users', '名', '名'),
}

def seed_v31(cur, tid: int) -> None:
    """N7 3차 — 서버 데이터 값(단계·종류·상태) 클라이언트 표시 매핑 키."""
    for key, (en, ja, zh) in UI_TRANSLATIONS_V31.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


UI_TRANSLATIONS_V30: dict[str, tuple[str, str, str]] = {
    # N7 — Next 포팅기 t() 신설 키 전수 시드 (라이브 스캔 후 709키 일괄 번역)
    "access.assetRole": ('Asset \\ Role', '資産 \\ ロール', '资产 \\ 角色'),
    "access.assetWord": ('Asset', '資産', '资产'),
    "access.cellCycle": ('Cell click = NONE→READ→WRITE', 'セルクリック = NONE→READ→WRITE', '点击单元格 = NONE→READ→WRITE'),
    "access.clickSelect": ('Click row = select', '行クリック=選択', '点击行=选择'),
    "access.deleteRole": ('Delete role', 'ロール削除', '删除角色'),
    "access.deleteRoleTip": ('Delete the custom role by entered name', '入力名のカスタムロールを削除', '删除输入名称的自定义角色'),
    "access.emptyUsers": ('No users', 'ユーザーがいません', '没有用户'),
    "access.loginId": ('Login ID', 'ログイン ID', '登录 ID'),
    "access.loginIdPh": ('Login ID', 'ログイン ID', '登录 ID'),
    "access.permMatrix": ('Permission matrix', '権限マトリクス', '权限矩阵'),
    "access.reactivate": ('Reactivate', '再活性', '重新启用'),
    "access.roleWord": ('role(s)', 'ロール', '角色'),
    "access.saveChanges": ('Save changes', '変更保存', '保存更改'),
    "access.selected": ('Selected', '選択', '已选'),
    "access.title": ('Users & Permissions', 'ユーザー・権限', '用户·权限'),
    "access.userLedger": ('User ledger', 'ユーザー台帳', '用户台账'),
    "access.userRegTitle": ('Register user', 'ユーザー登録', '用户注册'),
    "access.userUnit": ('user(s)', 'ユーザー', '用户'),
    "act.act": ('Actual', '実績', '实绩'),
    "act.amt": ('Amount', '金額', '金额'),
    "act.at": ('Recorded', '積載', '记录'),
    "act.cat": ('Category', '分類', '分类'),
    "act.est": ('Estimate (quote)', '推定(見積)', '估算(报价)'),
    "act.item": ('Item', '品目', '品目'),
    "act.proj": ('Project', 'プロジェクト', '项目'),
    "act.rate": ('Var. %', '差異率', '差异率'),
    "act.unit": ('Unit price', '単価', '单价'),
    "act.var": ('Variance', '差異', '差异'),
    "act.vcat": ('Cost category', '原価分類', '成本分类'),
    "anom.ack": ('Acknowledge', '確認', '确认'),
    "anom.at": ('Occurred', '発生', '发生'),
    "anom.escalateBtn": ('Escalate (unhandled HIGH)', 'エスカレーション (未処理 HIGH)', '升级 (未处理 HIGH)'),
    "anom.noEvents": ('No anomaly events', '異常イベントなし', '无异常事件'),
    "anom.ref": ('Ref', '参照', '参照'),
    "anom.resolve": ('Resolve', '解消', '解决'),
    "anom.scanBtn": ('Anomaly scan (cost·milestone)', '異常スキャン (原価・マイルストーン)', '异常扫描 (成本·里程碑)'),
    "anom.sev": ('Severity', '深刻度', '严重度'),
    "anom.source": ('Source', '出所', '来源'),
    "anom.status": ('Status', '状態', '状态'),
    "anom.title": ('Description', '内容', '内容'),
    "anom.transition": ('Transition', '遷移', '转移'),
    "appr.inboxHint": ('Open approval inbox', '承認箱を開く', '打开审批箱'),
    "appr.requestHint": ('Request approval — register to inbox (M-15-2) then decide (standard strip)', '承認要請 — 承認箱(M-15-2)登録後に決裁 (標準ストリップ)', '请求审批 — 登记到审批箱(M-15-2)后审批 (标准条)'),
    "appr.requested": ('Approval requested ✓ — registered to inbox', '承認要請 ✓ — 承認箱登録', '已请求审批 ✓ — 已登记审批箱'),
    "approval.empty": ('No pending approval requests', '承認待ちの要請はありません', '没有待审批的请求'),
    "approval.selectFirst": ('Select requests to decide', '決裁する要請を選択してください', '请选择要审批的请求'),
    "approval.selectedN": ('{n} selected', '選択 {n}件', '已选 {n} 件'),
    "approval.testPass": ('Pass', '通過', '通过'),
    "approval.waiting": ('Waiting', '待機', '等待'),
    "arr.canvasTitle": ('Arrangement layout — CAD source (SSR)', '構成配置 — CAD 正本 (SSR)', '配置布局 — CAD 正本 (SSR)'),
    "arr.codeCol": ('Arr. code', '構成 Code', '配置 Code'),
    "arr.codePh": ('Arr. code (ARR-…)', '構成 Code (ARR-…)', '配置 Code (ARR-…)'),
    "arr.compCodeCol": ('Code', 'コード', '代码'),
    "arr.compCodePh": ('Component code', '構成品コード', '组件代码'),
    "arr.compDelete": ('Delete component', '構成品削除', '删除组件'),
    "arr.compEmpty": ('No components', '構成品なし', '无组件'),
    "arr.compHeader": ('Components', '構成品', '组件'),
    "arr.compNameCol": ('Name', '名前', '名称'),
    "arr.componentsCol": ('Components', '構成品', '组件'),
    "arr.countUnit": (' items', '件', '件'),
    "arr.emptyList": ('No arrangement codes', '構成コードなし', '无配置代码'),
    "arr.layer": ('Layer', 'レイヤー', '图层'),
    "arr.nameCol": ('Name', '名前', '名称'),
    "arr.namePh": ('Name', '名前', '名称'),
    "arr.posCol": ('Position', '位置', '位置'),
    "arr.qtyCol": ('Qty', '数量', '数量'),
    "arr.regBtn": ('Add arrangement', '構成登録', '新增配置'),
    "arr.regTitle": ('Register arrangement', '構成登録', '配置登记'),
    "arr.selectHint": ('Click a row to manage components', '行をクリックして構成品を管理', '点击行以管理组件'),
    "arr.statusCol": ('Status', '状態', '状态'),
    "arr.title": ('Arrangement Set-Up (M-4-2)', 'Arrangement Set-Up (M-4-2)', 'Arrangement Set-Up (M-4-2)'),
    "assist.ask": ('Ask', '質問', '提问'),
    "assist.asking": ('Querying…', '照会中…', '查询中…'),
    "assist.code": ('Code', 'コード', '代码'),
    "assist.hint": ('Searches internal assets (product codes·docs·data tables·macros·parts) by question keywords and answers with references. AI-synthesized answers auto-enable when API credit is ready (live badge).', '質問キーワードで内部資産(製品コード・文書・データ Table・Macro・部品)を検索し根拠付きで回答。AI 合成回答は API クレジット準備時に自動有効化 (live バッジ)。', '按问题关键词搜索内部资产(产品代码·文档·数据表·宏·部件)并附依据回答。AI 合成回答将在 API 额度就绪后自动启用 (live 徽章)。'),
    "assist.kind": ('Kind', '種別', '种类'),
    "assist.none": ('No matching internal assets', '一致する内部資料なし', '无匹配的内部资料'),
    "assist.open": ('Open', '開く', '打开'),
    "assist.panel": ('Internal asset Q&A — grounded on code·doc·table·macro·part search', '内部資料 Q&A — コード・文書・Table・Macro・部品検索根拠', '内部资料问答 — 基于代码·文档·表·宏·部件搜索依据'),
    "assist.ph": ('e.g. Materials for KDCR 3-13? · Where is FanTechData?', '例: KDCR 3-13 関連資料は? · FanTechData はどこ?', '例: KDCR 3-13 相关资料? · FanTechData 在哪里?'),
    "assist.title": ('Internal Q&A', '社内 Q&A', '内部 Q&A'),
    "assist.titleCol": ('Name', '名前', '名称'),
    "bom.added": ('Added', '追加', '新增'),
    "bom.base": ('Base', '基準(base)', '基准(base)'),
    "bom.code": ('Code', 'コード', '代码'),
    "bom.qty": ('Qty', '数量', '数量'),
    "bom.removed": ('Removed', '削除', '删除'),
    "bom.target": ('Target', '対象(target)', '对象(target)'),
    "bomcmp.baseToTarget": ('Base→Target', '基準→対象', '基准→对象'),
    "bomcmp.changed": ('Changed', '変更', '变更'),
    "bomcmp.empty": ('No comparison — enter codes', '比較結果なし — コードを入力', '无比较结果 — 请输入代码'),
    "bomcmp.enterHint": ('Enter to compare', 'Enter で比較', '回车进行比较'),
    "bomcmp.hasDiff": ('Has differences', '差異あり', '有差异'),
    "bomcmp.identical": ('Identical BOM', '同一 BOM', 'BOM 相同'),
    "bomcmp.kept": ('Kept', '維持', '保留'),
    "bomcmp.nameCol": ('Name', '名前', '名称'),
    "bomcmp.title": ('BOM Compare (M-4-8)', 'BOM 比較 (M-4-8)', 'BOM 比较 (M-4-8)'),
    "cad.drawHint": ('Drag = draw · end/mid/center/intersection snap · Shift=Polar(45°) · Esc', 'ドラッグ=作図 · 端点/中点/中心/交点スナップ · Shift=Polar(45°) · Esc', '拖拽=绘制 · 端点/中点/中心/交点捕捉 · Shift=Polar(45°) · Esc'),
    "cad.edit": ('Edit', '編集', '编辑'),
    "cad.editHint": ('Drag entity=move · drag empty=box select · Shift+click=add · middle=pan · Delete=delete', 'エンティティドラッグ=移動 · 空白ドラッグ=ボックス選択 · Shift+クリック=追加 · 中クリック=パン · Delete=削除', '拖拽实体=移动 · 拖拽空白=框选 · Shift+点击=追加 · 中键=平移 · Delete=删除'),
    "cad.editTitle": ('Edit — drag entity to move · Delete to remove', '編集 — エンティティドラッグ移動 · Delete 削除', '编辑 — 拖拽实体移动 · Delete 删除'),
    "cad.gridTitle": ('Grid overlay', 'グリッドオーバーレイ', '网格叠加'),
    "cad.plotBtn": ('Scale PDF', '縮尺 PDF', '比例 PDF'),
    "cad.plotHint": ('1:scale vector PDF plot (A4 landscape)', '1:縮尺ベクター PDF 印刷 (A4 横)', '1:比例矢量 PDF 打印 (A4 横向)'),
    "cad.plotScale": ('Scale', '縮尺', '比例'),
    "cad.trimTitle": ('Trim/Extend — click boundary then target line end', 'トリム/延長 — 境界線クリック後に対象線端をクリック', '修剪/延伸 — 点击边界线后点击目标线端'),
    "cad.ucsTitle": ('UCS — click point as user origin (end/center snap, re-click=clear)', 'UCS — クリック点をユーザー原点に (端点・中心スナップ, 再クリック=解除)', 'UCS — 点击点设为用户原点 (端点·中心捕捉, 再点=清除)'),
    "cal.confirmDelete": ('Delete this holiday?', 'この休日を削除しますか?', '删除该节假日?'),
    "cal.date": ('Date', '日付', '日期'),
    "cal.dateCol": ('Date', '日付', '日期'),
    "cal.delete": ('Delete', '削除', '删除'),
    "cal.empty": ('No holidays', '休日がありません', '没有节假日'),
    "cal.name": ('Holiday name', '休日名', '节假日名称'),
    "cal.namePh": ('e.g. Liberation Day', '例: 光復節', '例: 光复节'),
    "cal.regTitle": ('Register holiday', '休日登録', '节假日登记'),
    "cal.register": ('＋ Register', '＋ 登録', '＋ 登记'),
    "cal.registering": ('Registering…', '登録中…', '登记中…'),
    "cal.title": ('Workday·Holiday Calendar', '稼働日・休日カレンダー', '工作日·节假日日历'),
    "codrel.diagram": ('Diagram — Mother · Child', '構成図 — Mother · Child', '构成图 — Mother · Child'),
    "common.backendError": ('Backend error', 'バックエンドエラー', '后端错误'),
    "common.clickRow": ('Click row = select', '行クリック=選択', '点击行=选择'),
    "common.redo": ('Redo', 'やり直し', '重做'),
    "common.register": ('Register', '登録', '登记'),
    "common.selected": ('Selected', '選択', '已选'),
    "common.status": ('Status', '状態', '状态'),
    "common.undo": ('Undo', '元に戻す', '撤销'),
    "company.grade": ('Rating', '評価等級', '评级'),
    "company.regTitle": ('Register company', '取引先登録', '往来单位登记'),
    "company.terms": ('Payment terms', '決済条件', '付款条件'),
    "company.termsPh": ('e.g. cash at month-end', '例: 月末現金', '例: 月末现金'),
    "company.type": ('Type', '種別', '类型'),
    "costact.alert": ('Alert', '警報', '警报'),
    "costact.empty": ('No actuals — record via form above', '実績なし — 上部フォームで積載', '无实绩 — 用上方表单录入'),
    "costact.itemPh": ('e.g. Casing', '例: Casing', '例: Casing'),
    "costact.optional": ('Optional', '任意', '可选'),
    "costact.projPh": ('Select PS-…', 'PS-… 選択', '选择 PS-…'),
    "costact.record": ('Record actual', '実績積載', '录入实绩'),
    "costact.recording": ('Recording…', '積載中…', '录入中…'),
    "costact.title": ('Cost Actuals · Variance', '原価実績・差異分析', '成本实绩·差异分析'),
    "costact.totalVar": ('Total variance', '総差異', '总差异'),
    "cpq.arrangement": ('Arrangement', '構成図 (Arrangement)', '构成图 (Arrangement)'),
    "cpq.bomExpandTitle": ('BOM expansion — {n} items (recursive CTE + slot_map)', 'BOM 展開 — {n}項目 (再帰 CTE + slot_map)', 'BOM 展开 — {n} 项 (递归 CTE + slot_map)'),
    "cpq.cadOffline": ('CAD server connection failed', 'CAD サーバー接続失敗', 'CAD 服务器连接失败'),
    "cpq.cmdBasePoint": ('Specify base point >', '基準点指定 >', '指定基点 >'),
    "cpq.cmdIdle": ('Command >', 'コマンド待機 >', '命令等待 >'),
    "cpq.cmdSelected": ('selected', '選択', '已选'),
    "cpq.drawing": ('Drawing…', '作図中…', '绘制中…'),
    "cpq.expanding": ('Expanding…', '展開中…', '展开中…'),
    "cpq.merge": ('Merge', '統合', '合并'),
    "cpq.mergeHint": ('Selected module + click target → merge (MERGE)', '選択モジュール + 対象クリック → 統合 (MERGE)', '所选模块 + 点击目标 → 合并 (MERGE)'),
    "cpq.noExpandItems": ('No expansion items', '展開項目なし', '无展开项'),
    "cpq.quote": ('Quotation', '見積案', '报价案'),
    "cpq.snapOn": ('Snap ON', 'スナップ ON', '捕捉 ON'),
    "cpq.specExcelHint": ('Spec Excel (Slot·Value 2 cols)', '仕様 Excel (Slot・Value 2列)', '规格 Excel (Slot·Value 两列)'),
    "cpq.split": ('Split', '分割', '拆分'),
    "cpq.splitHint": ('Split selected module ①·② (SPLIT, re-run=clear)', '選択モジュール分割 ①・② (SPLIT, 再実行=解除)', '拆分所选模块 ①·② (SPLIT, 重按=取消)'),
    "cpq.unpriced": ('Unpriced', '未確定', '未确定'),
    "dash.anlyAvg": ('Avg duration', '平均所要', '平均耗时'),
    "dash.anlyFail": ('Failed', '失敗', '失败'),
    "dash.anlyRate": ('Success rate', '成功率', '成功率'),
    "dash.anlyTotal": ('Total runs', '総 Run', '总 Run'),
    "dash.contribLabel": ('Contribution', '寄与', '贡献'),
    "dash.costDir": ('Direct expense', '直接経費', '直接费用'),
    "dash.costMat": ('Material', '材料費', '材料费'),
    "dash.costMfg": ('Manufacturing', '製造費', '制造费'),
    "dash.deptEventTitle": ('Events by department', '部署別 Event 状況', '部门事件状况'),
    "dash.legendContract": ('Contract amt', '契約額', '合同额'),
    "dash.legendContribMargin": ('Contrib. margin', '寄与マージン', '贡献毛利'),
    "dash.monthlyTitle": ('Monthly revenue·contribution margin — orders (ORDERED)', '月別売上・寄与マージン推移 — 受注(ORDERED)', '月度营收·贡献毛利趋势 — 订单(ORDERED)'),
    "dash.noActual": ('No actuals', '実績未積載', '无实绩'),
    "dash.noOrders": ('No order data', '受注データなし', '无订单数据'),
    "dash.runCostTitle": ('EDIM Run analysis · cumulative cost', 'EDIM Run 分析 · 累積原価', 'EDIM Run 分析 · 累计成本'),
    "dash.thresholdExceed": ('Over threshold', 'しきい値超過', '超过阈值'),
    "dash.title": ('ERP Dashboard', 'ERP Dashboard', 'ERP Dashboard'),
    "dash.totalVar": ('Total variance', '総差異', '总差异'),
    "dash.varCostTitle": ('Quote vs actual variance — cost', '見積 vs 実績差異 — 原価', '报价 vs 实绩差异 — 成本'),
    "design.bomEmpty": ('No assembly-seq BOM', '組立順序 BOM なし', '无装配顺序 BOM'),
    "design.none": ('None', 'なし', '无'),
    "design.relEmpty": ('No relations — cond1: vertical·horizontal·center / cond2: contact·coord·angle', '関係定義なし — 条件1: 垂直・水平・中心 / 条件2: 接触・座標・角度', '无关系定义 — 条件1: 垂直·水平·中心 / 条件2: 接触·坐标·角度'),
    "design.title": ('Design Editor (S-4-1-1) — KDCR 3-13', 'Design Editor (S-4-1-1) — KDCR 3-13', 'Design Editor (S-4-1-1) — KDCR 3-13'),
    "detail.actionCol": ('Action', '行為', '操作'),
    "detail.actionComment": ('Comment (handle/reassign)', '処理/再割当意見', '处理/重新分配意见'),
    "detail.active": ('Active', '適用中', '生效中'),
    "detail.actorCol": ('Actor', '処理者', '处理人'),
    "detail.allocFail": ('Allocation failed', '採番失敗', '编号失败'),
    "detail.alreadyDone": ('Already done', '既に完了', '已完成'),
    "detail.approvalHist": ('Approval history', '承認履歴', '审批历史'),
    "detail.assemblySeq": ('Assembly seq', '組立順序', '装配顺序'),
    "detail.attrs": ('Attributes', '属性', '属性'),
    "detail.autoAlloc": ('Auto allocate', '自動採番', '自动编号'),
    "detail.bindCol": ('Binding', 'バインディング', '绑定'),
    "detail.cadOpenHint": ('Open a drawing with ?fileId=<id> (Project Folder drill-down)', '?fileId=<id> で図面を開いてください (Project Folder ドリルダウン)', '请用 ?fileId=<id> 打开图纸 (Project Folder 下钻)'),
    "detail.cases": (' item(s)', '件', '件'),
    "detail.category": ('Category', '区分', '区分'),
    "detail.codeInfo": ('Code info', 'コード情報', '代码信息'),
    "detail.codeTitle": ('Code Detail', 'コード詳細', '代码详情'),
    "detail.deadline": ('Deadline', '期限', '期限'),
    "detail.delayed": ('Delayed', '遅延', '延迟'),
    "detail.dimBinding": ('Dimension binding', '寸法バインディング', '尺寸绑定'),
    "detail.docInfo": ('Document info', '文書情報', '文档信息'),
    "detail.drawingBom": ('Drawing BOM', '図面 BOM', '图纸 BOM'),
    "detail.entities": ('entities', 'エンティティ', '实体'),
    "detail.escalate": ('Escalate', 'エスカレーション', '升级'),
    "detail.eventActions": ('Event actions', 'イベント処理', '事件处理'),
    "detail.eventList": ('Event list', 'イベント一覧', '事件列表'),
    "detail.eventTitle": ('Event Detail', 'イベント詳細', '事件详情'),
    "detail.expired": ('Expired', '満了', '过期'),
    "detail.fileLabel": ('File', 'ファイル', '文件'),
    "detail.fromDate": ('Valid from', '適用開始', '生效开始'),
    "detail.itemCol": ('Item', '項目', '项目'),
    "detail.material": ('Material', '材質', '材质'),
    "detail.model3dHint": ('Product 3D viewer — GLB source embedded in original PPT (U29)', '製品 3D ビューア — 原本 PPT 内蔵 GLB 正本 (U29)', '产品 3D 查看器 — 原始 PPT 内嵌 GLB 正本 (U29)'),
    "detail.name": ('Name', '名前', '名称'),
    "detail.needBackend": ('Backend required', 'バックエンド必要', '需要后端'),
    "detail.noApprovalHist": ('No approval history', '承認履歴なし', '无审批历史'),
    "detail.noCol": ('No', '番号', '编号'),
    "detail.noEvent": ('No events', 'イベントなし', '无事件'),
    "detail.noFlow": ('No process flow info', '工程フロー情報なし', '无工序流程信息'),
    "detail.noNext": ('(no successor)', '(後続なし)', '(无后续)'),
    "detail.noPrev": ('(no predecessor)', '(先行なし)', '(无前置)'),
    "detail.noPrice": ('No registered price — warn target on pricing run', '登録単価なし — Pricing Run 時 warn 対象', '无登记单价 — Pricing Run 时为 warn 对象'),
    "detail.noRefs": ('No parent refs — top-level or unlinked code', '上位参照なし — 最上位または未連結コード', '无上级引用 — 顶级或未关联代码'),
    "detail.optional": ('Optional', '任意', '可选'),
    "detail.outputTitle": ('Output Doc Detail', '成果物文書詳細', '产出文档详情'),
    "detail.ownerCol": ('Owner', '担当', '负责人'),
    "detail.partNo": ('Part no', '部品番号', '零件编号'),
    "detail.partTitle": ('Part Detail', '部品詳細', '部件详情'),
    "detail.person": ('Persons', '人員', '人员'),
    "detail.priceHist": ('Price history', '単価履歴', '单价历史'),
    "detail.procFlow": ('Process flow', '工程フロー', '工序流程'),
    "detail.project": ('Project', 'プロジェクト', '项目'),
    "detail.reassign": ('Reassign', '再割当', '重新分配'),
    "detail.reassignId": ('Reassignee ID', '再割当担当 ID', '重新分配负责人 ID'),
    "detail.relatedCode": ('Related product code (drawing text·filename match, U10)', '関連製品コード (図面テキスト・ファイル名マッチ, U10)', '相关产品代码 (图纸文本·文件名匹配, U10)'),
    "detail.repDrawing": ('Rep. drawing', '代表図面', '代表图纸'),
    "detail.reqSlots": ('Required slot definition', '必須スロット定義', '必需槽位定义'),
    "detail.required": ('Required', '必須', '必需'),
    "detail.skill": ('Skill', '熟練度', '熟练度'),
    "detail.source": ('Source', 'ソース', '来源'),
    "detail.statusCol": ('Status', '状態', '状态'),
    "detail.statusFlow": ('Status flow', '状態フロー', '状态流'),
    "detail.titleCol": ('Title', 'タイトル', '标题'),
    "detail.toDate": ('Valid to', '終了', '结束'),
    "detail.typeLabel": ('Type', '種別', '类型'),
    "detail.valCol": ('Value', '値', '值'),
    "detail.weight": ('Weight', '重量', '重量'),
    "detail.workplace": ('Workplace', '作業場', '工作场所'),
    "detail.wtime": ('W.Time', '工数', '工时'),
    "di18n.editHint": ('Double-click translation cell → type → Enter (empty=delete)', '翻訳セルをダブルクリック → 入力 → Enter (空=削除)', '双击翻译单元格 → 输入 → Enter (空=删除)'),
    "di18n.empty": ('No items', '項目なし', '无项目'),
    "di18n.language": ('Language', '言語', '语言'),
    "di18n.saving": ('Saving…', '保存中…', '保存中…'),
    "di18n.source": ('Source', '原文', '原文'),
    "di18n.target": ('Target', '対象', '目标'),
    "di18n.title": ('Data i18n', 'データ多言語', '数据多语言'),
    "di18n.transUnit": (' translated', '翻訳', '翻译'),
    "di18n.translated": ('Translated', '翻訳済み', '已翻译'),
    "di18n.translation": ('Translation', '翻訳', '翻译'),
    "di18n.untranslated": ('Untranslated', '未翻訳', '未翻译'),
    "docmgmt.docNoPh": ('Doc no (DOC-…)', '文書番号 (DOC-…)', '文档编号 (DOC-…)'),
    "docmgmt.noDocs": ('No documents', '文書がありません', '无文档'),
    "docmgmt.pdfPreview": ('PDF preview', 'PDF プレビュー', 'PDF 预览'),
    "docmgmt.rowHint": ('Click row=select · double-click=detail', '行クリック=選択 · ダブルクリック=詳細', '点击行=选择 · 双击=详情'),
    "docmgmt.selected": ('Selected', '選択', '已选'),
    "docmgmt.status": ('Status', '状態', '状态'),
    "docmgmt.typePh": ('Type (DWG/QUO…)', '種別 (DWG/QUO…)', '类型 (DWG/QUO…)'),
    "dtable.bar": ('Bar', '棒', '柱状'),
    "dtable.chart": ('Chart', 'チャート', '图表'),
    "dtable.chartDlHint": ('Download chart SVG (for doc attachment)', 'チャート SVG ダウンロード (文書添付用)', '下载图表 SVG (用于文档附件)'),
    "dtable.chartDocTitle": ('Chart', 'チャート', '图表'),
    "dtable.chartHint": ('Chart wizard — column mapping → line/bar chart (slide 59)', 'グラフウィザード — 列マッピング → ライン/棒チャート (スライド 59)', '图表向导 — 列映射 → 折线/柱状图 (幻灯片 59)'),
    "dtable.chartPrintHint": ('Print chart — new window + OS print dialog', 'チャート印刷 — 新ウィンドウ + OS 印刷ダイアログ', '打印图表 — 新窗口 + 系统打印对话框'),
    "dtable.chartTitle": ('Chart wizard', 'グラフウィザード', '图表向导'),
    "dtable.delRow": ('Delete row', '行削除', '删除行'),
    "dtable.empty": ('No data rows', 'データ行なし', '无数据行'),
    "dtable.exportBtn": ('⬇ Export', '⬇ Export', '⬇ Export'),
    "dtable.importBtn": ('⬆ Import', '⬆ Import', '⬆ Import'),
    "dtable.line": ('Line', 'ライン', '折线'),
    "dtable.noSeries": ('Select series', 'シリーズを選択してください', '请选择系列'),
    "dtable.rowClickHint": ('Click row = open edit panel', '行クリック = 編集パネルを開く', '点击行 = 打开编辑面板'),
    "dtable.seriesPick": ('Series (numeric cols)', 'シリーズ (数値列)', '系列 (数值列)'),
    "duct.canvasTitle": ('Duct auto layout — CAD engine (SSR)', 'Duct 自動配置 — CAD 実エンジン (SSR)', '风管自动布置 — CAD 实引擎 (SSR)'),
    "duct.editHint": ('Manual adjust — materialize auto layout as editable (dwg_file) then move/delete/trim (U1/U2 edit track)', '手動調整 — 自動配置を編集対象化(dwg_file)後に移動/削除/トリム (U1/U2 編集トラック)', '手动调整 — 将自动布置转为可编辑(dwg_file)后移动/删除/修剪 (U1/U2 编辑轨)'),
    "duct.entity": ('Entities', 'エンティティ', '实体'),
    "dwg.detailPrefix": ('Drawing', '図面', '图纸'),
    "dwg.done": ('Done', '完了', '完成'),
    "dwg.drawingTypeLabel": ('Drawing type', '図面種別', '图纸类型'),
    "dwg.empty": ('No drawings', '図面がありません', '无图纸'),
    "dwg.kindApproval": ('APPROVAL', 'APPROVAL(承認用)', 'APPROVAL(批准用)'),
    "dwg.kindLabel": ('Kind', '区分', '区分'),
    "dwg.kindManufacturing": ('MANUFACTURING', 'MANUFACTURING(製作用)', 'MANUFACTURING(制作用)'),
    "dwg.kindStandard": ('STANDARD', 'STANDARD(標準)', 'STANDARD(标准)'),
    "dwg.pendingChip": ('Pending', '待機', '待定'),
    "dwg.supPanelTitle": ('Supersedure — replace this drawing with a new one', 'Supersedure — この図面を新図面に代替', 'Supersedure — 用新图纸替代此图纸'),
    "eco.at": ('Registered', '登録', '登记'),
    "eco.colTitle": ('Title', 'タイトル', '标题'),
    "eco.gridEmpty": ('No change requests (ECR)', '変更要請(ECR)なし', '无变更请求(ECR)'),
    "eco.impact": ('Impact', '影響', '影响'),
    "eco.newDrawingPh": ('New drawing no (if superseding)', '新図面番号 (代替時)', '新图纸编号 (替代时)'),
    "eco.reasonPh": ('Change reason', '変更事由', '变更事由'),
    "eco.regBtn": ('Register ECR', 'ECR 登録', '登记 ECR'),
    "eco.regTitle": ('Register ECR', 'ECR 登録', '登记 ECR'),
    "eco.screenTitle": ('Engineering Change (ECR/ECO)', '設計変更 (ECR/ECO)', '设计变更 (ECR/ECO)'),
    "eco.status": ('Status', '状態', '状态'),
    "eco.target": ('Target', '対象', '对象'),
    "eco.targetNoPh": ('Target no (drawing/code)', '対象番号 (図面/コード)', '对象编号 (图纸/代码)'),
    "eco.targetTypeLabel": ('Target type', '対象種別', '对象类型'),
    "eco.title": ('Change title', '変更タイトル', '变更标题'),
    "eco.typeCode": ('Code', 'コード', '代码'),
    "eco.typeDrawing": ('Drawing', '図面', '图纸'),
    "editor.blockCall": ('Insert', '呼び出し', '调用'),
    "editor.blockInsHint": ('Call from parent — Block INSERT at given coords', '上位呼び出し — 指定座標に Block INSERT', '上级调用 — 在指定坐标 Block INSERT'),
    "editor.blockNeedEdit": ('Run a CAD edit tool first to create an edit target (e.g. move)', '先に CAD 編集ツールを実行して編集対象を作成 (例: 移動)', '请先执行 CAD 编辑工具以创建编辑对象 (例: 移动)'),
    "editor.blockRegHint": ('Register selected entities as Block (replaced by in-place INSERT)', '選択エンティティを Block 登録 (原位置 INSERT 代替)', '将所选实体注册为 Block (原位置 INSERT 替代)'),
    "editor.paramSyncHint": ('Full parametric sync — auto redraw on dimension change (dim-part sync)', 'パラメトリック全面同期 — 寸法変更時自動再作図 (寸法-部品同期)', '参数化全面同步 — 尺寸变更时自动重绘 (尺寸-部件同步)'),
    "editor.redo": ('Redo (Ctrl+Y)', 'やり直し (Ctrl+Y)', '重做 (Ctrl+Y)'),
    "editor.undo": ('Undo (Ctrl+Z)', '元に戻す (Ctrl+Z)', '撤销 (Ctrl+Z)'),
    "fin.addFx": ('＋ FX rate', '＋ 為替', '＋ 汇率'),
    "fin.addTax": ('＋ Tax', '＋ 税', '＋ 税'),
    "fin.amountPh": ('Amount', '金額', '金额'),
    "fin.calc": ('Calculate', '計算', '计算'),
    "fin.calcTitle": ('Tax calculator (currency → tax·KRW conversion)', '税額計算機 (通貨 → 税額・KRW 換算)', '税额计算器 (货币 → 税额·KRW 换算)'),
    "fin.code": ('Code', 'コード', '代码'),
    "fin.codePh": ('Code', 'コード', '代码'),
    "fin.converted": ('Converted', '換算', '换算'),
    "fin.currency": ('Currency', '通貨', '货币'),
    "fin.delete": ('Delete', '削除', '删除'),
    "fin.emptyFx": ('No FX rates', '為替なし', '无汇率'),
    "fin.emptyTax": ('No tax codes', '税コードなし', '无税码'),
    "fin.fxTitle": ('FX master', '為替マスター', '汇率主数据'),
    "fin.noTax": ('No tax', '税なし', '无税'),
    "fin.rateKrw": ('Rate (KRW)', '為替 (KRW)', '汇率 (KRW)'),
    "fin.ratePct": ('Rate (%)', '税率 (%)', '税率 (%)'),
    "fin.ratePh": ('Rate (KRW)', '為替(KRW)', '汇率(KRW)'),
    "fin.taxAmount": ('Tax amount', '税額', '税额'),
    "fin.taxName": ('Tax name', '税名', '税名'),
    "fin.taxNamePh": ('Tax name', '税名', '税名'),
    "fin.taxTitle": ('Tax codes', '税コード', '税码'),
    "fin.title": ('Multi-currency·Tax Master', '多通貨・税マスター', '多币种·税务主数据'),
    "fin.total": ('Total', '合計', '合计'),
    "fin.validFrom": ('Valid from', '適用開始', '生效开始'),
    "fin.validFromPh": ('Valid-from date', '適用開始日', '生效开始日'),
    "folder.empty": ('No files', 'ファイルなし', '无文件'),
    "folder.typeCol": ('Type', '種別', '类型'),
    "hier.emptyGroups": ('No code groups', 'コードグループなし', '无代码组'),
    "inv.addResv": ('＋ Reserve', '＋ 予約', '＋ 预留'),
    "inv.at": ('At', '日時', '时间'),
    "inv.atpPanel": ('Available-to-promise (ATP)', '有効在庫 ATP', '可用库存 ATP'),
    "inv.available": ('Available', '有効', '可用'),
    "inv.emptyStock": ('No stock — shows after inbound', '在庫なし — 入庫処理で表示', '无库存 — 入库后显示'),
    "inv.inLabel": ('Inbound', '入庫', '入库'),
    "inv.inboundBtn": ('Inbound', '入庫', '入库'),
    "inv.inbounding": ('Processing…', '入庫中…', '入库中…'),
    "inv.item": ('Item', '品目', '品目'),
    "inv.itemCodePh": ('Item code', '品目コード', '品目代码'),
    "inv.itemPh": ('Item code', '品目コード', '品目代码'),
    "inv.location": ('Location', '位置', '位置'),
    "inv.movePanel": ('Movement history (lot trace)', '入出庫履歴 (Lot 追跡)', '出入库历史 (批次追踪)'),
    "inv.name": ('Item name', '品名', '品名'),
    "inv.noResvShort": ('No reservations', '予約なし', '无预留'),
    "inv.noTrace": ('No history', '履歴なし', '无历史'),
    "inv.optional": ('Optional', '任意', '可选'),
    "inv.pageTitle": ('Inventory', '在庫管理', '库存管理'),
    "inv.price": ('Unit price', '単価', '单价'),
    "inv.pricePh": ('STOCK auto', 'STOCK 自動', 'STOCK 自动'),
    "inv.qty": ('Qty', '数量', '数量'),
    "inv.recent": ('Recent', '最近', '最近'),
    "inv.refCol": ('Ref', '参照', '参照'),
    "inv.refPh": ('Ref (WO etc)', '参照 (WO 等)', '参照 (WO 等)'),
    "inv.release": ('Release', '解除', '释放'),
    "inv.reserved": ('Reserved', '予約', '预留'),
    "inv.resvPanel": ('Stock reservations', '在庫予約', '库存预留'),
    "inv.totalValue": ('Total value', '総評価額', '总估值'),
    "inv.type": ('Type', '区分', '区分'),
    "inv.unitPrice": ('Unit price', '単価', '单价'),
    "inv.updated": ('Updated', '更新', '更新'),
    "inv.value": ('Value', '評価額', '估值'),
    "m3d.title": ('Product 3D Viewer', '製品 3D ビューア', '产品 3D 查看器'),
    "macro.aiGen": ('AI generate', 'AI 生成', 'AI 生成'),
    "macro.aiGenHint": ('AI-04 Prompt→Macro — generate EDIM macro expr with Claude (sample w/o key/credit)', 'AI-04 Prompt→Macro — Claude で EDIM Macro 式生成 (キー/クレジットなしはサンプル)', 'AI-04 Prompt→Macro — 用 Claude 生成 EDIM 宏表达式 (无密钥/额度时为示例)'),
    "macro.applyCol": ('Apply', '適用', '适用'),
    "macro.editTitle": ('Edit macro', 'Macro 編集', '宏编辑'),
    "macro.empty": ('No macros', 'Macro なし', '无宏'),
    "macro.errorLabel": ('Error', 'エラー', '错误'),
    "macro.exprPh": ('=IF(A>0, A*B, 0) — Excel-compatible', '=IF(A>0, A*B, 0) — Excel 互換式', '=IF(A>0, A*B, 0) — Excel 兼容式'),
    "macro.fnInsertHint": ('Click = insert into expr', 'クリック = 式に挿入', '点击 = 插入表达式'),
    "macro.fnNone": ('No matching functions', '一致関数なし', '无匹配函数'),
    "macro.fnSearchPh": ('Search (e.g. sum, condition)', '検索 (例: 合計, 条件)', '搜索 (例: 求和, 条件)'),
    "macro.fnWizard": ('Function wizard', '関数ウィザード', '函数向导'),
    "macro.fnWizardHint": ('Function wizard — search·desc·insert (TBX-014)', '関数ウィザード — 検索・説明・挿入 (TBX-014)', '函数向导 — 搜索·说明·插入 (TBX-014)'),
    "macro.newHint": ('(click a row or enter a new name)', '(行クリックまたは新規名入力)', '(点击行或输入新名称)'),
    "macro.resultLabel": ('Result', '結果', '结果'),
    "macro.statusCol": ('Status', '状態', '状态'),
    "macro.testRun": ('▶ Test Run', '▶ Test Run', '▶ Test Run'),
    "master.addBtn": ('＋ Add code', '＋ コード登録', '＋ 新增代码'),
    "master.clickSelect": ('Click row = select', '行クリック=選択', '点击行=选择'),
    "master.codeCol": ('Code', 'コード', '代码'),
    "master.codePh": ('Code (KDP …)', 'コード (KDP …)', '代码 (KDP …)'),
    "master.createdAt": ('Registered', '登録日', '登记日'),
    "master.delete": ('Delete', '削除', '删除'),
    "master.empty": ('No product codes', '製品コードなし', '无产品代码'),
    "master.groupPh": ('Group (e.g. KOF)', 'グループ (KOF 等)', '组 (如 KOF)'),
    "master.name": ('Code name', 'コード名', '代码名称'),
    "master.refs": ('Refs', '参照', '引用'),
    "master.regTitle": ('Register product code', '製品コード登録', '产品代码登记'),
    "master.restore": ('Restore (DRAFT)', '復元(DRAFT)', '恢复(DRAFT)'),
    "master.selected": ('Selected', '選択', '已选'),
    "master.status": ('Status', '状態', '状态'),
    "master.title": ('Product Code Master', '製品コードマスター', '产品代码主数据'),
    "mobile.inboundBtn": ('Inbound (MI-002)', '入庫処理 (MI-002)', '入库处理 (MI-002)'),
    "mobile.itemCode": ('Item code', '品目コード', '品目代码'),
    "mobile.noApproval": ('No pending approval requests', '承認待ち要請なし', '无待审批请求'),
    "mobile.noTask": ('No tasks', '業務なし', '无任务'),
    "mobile.tabApproval": ('Approvals', '承認箱', '审批箱'),
    "mobile.tabInbound": ('Inbound', '入庫', '入库'),
    "mobile.tabTask": ('Tasks', '業務箱', '任务箱'),
    "mobile.title": ('Mobile preview', 'モバイルプレビュー', '移动端预览'),
    "ms.actualCol": ('Actual', '実績日', '实际日'),
    "ms.addBtn": ('＋ Add delivery', '＋ 納期登録', '＋ 新增交期'),
    "ms.delay": ('Status', '状態', '状态'),
    "ms.emptyList": ('No milestones', 'マイルストーンなし', '无里程碑'),
    "ms.header": ('Schedule·Milestones', '日程・マイルストーン', '日程·里程碑'),
    "ms.plannedCol": ('Planned', '計画日', '计划日'),
    "ms.plannedPh": ('Planned YYYY-MM-DD', '計画日 YYYY-MM-DD', '计划日 YYYY-MM-DD'),
    "ms.project": ('Project', 'プロジェクト', '项目'),
    "ms.projectPh": ('Project No', 'プロジェクト No', '项目 No'),
    "ms.regTitle": ('Register delivery', '納期登録', '交期登记'),
    "ms.stage": ('Stage', '段階', '阶段'),
    "ms.workdaysLeft": ('Workdays left', '営業日残', '剩余工作日'),
    "order.amount": ('Amount', '金額', '金额'),
    "order.customer": ('Customer', '顧客', '客户'),
    "order.orderDate": ('Order date', '受注日', '接单日'),
    "order.quoteNo": ('Quote no', '見積番号', '报价编号'),
    "order.stage": ('Stage', '段階', '阶段'),
    "order.status": ('Status', '状態', '状态'),
    "order.title": ('Sales Orders', '受注管理', '销售订单'),
    "panel.childDetail": ('Code detail (Tech·Variant·Drawing)', 'コード詳細 (Tech・Variant・図面)', '代码详情 (Tech·Variant·图纸)'),
    "panel.childErr": ('Query failed — check mother code', '照会失敗 — Mother コード確認', '查询失败 — 检查 Mother 代码'),
    "panel.importHint": ('Structured Excel import — row 1 header (Key+cols), duplicate keys update (E-4 Specification)', '定型 Excel Import — 1行ヘッダー(Key+列), Key 重複は更新 (E-4 Specification)', '结构化 Excel 导入 — 首行表头(Key+列), Key 重复则更新 (E-4 Specification)'),
    "panel.noChildren": ('No children', 'Child なし', '无 Child'),
    "panel.query": ('Query', '照会', '查询'),
    "parts.empty": ('No parts', '部品なし', '无部件'),
    "parts.regTitle": ('Register part', '部品登録', '部件登记'),
    "parts.stdCol": ('Std', '標準', '标准'),
    "parts.supNameCol": ('Supplier item name', '供給者品名', '供应商品名'),
    "parts.unitCol": ('Unit', '単位', '单位'),
    "po.amount": ('Amount', '金額', '金额'),
    "po.createBtn": ('＋ Create PO', '＋ 発注作成', '＋ 创建订单'),
    "po.createTitle": ('Create PO', '発注作成', '创建采购订单'),
    "po.emptyList": ('No purchase orders', '発注なし', '无采购订单'),
    "po.expected": ('Expected', '予定日', '预计日'),
    "po.expectedPh": ('Expected YYYY-MM-DD', '予定日 YYYY-MM-DD', '预计日 YYYY-MM-DD'),
    "po.item": ('Item', '品目', '品目'),
    "po.itemsPh": ('name,qty,price (one per line)\\ne.g. Impeller #450,2,120000', '品名,数量,単価 (1行1件)\\n例: インペラ #450,2,120000', '品名,数量,单价 (每行一条)\\n例: 叶轮 #450,2,120000'),
    "po.ledger": ('PO ledger', '発注台帳', '订单台账'),
    "po.note": ('Note', '備考', '备注'),
    "po.order": ('Order', '発注', '下单'),
    "po.orderDate": ('Order date', '発注日', '下单日'),
    "po.poNo": ('PO no', '発注番号', '订单编号'),
    "po.receiveBtn": ('Goods receipt (GR)', '入庫処理 (GR)', '入库处理 (GR)'),
    "po.receiveRate": ('GR rate', '入庫率', '入库率'),
    "po.receivedQty": ('Received', '既入庫', '已入库'),
    "po.recvQty": ('GR qty', '入庫数量', '入库数量'),
    "po.rem": ('Remaining', '残余', '剩余'),
    "po.selectHint": ('Click a row to manage lines·approval·GR', '行をクリックしてライン・承認・入庫(GR)を管理', '点击行以管理明细·审批·入库(GR)'),
    "po.status": ('Status', '状態', '状态'),
    "price.code": ('Code', 'コード', '代码'),
    "price.empty": ('No prices', '単価なし', '无单价'),
    "price.ledgerTitle": ('Price ledger', '単価台帳', '单价台账'),
    "price.validFromPh": ('Valid-from YYYY-MM-DD', '適用開始 YYYY-MM-DD', '生效开始 YYYY-MM-DD'),
    "prj.client": ('Client', '顧客社', '客户'),
    "prj.dueDate": ('Delivery', '納期', '交期'),
    "prj.empty": ('No projects', 'プロジェクトなし', '无项目'),
    "prj.itemLabel": ('Item', 'Item', 'Item'),
    "prj.ledgerTitle": ('Project ledger', 'プロジェクト台帳', '项目台账'),
    "prj.rowHint": ('Click a row to manage stage transition·delete', '行をクリックして営業段階遷移・削除を管理', '点击行以管理阶段转移·删除'),
    "prj.saveStage": ('Save stage', '段階保存', '保存阶段'),
    "prj.type": ('Type', '種別', '类型'),
    "procset.code": ('Code', 'コード', '代码'),
    "procset.empty": ('No process definitions', 'プロセス定義なし', '无流程定义'),
    "procset.title": ('Process definitions', 'プロセス定義', '流程定义'),
    "purch.empty": ('No PR items', '発注要請品目なし', '无采购申请品目'),
    "purch.selectedItems": ('{n} item(s) selected', '選択 {n}品目', '已选 {n} 个品目'),
    "purch.title": ('Purchase·PR', '購買・発注要請', '采购·请购'),
    "quality.activeCol": ('Active', '活性', '启用'),
    "quality.addRuleBtn": ('Add rule', 'ルール登録', '新增规则'),
    "quality.autoVerifyTitle": ('Auto verify measurements (D4)', '測定値自動判定 (D4)', '测量值自动判定 (D4)'),
    "quality.empty": ('No verification rules', '検証ルールなし', '无验证规则'),
    "quality.fail": ('Fail', '失敗', '失败'),
    "quality.measPh": ('e.g.\\nA=560\\nB=800\\nC=316', '例:\\nA=560\\nB=800\\nC=316', '例:\\nA=560\\nB=800\\nC=316'),
    "quality.pass": ('Pass', '通過', '通过'),
    "quality.rulesTitle": ('Verification rules', '検証ルール', '验证规则'),
    "quality.runVerify": ('Run verdict', '判定実行', '执行判定'),
    "quality.title": ('Design Verification Rules (D-4V)', '設計検証ルール (D-4V)', '设计验证规则 (D-4V)'),
    "quality.valueCol": ('Value', '値', '值'),
    "quality.verdictCol": ('Verdict', '判定', '判定'),
    "raw.editBtn": ('Edit', '修正', '修改'),
    "raw.empty": ('No materials', '材質なし', '无材质'),
    "raw.hazardPh": ('Hazard label', '危険表記', '危险标注'),
    "raw.regTitle": ('Register material', '材質登録', '材质登记'),
    "raw.standardPh": ('Spec (KS etc)', '規格 (KS 等)', '规格 (KS 等)'),
    "raw.title": ('Raw Material·GPI', 'Raw Material·GPI', 'Raw Material·GPI'),
    "rel.childUnit": ('child', 'child', 'child'),
    "rel.title": ('Code Relationship', 'Code Relationship', 'Code Relationship'),
    "rpt.bizType": ('Business type', '事業種別', '业务类型'),
    "rpt.breakdownTitle": ('Cost tree (slide 74)', '費用ツリー (スライド 74)', '费用树 (幻灯片 74)'),
    "rpt.directCost": ('Direct cost', '直接原価', '直接成本'),
    "rpt.finishedCode": ('Finished code', '完成品コード', '成品代码'),
    "rpt.kindsUnit": (' kinds', '種', '种'),
    "rpt.noPcr": ('No PCR (created after run cost confirm)', 'PCR なし (Run 原価確定後生成)', '无 PCR (Run 成本确认后生成)'),
    "rpt.pcrPdfHint": ('PCR profitability report PDF (RPT-07)', 'PCR 収益性レポート PDF (RPT-07)', 'PCR 盈利报告 PDF (RPT-07)'),
    "rpt.pcrTitle": ('PCR profitability report (RPT-07) — {n}', 'PCR 収益性レポート (RPT-07) — {n}件', 'PCR 盈利报告 (RPT-07) — {n} 件'),
    "rpt.revenue": ('Revenue', '売上', '营收'),
    "rpt.status": ('Status', '状態', '状态'),
    "run.action": ('Action', '行動', '操作'),
    "run.calc3": ('3 categories', '3分類', '3 分类'),
    "run.calcLines": ('Lines', '内訳', '明细'),
    "run.calcTotal": ('Total', '合計', '合计'),
    "run.calcType": ('Category', '分類', '分类'),
    "run.cleanup": ('Archive cleanup', '保管整理', '归档清理'),
    "run.costTitle": ('Cost detail (cst_calc)', '原価詳細 (cst_calc)', '成本明细 (cst_calc)'),
    "run.directTotal": ('Direct subtotal', '直接費計', '直接费小计'),
    "run.noCosts": ('Not recorded', '未積載', '未录入'),
    "run.noCostsHint": ('No cost detail — cst_calc recorded from this run onward', '原価詳細なし — この Run 実行分から cst_calc 積載', '无成本明细 — 自本次 Run 起录入 cst_calc'),
    "run.noQuotes": ('No confirmed quotes — confirm after PCR', '確定見積なし — PCR 生成後に見積確定', '无确认报价 — 生成 PCR 后确认报价'),
    "run.noTax": ('No tax', '税なし', '无税'),
    "run.out": ('Outputs', '成果物', '产出'),
    "run.outputs": ('Outputs', '成果物', '产出'),
    "run.pcrCreate": ('Create PCR', 'PCR 生成', '生成 PCR'),
    "run.pcrMargin": ('Contrib. margin', '寄与マージン', '贡献毛利'),
    "run.pcrTitle": ('PCR profitability → quotation (cst_pcr·cst_quotation)', 'PCR 収益性 → 見積 (cst_pcr·cst_quotation)', 'PCR 盈利 → 报价 (cst_pcr·cst_quotation)'),
    "run.pipelineTitle": ('Run Pipeline (C-1)', 'Run パイプライン (C-1)', 'Run 流水线 (C-1)'),
    "run.quoteConfirm": ('Confirm quote', '見積確定', '确认报价'),
    "run.quoteTax": ('Tax', '税額', '税额'),
    "run.quoteTotal": ('Amount', '金額', '金额'),
    "runs.byCol": ('By', '実行者', '执行人'),
    "runs.cleanupRun": ('Clean run', 'Run 整理', '清理 Run'),
    "runs.durCol": ('Dur.(s)', '所要(s)', '耗时(s)'),
    "runs.empty": ('No run history', 'Run 履歴なし', '无 Run 历史'),
    "runs.gcApply": ('Apply GC', 'GC 適用', '应用 GC'),
    "runs.gcPreview": ('MinIO GC preview', 'MinIO GC プレビュー', 'MinIO GC 预览'),
    "runs.keepLatest": ('Keep latest', '最新維持', '保留最新'),
    "runs.latest": ('Latest', '最新', '最新'),
    "runs.protectedHint": ('Latest/referenced runs are protected', '最新/参照 Run は保護されます', '最新/被引用 Run 受保护'),
    "runs.referenced": ('Referenced', '参照', '被引用'),
    "runs.selectHint": ('Click row = select', '行クリック=選択', '点击行=选择'),
    "runs.selected": ('Selected run', '選択 Run', '所选 Run'),
    "runs.startedCol": ('Started', '開始', '开始'),
    "runs.title": ('Run History·Cleanup', 'Run 履歴・整理', 'Run 历史·清理'),
    "search.codes": ('Codes', 'コード', '代码'),
    "search.docs": ('Docs', '文書', '文档'),
    "search.files": ('Files', 'ファイル', '文件'),
    "search.noResult": ('No results', '結果なし', '无结果'),
    "search.screens": ('Screens', '画面', '画面'),
    "shell.addFolder": ('Add folder', 'フォルダ追加', '添加文件夹'),
    "shell.changeBtn": ('Change', '変更', '变更'),
    "shell.changePw": ('Change password', 'パスワード変更', '修改密码'),
    "shell.clickHint": ('Click =', 'クリック =', '点击 ='),
    "shell.closeAllTabs": ('Close all tabs', '全タブを閉じる', '关闭所有标签'),
    "shell.closeTab": ('Close tab', 'タブを閉じる', '关闭标签'),
    "shell.currentPw": ('Current password', '現在のパスワード', '当前密码'),
    "shell.dataTable": ('Data Table (M-3-7)', 'データ Table (M-3-7)', '数据 Table (M-3-7)'),
    "shell.demoScenario": ('Demo scenario (PDF)', 'デモシナリオ (PDF)', '演示场景 (PDF)'),
    "shell.docsPortal": ('Docs portal', '文書ポータル (docs)', '文档门户 (docs)'),
    "shell.folderHint": ('Add folder — items below the marker belong to it (arrange with ↑↓)', 'フォルダ追加 — マーカー下の項目が所属 (↑↓ で配置)', '添加文件夹 — 标记下的项目归属其中 (用 ↑↓ 排列)'),
    "shell.folderPh": ('New folder name', '新フォルダ名', '新文件夹名'),
    "shell.newPw": ('New password', '新パスワード', '新密码'),
    "shell.nextTab": ('Next tab', '次のタブ', '下一标签'),
    "shell.prevTab": ('Previous tab', '前のタブ', '上一标签'),
    "so.backlog": ('Order backlog (ORDERED)', '受注残 (ORDERED)', '订单积压 (ORDERED)'),
    "so.clickSelect": ('Click row = select', '行クリック=選択', '点击行=选择'),
    "so.contractAmount": ('Contract amt', '契約額', '合同额'),
    "so.contractPh": ('Contract amount (order)', '契約金額 (受注)', '合同金额 (接单)'),
    "so.date": ('Date', '日付', '日期'),
    "so.deliveryPh": ('Delivery YYYY-MM-DD', '納期 YYYY-MM-DD', '交期 YYYY-MM-DD'),
    "so.emptyOrders": ('No orders', '受注なし', '无订单'),
    "so.emptyQuotes": ('No quotations', '見積なし', '无报价'),
    "so.expDelivery": ('Exp. delivery', '予想納期', '预计交期'),
    "so.lifecycleTitle": ('Quote lifecycle — DRAFT→SENT→ORDERED/LOST', '見積 Lifecycle — DRAFT→SENT→ORDERED/LOST', '报价生命周期 — DRAFT→SENT→ORDERED/LOST'),
    "so.lost": ('Lost', '失注', '丢单'),
    "so.orderConvert": ('Convert to order', '受注転換', '转为订单'),
    "so.orderUnit": (' orders', '受注', '订单'),
    "so.project": ('Project', 'プロジェクト', '项目'),
    "so.quoteAmount": ('Quote amt', '見積額', '报价额'),
    "so.selected": ('Selected', '選択', '已选'),
    "so.send": ('Send', '送付', '发送'),
    "subcode.dupFirstHint": ('Available after duplicate check passes', '重複検討通過後に要請可能', '重复检查通过后可请求'),
    "subcode.empty": ('No slots', 'Slot なし', '无槽位'),
    "subcode.exportBtn": ('⬇ Export', '⬇ Export', '⬇ Export'),
    "subcode.group": ('Group', 'グループ', '组'),
    "subcode.groupReg": ('Add group', 'グループ登録', '新增组'),
    "subcode.importBtn": ('⬆ Import', '⬆ Import', '⬆ Import'),
    "subcode.newItem": ('New item', '新規項目', '新项目'),
    "supp.evalHistory": ('Evaluation history', '評価履歴', '评估历史'),
    "supp.fulfillment": ('Fulfillment', '履行率', '履约率'),
    "supp.grade": ('Grade', '等級', '等级'),
    "supp.newEval": ('New evaluation', '評価登録', '新增评估'),
    "supp.noEval": ('No evaluations', '評価なし', '无评估'),
    "supp.period": ('Period', '期間', '期间'),
    "supp.poCount": ('POs', '発注', '订单数'),
    "supp.price": ('Price', '価格', '价格'),
    "supp.qty": ('Ordered/received', '発注/入庫', '下单/入库'),
    "supp.quality": ('Quality', '品質', '质量'),
    "supp.saveEval": ('Save evaluation', '評価保存', '保存评估'),
    "supp.scorecard": ('Supplier evaluation', '仕入先評価', '供应商评估'),
    "supp.selectHint": ('Click a SUPPLIER row to manage fulfillment metrics·evaluations', 'SUPPLIER 行をクリックして履行指標・評価を管理', '点击 SUPPLIER 行以管理履约指标·评估'),
    "supp.suggested": ('Suggested delivery score', '推奨配送スコア', '建议交付评分'),
    "supp.total": ('Total', '総点', '总分'),
    "task.alreadyDone": ('Task already completed', '既に完了した業務です', '任务已完成'),
    "task.complete": ('Mark complete', '完了処理', '完成处理'),
    "task.empty": ('No tasks', '業務なし', '无任务'),
    "task.processCol": ('Process', '工程', '工序'),
    "task.selectFirst": ('Select a task to complete', '完了処理する業務を選択してください', '请选择要完成的任务'),
    "task.selectHint": ('Click row=select · double-click=event detail', '行クリック=選択 · ダブルクリック=イベント詳細', '点击行=选择 · 双击=事件详情'),
    "task.selected": ('Selected', '選択', '已选'),
    "task.taskCol": ('Task', '業務', '任务'),
    "task.title": ('Task Inbox', '業務箱', '任务箱'),
    "techdata.densityPdf": ('Density calc PDF', '密度計算書 PDF', '密度计算书 PDF'),
    "techdata.engineTable": ('Engine performance table', 'エンジン性能表', '发动机性能表'),
    "techdata.fanPerfPdf": ('Fan performance PDF', 'Fan 性能表 PDF', 'Fan 性能表 PDF'),
    "techdata.model": ('Model', 'モデル', '型号'),
    "techdata.noData": ('No performance data for the condition', '該当条件の性能データなし', '无符合条件的性能数据'),
    "techdata.power": ('Power', '動力', '功率'),
    "techdata.rowHint": ('Click row=select · Enter to re-query', '行クリック=選定 · Enter 再照会', '点击行=选定 · 回车重新查询'),
    "techdata.sound": ('Sound', '騒音', '噪音'),
    "templet.defCol": ('Definition', '定義', '定义'),
    "templet.editTitle": ('Edit templet', 'Templet 編集', 'Templet 编辑'),
    "templet.empty": ('No templets', 'Templet なし', '无 Templet'),
    "templet.newHint": ('(click row or new)', '(行クリックまたは新規)', '(点击行或新建)'),
    "templet.statusCol": ('Status', '状態', '状态'),
    "templet.sysNoDelete": ('System templet cannot be deleted', 'システム Templet は削除不可', '系统 Templet 不可删除'),
    "templet.system": ('System', 'システム', '系统'),
    "templet.title": ('Templet Manager', 'Templet 管理', 'Templet 管理'),
    "templet.typeCol": ('Type', '種別', '类型'),
    "tmenu.addTitle": ('Add items', '項目追加', '添加项目'),
    "tmenu.allAdded": ('All items included', '全項目含む', '已包含全部项目'),
    "tmenu.clear": ('Clear (revert to full)', '指定解除(全体復帰)', '取消指定(恢复全部)'),
    "tmenu.clearMine": ('Clear my personal setting', '個人設定解除', '清除个人设置'),
    "tmenu.cleared": ('Tenant default cleared ✓ — full permitted menu restored', 'テナント既定解除 ✓ — 全権限メニュー復帰', '租户默认已清除 ✓ — 恢复全部权限菜单'),
    "tmenu.editTitle": ('Edit composition', '構成編集', '编辑构成'),
    "tmenu.empty": ('No items — add from the right', '項目なし — 右側から追加', '无项目 — 从右侧添加'),
    "tmenu.folderHint": ('Add folder — items below the marker belong to it (arrange with ↑↓)', 'フォルダ追加 — マーカー下の項目が所属 (↑↓ 配置)', '添加文件夹 — 标记下的项目归属其中 (↑↓ 排列)'),
    "tmenu.folderPh": ('New folder name', '新フォルダ名', '新文件夹名'),
    "tmenu.head": ('Header', 'ヘッダー', '顶部'),
    "tmenu.headMenu": ('Header dropdown', 'ヘッダードロップダウン', '顶部下拉'),
    "tmenu.left": ('Left', '左側', '左侧'),
    "tmenu.leftPanel": ('Left panel', '左パネル', '左侧面板'),
    "tmenu.legend": ('● = tenant default set · ○ = full menu. Applied: personal > tenant default > full.', '● = テナント既定指定 · ○ = 全メニュー。適用: 個人設定 > テナント既定 > 全体。', '● = 已设租户默认 · ○ = 全部菜单。生效: 个人设置 > 租户默认 > 全部。'),
    "tmenu.mineCleared": ('Personal setting cleared ✓ — tenant default now applies (after refresh)', 'この账号の個人設定解除 ✓ — テナント既定が適用 (更新後反映)', '此账号个人设置已清除 ✓ — 将应用租户默认 (刷新后生效)'),
    "tmenu.mineWarn": ('This account has a personal setting for this module, so the tenant default is not visible to you (personal > tenant)', 'この账号にはこのモジュールの個人設定があり、テナント既定は本人画面に表示されません (個人 > テナント優先)', '此账号存在该模块的个人设置, 租户默认不会显示在您的界面 (个人 > 租户优先)'),
    "tmenu.module": ('Module', 'モジュール', '模块'),
    "tmenu.saveTenant": ('🏢 Save tenant default', '🏢 テナント既定保存', '🏢 保存租户默认'),
    "tmenu.saved": ('Tenant default saved ✓ — {n} items (applies to users w/o personal setting)', 'テナント既定保存 ✓ — {n}項目 (個人設定なしの全ユーザー適用)', '租户默认已保存 ✓ — {n} 项 (适用于无个人设置的所有用户)'),
    "tmenu.scope": ('Target', '対象選択', '选择对象'),
    "tmenu.title": ('Tenant Menu Manager', 'テナントメニュー管理', '租户菜单管理'),
    "uidsn.aiHintShort": ('→ organize purpose / items / needed DB tables, then propose templet', '→ 用途 / 項目 / 必要 DB Table 整理後 Templet 提案', '→ 整理用途 / 项目 / 所需 DB 表后提出 Templet 建议'),
    "uidsn.previewCaption": ('{n} widgets — used as processing form after publish (approval) (tbx_ui_form v{v}).', 'ウィジェット {n}個 — 公開(承認)後に処理 Form として使用 (tbx_ui_form v{v})。', '{n} 个组件 — 发布(审批)后用作处理 Form (tbx_ui_form v{v})。'),
    "uidsn.previewRenderNote": ('(dynamic render, TBX-003)', '(動的レンダリング, TBX-003)', '(动态渲染, TBX-003)'),
    "uidsn.previewTitle": ('Form preview', 'Form プレビュー', 'Form 预览'),
    "variant.deprecate": ('Deprecate', '廃棄', '废弃'),
    "variant.editBtn": ('Edit', '修正', '修改'),
    "variant.empty": ('No constant values defined', '定義された定数値なし', '无已定义常量值'),
    "variant.group": ('Group', 'グループ', '组'),
    "variant.status": ('Status', '状態', '状态'),
    "variant.title": ('Variant Constants', 'バリアント定数', '变体常量'),
    "wh.deleteBtn": ('Delete', '削除', '删除'),
    "wh.empty": ('No warehouse locations', '倉庫位置なし', '无仓库位置'),
    "wh.inspChip": ('Inspect', '検査', '检验'),
    "wh.inspShort": ('Inspect', '検査', '检验'),
    "wh.pageTitle": ('Warehouse Locations', '倉庫位置', '仓库位置'),
    "wh.parentCode": ('Parent code', '上位コード', '上级代码'),
    "wh.parentPh": ('Parent code (empty=top)', '上位コード (なし=最上位)', '上级代码 (空=顶级)'),
    "wh.regTitle": ('Register location', '位置登録', '位置登记'),
    "wh.remarks": ('Note', '備考', '备注'),
    "wo.assignee": ('Assignee', '担当', '负责人'),
    "wo.assigneePh": ('Assignee ID', '担当 ID', '负责人 ID'),
    "wo.done": ('Done', '完了', '完成'),
    "wo.drawingNoPh": ('Drawing no', '図面番号', '图纸编号'),
    "wo.empty": ('No work orders', '作業指示なし', '无工单'),
    "wo.header": ('Work Orders', '作業指示', '工单'),
    "wo.issueBtn": ('＋ Issue', '＋ 発行', '＋ 签发'),
    "wo.issueTitle": ('Issue work order', '作業指示発行', '签发工单'),
    "wo.issuedDate": ('Issued', '指示日', '签发日'),
    "wo.projectNoPh": ('Project No', 'プロジェクト No', '项目 No'),
    "wo.start": ('Start', '着手', '开工'),
    "wo.status": ('Status', '状態', '状态'),
    "wo.titleCol": ('Title', 'タイトル', '标题'),
    "wo.titlePh": ('Work order title', '作業指示タイトル', '工单标题'),
    "wo.transition": ('Transition', '遷移', '转移'),
    "wp.drawingLabel": ('Drawing', '図面', '图纸'),
    "wp.empty": ('No material rows', '資材行なし', '无材料行'),
    "wp.itemCol": ('Material', '資材', '材料'),
    "wp.remarksCol": ('Note', '備考', '备注'),
    "wp.save": ('Save', '保存', '保存'),
    "wp.saving": ('Saving…', '保存中…', '保存中…'),
    "wp.supplierCol": ('Supplier', '仕入先', '供应商'),
    "wp.timeCol": ('W.Time', 'W.Time', 'W.Time'),
    "wp.title": ('Work Process MAKE/BUY (G3-c)', '作業工程 MAKE/BUY (G3-c)', '作业工序 MAKE/BUY (G3-c)'),
    "wp.warehouseCol": ('Warehouse', '倉庫', '仓库'),
    "xreview.commentPh": ('Review comment (required to reject)', '検討意見 (差戻し時必須)', '评审意见 (驳回时必填)'),
    "xreview.empty": ('No X-codes awaiting review', '検討待ちの X-code なし', '无待评审的 X-code'),
    "xreview.project": ('Project', 'プロジェクト', '项目'),
    "xreview.projectName": ('Project name', 'プロジェクト名', '项目名'),
    "xreview.requestedAt": ('Requested at', '要請日時', '请求时间'),
    "xreview.requestedBy": ('Requester', '要請者', '请求人'),
    "xreview.selectHint": ('Click a row to select review target', '行をクリックして検討対象を選択', '点击行以选择评审对象'),
    "xreview.slots": ('Composition (slots)', '構成(スロット)', '构成(槽位)'),
    "xreview.target": ('Target', '対象', '对象'),
    "xreview.waiting": ('Waiting', '待機', '等待'),
}

def seed_v30(cur, tid: int) -> None:
    """N7 — Next 포팅기 t() 신설 709키 일괄 시드."""
    for key, (en, ja, zh) in UI_TRANSLATIONS_V30.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


# 2.5 — 표준 업무 프로세스 단계명 (좌측 패널). 노드명은 테넌트가 자유롭게 바꿀 수 있으므로
# 키를 시드 이름 자체로 두고, 개명된 노드는 저장된 이름 그대로 노출(폴백)한다.
UI_TRANSLATIONS_V33: dict[str, tuple[str, str, str]] = {
    "process.node.영업·견적": ('Sales & Quotation', '営業・見積', '销售·报价'),
    "process.node.프로젝트 등록": ('Project Registration', 'プロジェクト登録', '项目登记'),
    "process.node.사양 선택 (CPQ)": ('Configuration (CPQ)', '仕様選択 (CPQ)', '规格选择 (CPQ)'),
    "process.node.견적 산출": ('Quotation Run', '見積算出', '报价计算'),
    "process.node.견적서 발행": ('Quotation Issue', '見積書発行', '报价单发行'),
    "process.node.설계·기술": ('Engineering', '設計・技術', '设计·技术'),
    "process.node.Sub Code 관리": ('Sub Code Management', 'Sub Code 管理', 'Sub Code 管理'),
    "process.node.제품 코드": ('Product Code', '製品コード', '产品编码'),
    "process.node.BOM 관계": ('BOM Relationship', 'BOM 関係', 'BOM 关系'),
    "process.node.도면 관리": ('Drawing Management', '図面管理', '图纸管理'),
    "process.node.생산·구매": ('Production & Purchasing', '生産・購買', '生产·采购'),
    "process.node.작업 지시": ('Work Order', '作業指示', '工作指令'),
    "process.node.소요 계획": ('Requirement Planning', '所要計画', '需求计划'),
    "process.node.발주": ('Purchase Order', '発注', '采购下单'),
    "process.node.재고": ('Inventory', '在庫', '库存'),
    "process.node.품질·출하": ('Quality & Shipping', '品質・出荷', '质量·出货'),
    "process.node.검사": ('Inspection', '検査', '检验'),
    "process.node.ERP Handoff": ('ERP Handoff', 'ERP Handoff', 'ERP Handoff'),
    "process.node.산출물 패키지": ('Output Package', '成果物パッケージ', '产出物包'),
    "process.node.공통": ('Common', '共通', '公共'),
    "process.node.승인함": ('Approval Inbox', '承認箱', '审批箱'),
    "process.node.Run 이력·Snapshot": ('Run History & Snapshot', 'Run 履歴・Snapshot', 'Run 历史·Snapshot'),
    # 1.5 정보그룹 라벨 — 백엔드 INFO_GROUPS 가 원천이라 화면 하드코딩 키가 없던 구간
    "info.group.cost": ('Cost (actual · PCR · breakdown)', '原価 (実績・PCR・明細)', '成本 (实绩·PCR·明细)'),
    "info.group.price": ('Unit price (purchase · sales)', '単価 (購買・販売)', '单价 (采购·销售)'),
    "info.group.quote": ('Quotation amount', '見積金額', '报价金额'),
    "info.group.partner": ('Customer · supplier name', '取引先・仕入先名', '客户·供应商名称'),
    # 2.2a S-1-1 Slot 값 승인 열
    "subcode.approveCol": ('Approve', '承認', '审批'),
    "subcode.approveHint": ('Approve pending values so they can be composed',
                            '未承認値を承認して組合せ対象にする', '审批未批准值以纳入组合'),
    "common.approve": ('Approve', '承認', '审批'),
    # 2.7 #40 — BOM 전개 근거 대조
    # 4.0 #14 Head Registry · 4.1 #16 Accordion Template Host
    "head.label": ('Head', 'Head', 'Head'),
    "head.hint": ('Head — work area visible per permission; selecting moves to its default screen',
                  'Head — 権限に応じて見える業務領域。選択すると既定画面へ移動',
                  'Head — 按权限可见的业务区域；选择后跳转至其默认画面'),
    "head.adminTitle": ('Head Management', 'Head 管理', 'Head 管理'),
    "head.code": ('Code', 'コード', '编码'),
    "head.name": ('Head name', 'Head 名', 'Head 名称'),
    "head.type": ('Type', '種別', '类型'),
    "head.minLevel": ('Min level', '最小レベル', '最低级别'),
    "head.status": ('Status', '状態', '状态'),
    "head.bindings": ('Bindings', 'バインディング', '绑定'),
    "head.order": ('Order', '順序', '顺序'),
    "head.addBtn": ('+ New Head', '＋ Head 登録', '＋ 新建 Head'),
    "head.regTitle": ('Register Head (DRAFT)', 'Head 登録 (DRAFT)', '登记 Head (DRAFT)'),
    "head.seed": ('Create standard Heads', '標準 Head 作成', '创建标准 Head'),
    "head.review": ('Request approval', '承認依頼', '提交审批'),
    "head.publish": ('Publish', '公開', '发布'),
    "head.withdraw": ('Withdraw (DRAFT)', '取り下げ (DRAFT)', '撤回 (DRAFT)'),
    "head.noCenter": ('No center binding — cannot publish (#19)',
                      'center バインディングなし — 公開不可 (#19)', '无 center 绑定 — 无法发布 (#19)'),
    "head.panelBind": ('Panel bindings', 'パネルバインディング', '面板绑定'),
    "head.panel": ('Panel', 'パネル', '面板'),
    "head.targetKind": ('Kind', '種別', '类型'),
    "head.targetRef": ('Target', '対象', '目标'),
    "head.bindLabel": ('Display name (optional)', '表示名 (任意)', '显示名称 (可选)'),
    "head.bindLabelCol": ('Display name', '表示名', '显示名称'),
    "head.bindAdd": ('+ Binding', '＋ バインディング', '＋ 绑定'),
    "head.noBindings": ('No bindings — at least one CENTER is required to publish',
                        'バインディングなし — 公開には CENTER が最低1つ必要',
                        '无绑定 — 发布至少需要一个 CENTER'),
    "head.publishedLocked": ('Published — withdraw to edit', '公開中 — 取り下げ後に編集', '已发布 — 撤回后编辑'),
    "head.empty": ('No Heads — start with standard Heads', 'Head なし — 標準 Head から開始', '暂无 Head — 从标准 Head 开始'),
    "head.clickSelect": ('Click a row to select', '行クリックで選択', '点击行选择'),
    # 표준 Head 이름 — 프로세스 단계명(2.5)과 같은 규약: 키는 시드 이름 자체,
    # 테넌트가 개명하면 저장된 이름 그대로 폴백된다
    "head.name.업무 (사용자)": ('Operations (User)', '業務 (ユーザー)', '业务 (用户)'),
    "head.name.설정 (Set-up)": ('Set-up', '設定 (Set-up)', '设置 (Set-up)'),
    "head.name.관리 (기업 관리자)": ('Administration', '管理 (企業管理者)', '管理 (企业管理员)'),
    "head.name.플랫폼 (EDIM)": ('Platform (EDIM)', 'プラットフォーム (EDIM)', '平台 (EDIM)'),
    # 4.6 #29 — Slot 매핑 편집기
    "subcode.gtShortSpec": ('Spec', '仕様', '规格'),
    "subcode.gtShortRaw": ('Raw', '原材料', '原材料'),
    "subcode.gtShortGpi": ('Purchased', '購買品', '采购品'),
    "subcode.gtShortProduct": ('Product', '製品', '产品'),
    "codrel.undefined": ('undefined', '未定義', '未定义'),
    "subcode.gtSpec": ('Specification', '仕様 (Specification)', '规格 (Specification)'),
    "subcode.gtRaw": ('Raw Material', '原材料 (Raw Material)', '原材料 (Raw Material)'),
    "subcode.gtGpi": ('Purchased (GPI)', '購買品 (GPI)', '采购品 (GPI)'),
    "subcode.gtProduct": ('Product', '製品 (Product)', '产品 (Product)'),
    "subcode.autoSlot": ('auto', '自動', '自动'),
    "subcode.autoSlotHint": ('Leave blank to auto-assign the Item Head (#26)',
                             '空欄なら Item Head を自動採番 (#26)', '留空则自动分配 Item Head (#26)'),
    "codrel.slotMapTitle": ('Slot mapping — Mother condition → Child expansion basis (#29)',
                            'Slot マッピング — Mother 選択条件 → Child 展開基準 (#29)',
                            'Slot 映射 — Mother 选择条件 → Child 展开基准 (#29)'),
    "codrel.childSlot": ('Child Slot', 'Child Slot', 'Child Slot'),
    "codrel.motherSlot": ('Inherit Mother Slot', 'Mother Slot 継承', '继承 Mother Slot'),
    "codrel.fixedValue": ('Fixed value', '固定値', '固定值'),
    "codrel.mapAdd": ('+ Mapping', '＋ マッピング', '＋ 映射'),
    "codrel.source": ('Value source', '値の出所', '值来源'),
    "codrel.inherit": ('Mother', 'Mother', 'Mother'),
    "codrel.fixed": ('Fixed', '固定', '固定'),
    "codrel.noMap": ('No mapping — child slots expand empty', 'マッピングなし — Child スロットは空で展開',
                     '无映射 — Child 槽位将为空展开'),
    "common.or": ('or', 'または', '或'),
    "head.design": ('Display', '表示', '显示'),
    "head.visible": ('Show in user list', 'ユーザー一覧に表示', '在用户列表显示'),
    "head.pin": ('Pin to top', '上部固定', '置顶'),
    "head.kpi": ('KPI (display)', 'KPI (表示)', 'KPI (显示)'),
    "head.kpi.runs": ('Runs in progress', '進行中 Run', '进行中 Run'),
    "head.kpi.approvals": ('Pending approvals', '承認待ち', '待审批'),
    "head.kpi.todos": ('My To-Do', '自分の To-Do', '我的待办'),
    "head.kpi.projects": ('Active projects', '進行プロジェクト', '活跃项目'),
    "head.kpi.handoffs": ('ERP pending', 'ERP 受信待ち', 'ERP 待接收'),
    "head.kpi.anomalies": ('Anomaly alerts', '異常警告', '异常告警'),
    "panel.templates": ('Template', 'Template', 'Template'),
    "panel.expandAll": ('Expand all', 'すべて展開', '全部展开'),
    "panel.collapseAll": ('Collapse all', 'すべて折りたたむ', '全部折叠'),
    "panel.allCollapsed": ('All templates collapsed — click a title to expand',
                           'すべて折りたたみ — タイトルをクリックで展開', '全部折叠 — 点击标题展开'),
    "basis.btn": ('Check BOM basis', '展開根拠の照合', '核对BOM依据'),
    "basis.hint": ('Which relationship revisions this Run expanded on · whether re-running yields the same BOM',
                   'この Run がどの関係 Revision で展開したか・再実行で同じ BOM か',
                   '此 Run 基于哪些关系修订展开·重跑是否得到相同 BOM'),
    "basis.title": ('BOM expansion basis', 'BOM 展開根拠', 'BOM 展开依据'),
    "basis.stable": ('Basis unchanged — re-run yields the same BOM', '根拠同一 — 再実行でも同じ BOM', '依据一致 — 重跑得到相同 BOM'),
    "basis.moved": ('Basis changed', '根拠が変更されました', '依据已变更'),
    "basis.edges": ('Relationships', '関係', '关系'),
    "basis.added": ('added', '追加', '新增'),
    "basis.removed": ('removed', '削除', '移除'),
    "basis.revised": ('Rev', 'Rev', 'Rev'),
    "basis.sameHint": ('The pinned relationship revision set matches the current one',
                       '固定された関係 Revision 集合が現在と同一です', '已固定的关系修订集合与当前一致'),
    # 3.3 #53 — Project Folder 역할 칩
    "kind.output": ('Deliverable', '成果物', '交付物'),
    "kind.source": ('Source drawing', '作図原本', '作图原件'),
    "folder.immutable": ('Run deliverable — cannot be edited or overwritten (immutable)',
                         'Run 成果物 — 編集・上書き不可（不変）', 'Run 交付物 — 不可编辑或覆盖（不变）'),
    "basis.movedHint": ("Even if the basis moved this Run's BOM does not change — it means a re-run may differ",
                        '根拠が動いてもこの Run の BOM は変わりません — 再実行で結果が変わり得るという意味です',
                        '依据变动不会改变此 Run 的 BOM — 意味着重跑结果可能不同'),
}

UI_TRANSLATIONS_V32: dict[str, tuple[str, str, str]] = {
    "common.cancel": ('Cancel', 'キャンセル', '取消'),
    "company.statusCol": ('Status', '状態', '状态'),
    "company.editTitle": ('Edit company', '取引先修正', '往来单位修改'),
    "company.empty": ('No companies', '取引先なし', '无往来单位'),
    "enum.active": ('Active', '有効', '有效'),
    "devreq.btnHint": ('Submit requirement (dev server only)', '要求事項受付 (開発サーバー専用)', '需求受理 (仅开发服务器)'),
    "devreq.title": ('Requirements — dev server only', '要求事項受付 — 開発サーバー専用', '需求受理 — 仅开发服务器'),
    "devreq.tabNew": ('New', '登録', '登记'),
    "devreq.tabList": ('List', '一覧', '列表'),
    "devreq.fldTitle": ('Title', 'タイトル', '标题'),
    "devreq.titlePh": ('e.g. add currency select to price dialog', '例: 単価登録ダイアログに通貨選択追加', '例: 在单价登记对话框添加货币选择'),
    "devreq.fldCategory": ('Category', '分類', '分类'),
    "devreq.fldPriority": ('Priority', '優先度', '优先级'),
    "devreq.fldContent": ('Details', '詳細内容', '详细内容'),
    "devreq.contentPh": ('What to change and how — include repro steps·expected behavior', '何をどう変えるか具体的に — 再現手順・期待動作を含む', '具体说明改什么怎么改 — 含复现步骤·期望行为'),
    "devreq.fldShot": ('Screenshot', 'スクリーンショット', '截图'),
    "devreq.pickImage": ('Pick image', '画像選択', '选择图片'),
    "devreq.pasteHint": ('or capture then Ctrl+V paste (png/jpg/gif/webp · 10MB)', 'またはキャプチャ後 Ctrl+V 貼り付け (png/jpg/gif/webp · 10MB)', '或截图后 Ctrl+V 粘贴 (png/jpg/gif/webp · 10MB)'),
    "devreq.remove": ('Remove', '除去', '移除'),
    "devreq.ctxNote": ('Current screen', '現在画面', '当前画面'),
    "devreq.ctxNote2": (' is saved as context. Submitted items are handled in batch rounds.', ' がコンテキストとして保存されます。受付分は処理ラウンドで一括反映されます。', ' 将作为上下文一并保存。受理的需求将在处理轮次中统一落实。'),
    "devreq.titleRequired": ('Enter a title', 'タイトルを入力してください', '请输入标题'),
    "devreq.saving": ('Registering…', '登録中…', '登记中…'),
    "devreq.loading": ('Loading…', '読み込み中…', '加载中…'),
    "devreq.empty": ('No requirements submitted.', '受付済み要求事項なし。', '暂无受理的需求。'),
    "devreq.colTitle": ('Title', 'タイトル', '标题'),
    "devreq.colCat": ('Cat.', '分類', '分类'),
    "devreq.colPri": ('Pri', '優先', '优先'),
    "devreq.colStatus": ('Status', '状態', '状态'),
    "devreq.colReq": ('Requester', '要請者', '请求人'),
    "devreq.colManage": ('Manage', '処理', '处理'),
    "devreq.noContent": ('(no details)', '(詳細なし)', '(无详细内容)'),
    "devreq.resolution": ('Resolution', '処理', '处理'),
    "shell.newPwConfirm": ('Confirm new password', '新パスワード確認', '确认新密码'),
    "shell.pwMismatch": ('New password confirmation does not match', '新パスワード確認が一致しません', '新密码确认不一致'),
    "search.companies": ('Suppliers', '仕入先', '供应商'),
    "search.warehouses": ('Warehouses', '倉庫', '仓库'),
    "search.users": ('Users', 'ユーザー', '用户'),
    "search.parts": ('Parts', '部品', '部件'),
    "search.projects": ('Projects', 'プロジェクト', '项目'),
    "audit.fieldCol": ('Field', 'フィールド', '字段'),
    "folder.exportPkg": ('⬇ Delivery package', '⬇ 納品パッケージ', '⬇ 交付包'),
    "folder.exportPkgHint": ('Customer delivery — outputs-only ZIP + manifest (internal received files excluded, E2)', '顧客納品用 — 成果物のみ ZIP + マニフェスト (内部受領資料除外, E2)', '客户交付用 — 仅产出 ZIP + 清单 (排除内部接收资料, E2)'),
    "cpq.blocksDxfDone": ('Block diagram DXF ⬇ (from-blocks engine)', 'ブロック図 DXF ⬇ (from-blocks エンジン)', '块图 DXF ⬇ (from-blocks 引擎)'),
    "cpq.blocksDxfFail": ('DXF download failed', 'DXF ダウンロード失敗', 'DXF 下载失败'),
    "cpq.blocksDxfHint": ('Download block diagram DXF (from-blocks engine)', 'ブロック図 DXF ダウンロード (from-blocks エンジン)', '下载块图 DXF (from-blocks 引擎)'),
    "prj.dupCheck": ('Dup check', '重複検討', '查重'),
    "prj.dupCheckHint": ('S-3-5 duplicate check — query same/similar names', 'S-3-5 重複検討 — 同一/類似名の実照会', 'S-3-5 查重 — 相同/相似名称实查询'),
    "prj.dupFound": ('Possible duplicates', '重複疑い', '疑似重复'),
    "prj.dupNone": ('No duplicates ✓', '重複なし ✓', '无重复 ✓'),
    "cad.dxfDlHint": ('Download source file', '原本ファイルダウンロード', '下载源文件'),
    "access.roles": ('Roles (multi, sys_user_role)', 'ロール (複数, sys_user_role)', '角色 (多个, sys_user_role)'),
    "macro.refsLoading": ('Loading refs…', '参照照会中…', '查询引用中…'),
    "macro.noRefs": ('No refs', '参照なし', '无引用'),
    "access.adminOnly": ('User registration is ADMIN only', 'ユーザー登録は ADMIN 専用', '用户注册仅限 ADMIN'),
    "approval.readOnly": ('Read-only', '読み取り専用', '只读'),
    "approval.viewMine": ('My requests', '自分の要請', '我的请求'),
    "approval.viewMineHint": ('Show only my pending requests', '自分が要請した待機件のみ表示', '仅显示我请求的待办'),
    "approval.emptyMine": ('No pending requests of mine', '自分が要請した待機件なし', '没有我请求的待办'),
    "parts.editTitle": ('Edit part', '部品修正', '部件修改'),
    "wh.editTitle": ('Edit location', '位置修正', '位置修改'),
    "wp.codingOffline": ('Evaluation unavailable — backend required', '評価不可 — バックエンド必要', '无法评估 — 需要后端'),
    "wp.codingResult": ('Mfg cost formula evaluated ✓', '製造費算式評価 ✓', '制造费公式评估 ✓'),
    "wp.minUnit": ('min', '分', '分'),
    "wp.rateLabel": ('Rate (KRW/h)', '賃率(円/h)', '工时费率(元/h)'),
    "wp.codingHint": ('Evaluate mfg cost formula — WT=ΣW.Time · PERSONS=Σpersons (DWG-021, ENG-01)', '製造費算式実評価 — WT=ΣW.Time · PERSONS=Σ人員 (DWG-021, ENG-01)', '制造费公式实评估 — WT=ΣW.Time · PERSONS=Σ人员 (DWG-021, ENG-01)'),
    "wh.inspCycle": ('Inspection cycle', '検査周期', '检验周期'),
    "wh.inspCyclePh": ('e.g. 6 months (empty = no change)', '例: 6ヶ月 (空 = 変更なし)', '例: 6个月 (空 = 不变更)'),
    "so.delHint": ('Only DRAFT quotations can be deleted (issued/approved protected)', 'DRAFT 見積のみ削除可能 (発行/承認は保護)', '仅可删除 DRAFT 报价 (已发行/已批准受保护)'),
    "cpq.delSelHint": ('Delete selected quotation draft — 409 protected if Run history exists', '選択した見積案を削除 — Run 履歴があれば 409 保護', '删除所选报价案 — 存在 Run 记录时 409 保护'),
    "rel.draftLabel": ('DRAFT (pre-approval — ✕ to withdraw)', 'DRAFT (承認前 — ✕ で回収)', 'DRAFT (批准前 — ✕ 撤回)'),
    "rel.draftDelHint": ('Delete DRAFT relation (approved relations protected)', 'DRAFT 関係を削除 (承認済み関係は保護)', '删除 DRAFT 关系 (已批准关系受保护)'),
    "access.invite": ('Invite', '招待', '邀请'),
    "access.inviteHint": ('Send invite — no mail server: in-app notification', '招待案内 — メールサーバー未設定: アプリ内通知', '发送邀请 — 未配置邮件服务器: 应用内通知'),
    "master.batchCount": (' selected', '件', '件'),
    "master.batchApply": ('Batch transition', '一括遷移', '批量转换'),
    "master.batchDelete": ('Batch delete', '一括削除', '批量删除'),
    "company.batchOn": ('Batch activate', '一括有効', '批量启用'),
    "company.batchOff": ('Batch deactivate', '一括無効', '批量停用'),
    "ms.sumOverdue": ('Overdue', '遅延', '逾期'),
    "ms.sumDueSoon": ('Due soon', '間近', '临近'),
    "cal.calcTitle": ('Workday calculator', '営業日計算機', '工作日计算器'),
    "cal.calcRange": ('Workdays in range', '区間営業日', '区间工作日'),
    "cal.calcDue": ('Due after N workdays', 'N営業日後の期日', 'N个工作日后到期'),
    "cal.workdaysOut": ('Workdays', '営業日', '工作日'),
    "cal.dayUnit": (' days', '日', '天'),
    "cal.dueOut": ('Due', '期日', '到期'),
    "cal.workdayUnit": (' workdays', '営業日', '个工作日'),
    "cal.calcHint": ('Excludes weekends & registered holidays', '週末・登録休日を除外', '不含周末及已登记假日'),
    "eco.detailTitle": ('ECO Detail', 'ECO 詳細', 'ECO 详情'),
    "eco.reason": ('Reason', '事由', '事由'),
    "eco.appliedAt": ('Applied', '適用', '应用'),
    "eco.impactTitle": ('Impact analysis (impact_data)', '影響分析 (impact_data)', '影响分析 (impact_data)'),
    "eco.noImpact": ('No impact data', '影響データなし', '无影响数据'),
    "dt.impactTitle": ('Impact', '影響度', '影响度'),
    "dt.noImpact": ('No referencing macros — safe to edit', '参照マクロなし — 自由に変更可能', '无引用宏 — 可自由修改'),
    "dt.impactHint": ('— referencing macros exist: value changes affect formula results', '— 参照マクロあり: 値変更は算式結果に影響', '— 存在引用宏: 修改值将影响公式结果'),
    "price.resolveTitle": ('Price resolve', '単価解決', '单价解析'),
    "price.resolveCodePh": ('Code (e.g. KDF 32)', 'コード (例: KDF 32)', '代码 (例: KDF 32)'),
    "price.resolveHint": ('Effective price at date — source priority (purchase > quote > standard)', '基準日の有効単価 — ソース優先順位 (購買>見積>標準)', '基准日有效单价 — 来源优先级 (采购>报价>标准)'),
    "rpt.actualCol": ('Actual', '実績', '实绩'),
    "rpt.actualHint": ('Recompute with actuals (D-6) — replace direct costs with purchase actuals', '実績反映再計算 (D-6) — 直接費を購買実績に置換', '按实绩重算 (D-6) — 直接费用替换为采购实绩'),
    "rpt.actualTitle": ('Actual-based recalculation (D-6)', '実績反映再計算 (D-6)', '实绩重算 (D-6)'),
    "rpt.noActual": ('No purchase actuals — estimate only', '購買実績なし — 推定のみ表示', '无采购实绩 — 仅显示估算'),
    "rpt.actualBasis": ('Actuals', '実績', '实绩'),
    "rpt.estimate": ('Estimate', '推定', '估算'),
    "detail.bomAdd": ('Add BOM line', 'BOM 追加', '添加 BOM'),
    "detail.bomPartPh": ('Part no (registered)', '部品番号 (台帳登録分)', '零件号 (已登记)'),
    "detail.bomNotePh": ('Assembly note', '組立備考', '装配备注'),
    "detail.bomDelHint": ('Delete BOM line', 'BOM 行を削除', '删除 BOM 行'),
    "notif.title": ('Notifications', '通知', '通知'),
    "notif.allRead": ('Mark all read', 'すべて既読', '全部已读'),
    "notif.announceBtn": ('Announce', 'お知らせ', '公告'),
    "notif.announceHint": ('Send announcement — in-app notification to all active users (ADMIN)', 'お知らせ送信 — 全ユーザーへアプリ内通知 (ADMIN)', '发送公告 — 向全体用户推送应用内通知 (ADMIN)'),
    "notif.announceTitlePh": ('Announcement title', 'お知らせタイトル', '公告标题'),
    "notif.announceLinkPh": ('Link (optional, /path)', 'リンク (任意, /パス)', '链接 (可选, /路径)'),
    "notif.announceSend": ('Send', '送信', '发送'),
    "notif.announceSent": ('Announcement sent ✓', 'お知らせ送信 ✓', '公告已发送 ✓'),
    "di18n.exportHint": ('Source + en/ja/zh translation table XLSX (untranslated = blank)', '原文+en/ja/zh 翻訳表 XLSX (未訳=空欄)', '原文+en/ja/zh 翻译表 XLSX (未译=空白)'),
    "di18n.importHint": ('Bulk import translation XLSX — headers ID·en/ja/zh, blank = no change', '翻訳 XLSX 一括インポート — ヘッダー ID·en/ja/zh, 空欄=変更なし', '批量导入翻译 XLSX — 表头 ID·en/ja/zh, 空白=不变更'),
    "di18n.importDone": ('Bulk import ✓', '一括インポート ✓', '批量导入 ✓'),
    "di18n.cellUnit": (' cells', 'セル', '单元格'),
    "di18n.rejected": ('rejected', '拒否', '拒绝'),
    "cpq.sessionReset": ('Session reset ✓ — temporary selections restored (saved drafts/outputs kept)', 'セッション初期化 ✓ — 一時選択を復元 (保存済み案・成果物は維持)', '会话重置 ✓ — 临时选择已还原 (已保存方案/产出保留)'),
    "cpq.sessionResetHint": ('CPQ Session Reset — restore temporary selections to defaults', 'CPQ Session Reset — 一時選択を既定値へ復元', 'CPQ 会话重置 — 临时选择还原为默认值'),
    "cpq.resetBtn": ('Reset session', 'セッション初期化', '会话重置'),
    "handoff.title": ('ERP Handoff — only approved packages are received', 'ERP Handoff — 承認済み Package のみ受信', 'ERP Handoff — 仅接收已批准的 Package'),
    "handoff.createBtn": ('Create Handoff', 'Handoff 作成', '创建 Handoff'),
    "handoff.createHint": ('Creates approval request after validation (BOM/cost/quotation) — re-creation makes a new version', '検証通過後に承認要求を作成 — 再作成は新 Version', '通过校验后创建审批请求 — 重新创建为新版本'),
    "handoff.grade": ('Validation', '検証', '校验'),
    "handoff.at": ('Created', '生成', '创建'),
    "handoff.acceptBtn": ('Accept in ERP', 'ERP 受信', 'ERP 接收'),
    "handoff.acceptHint": ('Accept into ERP — project work starts', 'ERP 受信 — プロジェクト業務開始', 'ERP 接收 — 项目业务开始'),
    "handoff.waitApproval": ('Awaiting approval', '承認待ち', '待审批'),
    "handoff.empty": ('No handoffs — create from a SUCCESS run', 'Handoff なし — SUCCESS Run から作成', '暂无 Handoff — 从 SUCCESS Run 创建'),
    "subcode.dryRun": ('Preview only', '検討のみ', '仅预览'),
    "subcode.dryRunHint": ('Diff review — preview add/update/reject without applying', 'Diff 検討 — 反映せず追加/更新/拒否をプレビュー', 'Diff 预览 — 不应用，仅查看新增/更新/拒绝'),
    "folder.pkgTitle": ('Output Package', 'Output Package', 'Output Package'),
    "folder.pkgOutputs": ('outputs', '成果物', '产出'),
    "shell.mfa": ('MFA setup (OTP)', 'MFA 設定 (OTP)', 'MFA 设置 (OTP)'),
    "mfa.on": ('Enabled', '有効', '已启用'),
    "mfa.off": ('Disabled', '無効', '未启用'),
    "mfa.pending": ('Setting up (inactive)', '設定中 (未有効)', '设置中 (未启用)'),
    "mfa.hint": ('Optional — when enabled, login requires a 6-digit TOTP code from your authenticator app', '任意 — 有効化するとログイン時に認証アプリの6桁コードが必要', '可选 — 启用后登录需输入验证器应用的6位代码'),
    "mfa.setupBtn": ('1) Issue secret', '① シークレット発行', '① 生成密钥'),
    "mfa.setupDone": ('Secret issued — register in authenticator, then activate with code', 'シークレット発行 — 認証アプリ登録後コードで有効化', '密钥已生成 — 注册到验证器后用代码激活'),
    "mfa.codePh": ('6-digit app code', 'アプリの6桁コード', '应用的6位代码'),
    "mfa.enableBtn": ('2) Activate', '② 有効化', '② 激活'),
    "mfa.disableBtn": ('Disable', '解除', '停用'),
    "hier.attrTitle": ('Edit attributes', '属性編集', '编辑属性'),
    "hier.remark": ('Remark', '備考', '备注'),
    "hier.color": ('Color', '色', '颜色'),
    "hier.colorNone": ('Default', '既定', '默认'),
    "hier.lock": ('Lock+save', 'ロック+保存', '锁定+保存'),
    "hier.unlock": ('Unlock+save', '解除+保存', '解锁+保存'),
    "hier.lockHint": ('Locked nodes are protected from edit/move/delete (409)', 'ロックノードは編集/移動/削除から保護 (409)', '锁定节点受编辑/移动/删除保护 (409)'),
    "hier.impactTitle": ('Impact analysis', '影響分析', '影响分析'),
    "hier.impactHint": ('Moving keeps referencing assets linked automatically; delete is blocked while references exist', '移動時は参照資産のリンクを自動維持、参照があると削除はブロックされます', '移动时自动保持引用资产链接；存在引用时删除将被阻止'),
    "hier.descendants": ('Descendants', '下位ノード', '子节点'),
    "hier.refTotal": ('Referencing assets', '参照資産', '引用资产'),
    "hier.noRefs": ('No referencing assets — safe to move or delete', '参照資産なし — 安全に移動・削除できます', '无引用资产 — 可安全移动或删除'),
    "detail.whereUsedDeep": ('Full where-used (multi-level)', '全遡及展開 (多段 Where-Used)', '完整反查 (多级 Where-Used)'),
    "tenant.title": ('Customer tenants (platform)', '顧客テナント管理 (プラットフォーム)', '客户租户管理 (平台)'),
    "tenant.unit": ('tenants', '社', '家'),
    "tenant.newTitle": ('Register customer (onboarding)', '顧客登録 (オンボーディング)', '客户注册 (入驻)'),
    "tenant.codePh": ('Code (latin)', 'コード (英字)', '代码 (英文)'),
    "tenant.namePh": ('Customer name', '顧客名', '客户名称'),
    "tenant.adminPh": ('Admin login', '管理者ID', '管理员工号'),
    "tenant.adminNamePh": ('Admin name', '管理者名', '管理员姓名'),
    "tenant.pwPh": ('Initial password (6+)', '初期パスワード (6文字+)', '初始密码 (6位+)'),
    "tenant.create": ('＋ Create customer', '＋ 顧客作成', '＋ 创建客户'),
    "tenant.hint": ('Creates the admin account and base hierarchy nodes (/C /M /T) so the tenant is usable immediately',
                    '管理者アカウントと基本ノード(/C /M /T)を同時に作成し、すぐ利用できます',
                    '同时创建管理员账号与基础节点(/C /M /T)，创建后即可使用'),
    "tenant.listTitle": ('Customer list', '顧客一覧', '客户列表'),
    "tenant.code": ('Code', 'コード', '代码'),
    "tenant.name": ('Customer', '顧客', '客户'),
    "tenant.status": ('Status', '状態', '状态'),
    "tenant.users": ('Users', 'ユーザー', '用户'),
    "tenant.codes": ('Product codes', '製品コード', '产品代码'),
    "tenant.projects": ('Projects', 'プロジェクト', '项目'),
    "tenant.created": ('Created', '作成', '创建'),
    "tenant.suspend": ('Suspend', '利用停止', '停用'),
    "tenant.resume": ('Resume', '利用再開', '恢复'),
    "tenant.platformHint": ('The platform tenant cannot be suspended', 'プラットフォームテナントは停止できません', '平台租户不可停用'),
    "info.matrixTitle": ('Information access (view · masking)', '情報アクセス権限 (閲覧・マスキング)', '信息访问权限 (查看·脱敏)'),
    "info.hint": ('Separate from work permissions — amounts and partners can be hidden while records stay viewable. Any mode other than full blocks downloads',
                  '作業権限とは別 — 閲覧可でも金額・取引先を隠せます。full 以外はダウンロード不可',
                  '与作业权限分离 — 可查看记录但隐藏金额与往来单位。full 以外模式禁止下载'),
    "info.group": ('Information group', '情報グループ', '信息组'),
    "info.mine": ('My mode', '自分のモード', '我的模式'),
    "info.tempTitle": ('Temporary access (time-boxed)', '一時閲覧 (期間限定)', '临时查看 (限期)'),
    "info.loginPh": ('Login', '社員番号', '工号'),
    "info.hours": ('Hours', '時間', '小时'),
    "info.reasonPh": ('Reason (audited)', '理由 (監査記録)', '事由 (审计记录)'),
    "info.grant": ('Grant', '一時付与', '临时授予'),
    "info.user": ('User', 'ユーザー', '用户'),
    "info.mode": ('Mode', 'モード', '模式'),
    "info.until": ('Expires', '期限', '到期'),
    "info.revoke": ('Revoke', '回収', '回收'),
    "info.noTemp": ('No grants yet', '付与履歴なし', '暂无授予记录'),
    "info.maskedHint": ('Masked — information access permission required', 'マスキング済 — 情報アクセス権限が必要', '已脱敏 — 需要信息访问权限'),
    "snap.title": ('Snapshot freeze · reproduce', 'Snapshot 固定・再現', 'Snapshot 固化·复现'),
    "snap.notePh": ('Note (optional)', 'メモ (任意)', '备注 (可选)'),
    "snap.freeze": ('Freeze', '固定', '固化'),
    "snap.hint": ('Once frozen the snapshot never changes — verification only reports drift from current data',
                  '固定後は不変 — 検証は現在データとの差分のみ報告',
                  '固化后内容不变 — 校验仅报告与当前数据的差异'),
    "snap.integrity": ('Integrity', '無欠性', '完整性'),
    "snap.intact": ('OK', '正常', '正常'),
    "snap.broken": ('Broken', '破損', '损坏'),
    "snap.source": ('Source', '原本', '原始'),
    "snap.exists": ('present', '存在', '存在'),
    "snap.deleted": ('deleted', '削除済', '已删除'),
    "snap.drift": ('Drift', '差分', '差异'),
    "snap.same": ('Identical to current data — fully reproducible', '現在データと同一 — そのまま再現可能', '与当前数据一致 — 可完整复现'),
    "snap.code": ('Snapshot', 'Snapshot', 'Snapshot'),
    "snap.project": ('Project', 'プロジェクト', '项目'),
    "snap.created": ('Frozen', '固定', '固化'),
    "snap.handed": ('linked', '連携', '关联'),
    "snap.verify": ('Verify', '検証', '校验'),
    "snap.empty": ('No snapshots yet — enter a run number and freeze', 'Snapshot なし — Run 番号を入力して固定', '暂无 Snapshot — 输入 Run 编号后固化'),
    "process.title": ('Work process', '業務プロセス', '业务流程'),
    "process.editHint": ('Edit process — steps, order, linked screen', 'プロセス編集 — 段階・順序・連携画面', '编辑流程 — 步骤·顺序·关联画面'),
    "process.toMenu": ('Switch to menu view', 'メニュー表示に切替', '切换到菜单视图'),
    "process.toProcess": ('Switch to process view', 'プロセス表示に切替', '切换到流程视图'),
    "process.empty": ('No work process defined yet', '業務プロセス未定義', '尚未定义业务流程'),
    "process.seed": ('Create standard process', '標準プロセス作成', '创建标准流程'),
    "process.emptyHint": ('After creating it you can freely change steps, order and linked screens with ✎',
                          '作成後 ✎ で段階・順序・連携画面を自由に変更できます',
                          '创建后可用 ✎ 自由修改步骤·顺序·关联画面'),
    "process.unbound": ('unlinked', '未連携', '未关联'),
    "process.up": ('Move up', '上へ', '上移'),
    "process.down": ('Move down', '下へ', '下移'),
    "process.rename": ('Rename · change screen', '名称・画面変更', '重命名·更换画面'),
    "process.renamePrompt": ('Step name', '段階名', '步骤名称'),
    "process.hrefPrompt": ('Linked screen path (e.g. /erp/projects, empty = group)', '連携画面パス (例 /erp/projects、空=グループ)', '关联画面路径 (例 /erp/projects，留空=分组)'),
    "process.addChild": ('Add sub-step', '下位段階追加', '添加子步骤'),
    "process.addRoot": ('Top-level step', '最上位段階', '顶层步骤'),
    "process.del": ('Delete', '削除', '删除'),
    "process.namePh": ('Step name', '段階名', '步骤名称'),
    "cpq.arrHint": ('Select arrangement — layout blocks are replaced', 'Arrangement 選択 — 配置ブロックが入れ替わります', '选择 Arrangement — 布局块将被替换'),
    "cpq.arrLoaded": ('blocks', 'ブロック', '构成块'),
    "cpq.arrEmpty": ('No components — showing default blocks', '構成なし — 既定ブロック表示', '无组件 — 显示默认块'),
    "access.tenantExport": ('Tenant export', 'テナント export', '租户导出'),
    "access.tenantExportHint": ('Tenant data export — core tables JSON ZIP (offboarding/backup, audited)', 'テナントデータ export — コアテーブル JSON ZIP (オフボーディング/バックアップ, 監査記録)', '租户数据导出 — 核心表 JSON ZIP (离场/备份, 记录审计)'),
}

def seed_v32(cur, tid: int) -> None:
    """F5/devreq 이식분 — 거래처 수정·요구사항 접수 UI 키."""
    for key, (en, ja, zh) in UI_TRANSLATIONS_V32.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


def seed_v33(cur, tid: int) -> None:
    """#28 — 일반 코드 그룹 GEN (Slot 미정의).

    Slot 이 정의된 그룹은 승인된 Sub Code 조합으로만 제품 코드를 만든다(자유텍스트 금지).
    반면 구매품·일회성 마스터 코드는 사양 조합의 산물이 아니므로 수기 등록 경로가 필요하다.
    GEN 은 그 정식 창구 — Slot 을 두지 않아 조합 대상에서 제외된다."""
    cur.execute("SELECT 1 FROM code_group WHERE tenant_id=%s AND group_code='GEN'", (tid,))
    if cur.fetchone():
        return
    cur.execute("SELECT address FROM sys_hierarchy WHERE tenant_id=%s AND address='/C'", (tid,))
    addr = "/C/GEN" if cur.fetchone() else "/C"
    cur.execute(
        """INSERT INTO code_group (tenant_id, group_code, group_name, group_type,
           hierarchy_address, approval_status)
           VALUES (%s,'GEN','일반 코드 (Slot 미정의)','PRODUCT',%s,'APPROVED')""", (tid, addr))
    logger.info("seed v33 — GEN 일반 코드 그룹 (수기 등록 창구, #28)")


def seed_v34(cur, tid: int) -> None:
    """2.5 — 표준 프로세스 단계명 번역 (2.0 좌측 패널 도입 후 EN/JA/ZH 미번역이던 구간)."""
    for key, (en, ja, zh) in UI_TRANSLATIONS_V33.items():
        for locale, text in (("en", en), ("ja", ja), ("zh", zh)):
            cur.execute(
                """UPDATE sys_translation SET text=%s
                   WHERE tenant_id=%s AND entity_type='UI' AND locale=%s AND field=%s""",
                (text, tid, locale, key))
            if cur.rowcount == 0:
                cur.execute(
                    """INSERT INTO sys_translation (tenant_id, locale, entity_type, entity_id, field, text)
                       VALUES (%s,%s,'UI',0,%s,%s)""", (tid, locale, key, text))


def seed_v26(cur, tid: int) -> None:
    """D2 — 데모 재고 (inv_stock, 멱등: 마이그레이션 0004 후·비어있을 때만)."""
    cur.execute("SELECT to_regclass('public.inv_stock')")
    if not cur.fetchone()[0]:
        return
    cur.execute("SELECT 1 FROM inv_stock WHERE tenant_id=%s LIMIT 1", (tid,))
    if cur.fetchone():
        return
    for code, name, loc, qty in [("EWT-3", "Water Trap", "GEN-A01", 5),
                                 ("KDC-1", "Casing 강판", "WH-A-GEN", 2)]:
        cur.execute(
            """INSERT INTO inv_stock (tenant_id, item_code, item_name, location_code, quantity)
               VALUES (%s,%s,%s,%s,%s)""", (tid, code, name, loc, qty))
        cur.execute(
            """INSERT INTO inv_movement (tenant_id, item_code, location_code, movement_type,
               quantity, ref_type) VALUES (%s,%s,%s,'IN',%s,'SEED')""", (tid, code, loc, qty))
    logger.info("seed v26 — 데모 재고 2건")
